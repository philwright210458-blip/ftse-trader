import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace

import pandas as pd

import ftse_live_mock as live
import ftse_trade_scanner as sc


def setup(ticker, score, direction="Long", conflict=False, stale=False):
    zone = SimpleNamespace(low=99.0, high=100.0, touch_count=4)
    return sc.TradeSetup(company=f"{ticker} plc", ticker=ticker, index_label="FTSE 100", setup_type="Breakout",
                         direction=direction, zone=zone, entry=100.0, stop=99.0, target=103.0, note="",
                         quality_score=score, conflict=conflict, stale=stale)


def info(run_time, day="2026-10-19", mode="watch"):
    data_time = pd.Timestamp(f"{day} 10:00", tz=sc.UK_TZ)
    return sc.RunInfo(data_time, "10:15", day, run_time, mode, 15, 0)


class LiveMockTests(unittest.TestCase):
    def setUp(self):
        self.dir = Path(tempfile.mkdtemp())

    def test_best_scores_first_up_to_daily_limit_one_per_share(self):
        rows = live.record_live_mock(self.dir, info("2026-10-19 10:18:00"),
                                     [setup("AAA.L", 60), setup("BBB.L", 80), setup("AAA.L", 90, "Short"),
                                      setup("CCC.L", 70, conflict=True), setup("DDD.L", 95, stale=True)], 2)
        self.assertEqual([r["ticker"] for r in rows], ["AAA.L", "BBB.L"])
        self.assertEqual(rows[0]["direction"], "Short")       # highest-scoring idea for AAA
        later = live.record_live_mock(self.dir, info("2026-10-19 10:33:00"), [setup("EEE.L", 99)], 2)
        self.assertEqual(later, [])                            # daily limit already reached
        self.assertEqual(len(live.read_live(self.dir / live.LIVE_FILENAME)), 2)

    def test_numbering_continues_through_the_day(self):
        live.record_live_mock(self.dir, info("2026-10-19 10:03:00", mode="full"), [setup("AAA.L", 60)], 3)
        rows = live.record_live_mock(self.dir, info("2026-10-19 11:01:00"), [setup("AAA.L", 70), setup("BBB.L", 50)], 3)
        self.assertEqual([(r["ticker"], r["trade_no"]) for r in rows], [("BBB.L", 2)])

    def test_no_entries_after_cut_off_or_on_a_later_day(self):
        self.assertEqual(live.record_live_mock(self.dir, info("2026-10-19 15:46:00"), [setup("AAA.L", 60)]), [])
        self.assertEqual(live.record_live_mock(self.dir, info("2026-10-20 20:28:00"), [setup("AAA.L", 60)]), [])

    def test_key_matches_outcome_tracker_key(self):
        run = info("2026-10-19 10:18:00")
        row = live.record_live_mock(self.dir, run, [setup("AAA.L", 60)])[0]
        expected = run.data_time.isoformat() + "|AAA.L|Long|2026-10-19 10:18:00|watch"
        self.assertEqual(row["key"], expected)



class MockProgressTests(unittest.TestCase):
    def candles(self, rows):
        index = pd.date_range("2026-10-21 10:00", periods=len(rows), freq="15min", tz="Europe/London")
        return pd.DataFrame(rows, columns=["open", "high", "low", "close"], index=index)

    def trade(self, direction="Long"):
        return {"signal_time": "2026-10-21 10:03:48", "direction": direction, "stop": 99.0, "target": 103.0}

    def test_open_trade_uses_next_candle_open_and_latest_close(self):
        from ftse_live_mock import mock_progress
        data = self.candles([(100, 100.2, 99.8, 100), (100.1, 100.6, 99.9, 100.5), (100.5, 101.0, 100.2, 100.8)])
        p = mock_progress(self.trade(), data)
        self.assertEqual((p["fill"], p["status"], p["price"]), (100.1, "Open", 100.8))
        self.assertAlmostEqual(p["net"], round((100.8 - 100.1) / 100.1 * 100 - 0.1, 2))

    def test_stop_counted_first_when_both_hit(self):
        from ftse_live_mock import mock_progress
        data = self.candles([(100, 100.2, 99.8, 100), (100.0, 103.5, 98.5, 101)])
        p = mock_progress(self.trade(), data)
        self.assertEqual((p["status"], p["price"], p["time"]), ("Stopped out", 99.0, "10:30"))

    def test_sell_target_and_gap(self):
        from ftse_live_mock import mock_progress
        short = {"signal_time": "2026-10-21 10:03:48", "direction": "Short", "stop": 101.0, "target": 97.0}
        hit = mock_progress(short, self.candles([(100, 100, 100, 100), (100, 100.5, 96.5, 97)]))
        self.assertEqual((hit["status"], hit["price"], hit["costs"]), ("Target hit", 97.0, 0.10))
        gap = mock_progress(short, self.candles([(100, 100, 100, 100), (101.5, 102, 101.2, 101.6)]))
        self.assertTrue(gap["status"].startswith("Skipped"))

    def test_waiting_and_no_data(self):
        from ftse_live_mock import mock_progress
        self.assertEqual(mock_progress(self.trade(), self.candles([(100, 100, 100, 100)]))["status"], "Waiting for entry candle")
        self.assertEqual(mock_progress(self.trade(), None)["status"], "Price check unavailable")

    def test_scan_workbook_has_only_the_mock_sheet(self):
        import ftse_trade_scanner as sc
        info = sc.RunInfo(day="2026-10-21", as_of="10:30", data_time=pd.Timestamp("2026-10-21 10:15", tz=sc.UK_TZ),
                          run_time="2026-10-21 10:31:00", mode="watch", scanned=15, failed=0)
        row = {"trade_no": 1, "signal_time": "2026-10-21 10:03:48", "company": "Alpha plc", "ticker": "A.L",
               "setup": "Breakout", "direction": "Long", "stop": 99.0, "target": 103.0,
               "progress": {"fill": 100.1, "status": "Target hit", "time": "10:30", "price": 103.0, "gross": 2.9,
                            "costs": 0.6, "net": 2.3}}
        book = sc.build_workbook(info, [], [], [], pd.DataFrame(columns=sc.SETUP_COLUMNS),
                                 pd.DataFrame(columns=sc.RUN_COLUMNS), {}, [row])
        self.assertEqual(book.sheetnames, ["Today's Mock Trades"])
        ws = book["Today's Mock Trades"]
        self.assertIn("watch scan, data to 10:30 UK", ws.cell(1, 1).value)
        self.assertIn("Mock trades made today", ws.cell(2, 1).value)
        self.assertEqual([ws.cell(4, c).value for c in (3, 5, 6, 11, 16)], ["Alpha plc", "Reclaim", "Buy", "Target hit", 2.3])
        self.assertEqual(ws.cell(3, 10).value, "Stop %")
        self.assertAlmostEqual(ws.cell(4, 10).value, round((100.1 - 99.0) / 100.1 * 100, 2))   # room at the entry price
        self.assertLessEqual(max(ws.column_dimensions[c].width for c in "ABCDEFGHIJKLMNOP"), 22)

    def test_open_trade_closes_at_end_of_day(self):
        from ftse_live_mock import mock_progress
        index = pd.date_range("2026-10-21 16:00", periods=2, freq="15min", tz="Europe/London")
        data = pd.DataFrame([(100, 100.2, 99.8, 100), (100.1, 100.6, 99.9, 100.5)],
                            columns=["open", "high", "low", "close"], index=index)
        row = dict(self.trade(), signal_time="2026-10-21 15:59:00")
        p = mock_progress(row, data)
        self.assertEqual((p["status"], p["time"], p["price"]), ("Closed at end of day", "16:30", 100.5))
        early = mock_progress(row, data.iloc[:1])
        self.assertEqual((early["status"], early["time"]), ("Open", "16:15"))

    def test_refresh_after_close_rewrites_mock_sheet(self):
        import tempfile
        from unittest.mock import patch
        import ftse_trade_scanner as sc
        with tempfile.TemporaryDirectory() as folder:
            folder = Path(folder)
            pd.DataFrame([{**{c: "" for c in sc.RUN_COLUMNS}, "day": "2026-10-21"}])[sc.RUN_COLUMNS].to_csv(folder / "runs.csv", index=False)
            rows = [{"trade_no": 1, "signal_time": "2026-10-21 10:03:48", "company": "Alpha plc", "ticker": "A.L",
                     "setup": "Range", "direction": "Short", "stop": 101.0, "target": 97.0,
                     "progress": {"fill": 100.0, "status": "Closed at end of day", "time": "16:30", "price": 99.0,
                                  "gross": 1.0, "costs": 0.1, "net": 0.9}}]
            now = pd.Timestamp("2026-10-21 17:05", tz=sc.UK_TZ).to_pydatetime()
            with patch.object(sc, "todays_mock_trades", return_value=rows):
                self.assertEqual(sc.refresh_mock_workbook(folder, folder / "s.xlsx", "15m", now), 0)
            from openpyxl import load_workbook
            ws = load_workbook(folder / "s.xlsx")["Today's Mock Trades"]
            self.assertIn("after the market close", ws.cell(1, 1).value)
            self.assertEqual(ws.cell(4, 11).value, "Closed at end of day")
            other = pd.Timestamp("2026-10-22 17:05", tz=sc.UK_TZ).to_pydatetime()
            (folder / "s.xlsx").unlink()
            self.assertEqual(sc.refresh_mock_workbook(folder, folder / "s.xlsx", "15m", other), 0)
            self.assertFalse((folder / "s.xlsx").exists())

    def test_closed_trades_free_their_slots(self):
        self.dir = Path(tempfile.mkdtemp())
        first = live.record_live_mock(self.dir, info("2026-10-19 10:18:00"),
                                      [setup(t, 90 - n) for n, t in enumerate(["A.L", "B.L", "C.L", "D.L", "E.L", "F.L"])],
                                      10, max_open=5, open_now=0)
        self.assertEqual([r["ticker"] for r in first], ["A.L", "B.L", "C.L", "D.L", "E.L"])   # several from one scan
        self.assertEqual(live.record_live_mock(self.dir, info("2026-10-19 10:33:00"), [setup("G.L", 99)],
                                               10, max_open=5, open_now=5), [])               # all 5 still open
        later = live.record_live_mock(self.dir, info("2026-10-19 12:03:00"), [setup("G.L", 99), setup("H.L", 98),
                                      setup("A.L", 97)], 10, max_open=5, open_now=3)          # 2 have closed
        self.assertEqual([(r["ticker"], r["trade_no"]) for r in later], [("G.L", 6), ("H.L", 7)])  # A.L already traded today

    def test_daily_ceiling_still_applies(self):
        self.dir = Path(tempfile.mkdtemp())
        live.record_live_mock(self.dir, info("2026-10-19 10:18:00"), [setup(f"S{n}.L", 90 - n) for n in range(9)],
                              10, max_open=0)
        rows = live.record_live_mock(self.dir, info("2026-10-19 11:03:00"), [setup("X.L", 99), setup("Y.L", 98)],
                                     10, max_open=5, open_now=0)
        self.assertEqual([r["ticker"] for r in rows], ["X.L"])

    def test_skipped_trades_do_not_use_up_the_daily_limit(self):
        self.dir = Path(tempfile.mkdtemp())
        live.record_live_mock(self.dir, info("2026-10-19 10:18:00"), [setup(f"S{n}.L", 90 - n) for n in range(10)], 10, max_open=0)
        none_left = live.record_live_mock(self.dir, info("2026-10-19 11:03:00"), [setup("X.L", 99)], 10, max_open=0)
        self.assertEqual(none_left, [])
        rows = live.record_live_mock(self.dir, info("2026-10-19 11:18:00"), [setup("X.L", 99), setup("Y.L", 98),
                                     setup("Z.L", 97)], 10, max_open=0, not_entered=2)   # 2 of the 10 were skipped
        self.assertEqual([r["ticker"] for r in rows], ["X.L", "Y.L"])
        statuses = [{"progress": {"status": s}} for s in ("Skipped: stop too close at entry", "Open",
                                                           "Skipped: opened past the stop or target", "Stopped out")]
        self.assertEqual(live.skipped_at_entry(statuses), 2)

    def test_still_open_counts_unknown_as_open(self):
        rows = [{"progress": {"status": s}} for s in ("Open", "Stopped out", "Target hit", "Closed at end of day",
                                                       "Skipped: opened past the stop or target", "Waiting for entry candle",
                                                       "Price check unavailable")]
        self.assertEqual(live.still_open(rows), 3)


if __name__ == "__main__":
    unittest.main()
