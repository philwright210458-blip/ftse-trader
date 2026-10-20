"""Claude's model: a monthly momentum portfolio of FTSE 350 shares (daily data).

A different approach from the 15-minute scanner, built on one of the most
studied effects in share prices: shares that have risen most over the past
year tend to keep doing better than average for the next few months (the
"momentum" effect, found in UK shares as well as US ones). Trades last weeks
or months, so dealing costs matter far less than in day trading.

Rules (set once from the published research, not tuned to these results):

1. Once a month, after the close on the last trading day, rank every share
   (investment trusts excluded) by its price change from 12 months ago to
   1 month ago. The most recent month is skipped because it tends to reverse.
2. Only shares trading above their 200-day average price can be bought or kept.
3. Market filter: if the FTSE 100 is below its own 200-day average, sell
   everything and hold cash until it recovers.
4. Hold the 10 best-ranked shares. A share already held is kept while it stays
   in the top 20, which avoids selling and re-buying for small rank changes.
5. Trades happen at the next morning's opening price. Costs: 0.10% dealing on
   every buy and sell, plus 0.5% stamp duty on every buy.

Dividends are added back from Yahoo's dividend records, so returns include them.
Run on a machine that can reach Yahoo (the "FTSE Momentum Backtest" workflow):

    python ftse_momentum.py --exclude-trusts --out ftse_momentum_backtest.xlsx

Big caution: the share list is today's FTSE 350. Companies that collapsed or
were taken over and left the index are missing, which flatters any backtest
that buys shares from it ("survivorship bias"). So the model is also compared
with simply holding every share on the same list equally, which has the same
bias: the gap between the two is the fairer measure of the model itself.
"""

from __future__ import annotations

import argparse
import math
import sys
import time
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import ftse_trade_scanner as sc

MARKET_TICKER = "^FTSE"
BENCHMARK_TICKER = "ISF.L"      # iShares Core FTSE 100 ETF: the FTSE 100 including dividends
DEALING, STAMP = 0.001, 0.005
LOOKBACK, SKIP, TREND_DAYS = 252, 21, 200


@dataclass
class Settings:
    top: int = 10
    keep: int = 20
    market_filter: bool = True
    pot: float = 1000.0
    start: str = "2011-01-01"


@dataclass
class Result:
    values: pd.Series                       # portfolio value on every trading day
    trades: list[dict] = field(default_factory=list)
    holdings: dict[str, float] = field(default_factory=dict)
    costs: float = 0.0
    months_in_cash: int = 0
    rebalances: int = 0


# ----------------------------------------------------------------------
# Data
# ----------------------------------------------------------------------

def download(tickers: list[str], start: str, chunk: int = 40,
             pause: float = 2.0) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """Daily opens and closes (split-adjusted, not dividend-adjusted) and dividends, one column per ticker.

    Dividends are added back by this script (see with_dividends) rather than trusting Yahoo's
    adjusted prices, because Yahoo sometimes records UK dividends in pounds against prices in pence."""
    import yfinance as yf
    parts = {"Open": [], "Close": [], "Dividends": []}
    for first in range(0, len(tickers), chunk):
        batch = tickers[first:first + chunk]
        raw = yf.download(batch, start=start, interval="1d", auto_adjust=False, actions=True, progress=False,
                          group_by="column", threads=True)
        if raw is None or raw.empty:
            continue
        for name in parts:
            if isinstance(raw.columns, pd.MultiIndex):
                part = raw[name] if name in raw.columns.get_level_values(0) else pd.DataFrame(0.0, index=raw.index, columns=batch)
            else:
                part = raw[[name]].set_axis(batch, axis=1) if name in raw.columns else pd.DataFrame(0.0, index=raw.index, columns=batch)
            parts[name].append(part)
        print(f"  ...{min(first + chunk, len(tickers))}/{len(tickers)} downloaded")
        time.sleep(pause)
    frames = []
    for name in ("Open", "Close", "Dividends"):
        frame = pd.concat(parts[name], axis=1) if parts[name] else pd.DataFrame()
        frame.index = pd.DatetimeIndex(frame.index).tz_localize(None).normalize()
        frames.append(frame.sort_index())
    return frames[0], frames[1], frames[2].fillna(0.0)


def corrected_dividends(close: pd.DataFrame, dividends: pd.DataFrame) -> tuple[pd.DataFrame, int]:
    """Dividends per share in the same units as the price, with ex-dates on the price's dates.

    A payment worth less than 0.2% of the share price that becomes 0.2%-15% when multiplied by 100 is
    taken to be recorded in pounds against a price in pence, and is corrected. Returns (dividends,
    number of shares corrected)."""
    result = pd.DataFrame(0.0, index=close.index, columns=close.columns)
    corrected = 0
    for ticker in close.columns:
        if ticker not in dividends.columns:
            continue
        paid = dividends[ticker].reindex(close.index).fillna(0.0)
        paid = paid[paid > 0]
        if paid.empty:
            continue
        before = close[ticker].ffill().shift(1)
        fixed_any = False
        for stamp, amount in paid.items():
            price = before.get(stamp)
            if not (isinstance(price, float) and math.isfinite(price) and price > 0):
                continue
            if amount / price < 0.002 and 0.002 <= amount * 100 / price <= 0.15:
                amount, fixed_any = amount * 100, True
            if 0 < amount / price < 0.5:
                result.at[stamp, ticker] = amount
        corrected += int(fixed_any)
    return result, corrected


def with_dividends(open_: pd.DataFrame, close: pd.DataFrame, dividends: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame, dict]:
    """Total-return prices: earlier prices are scaled down by each dividend, as if it were reinvested."""
    open_, close = open_.copy(), close.copy()
    paid, corrected = corrected_dividends(close, dividends)
    report = {"with_dividends": int((dividends.reindex(columns=close.columns).fillna(0) > 0).any().sum()),
              "corrected": corrected}
    for ticker in close.columns:
        payments = paid[ticker][paid[ticker] > 0]
        if payments.empty:
            continue
        before = close[ticker].ffill().shift(1)
        factor = pd.Series(1.0, index=close.index)
        for stamp, amount in payments.items():
            factor[close.index < stamp] *= 1 - amount / before[stamp]
        close[ticker] = close[ticker] * factor
        if ticker in open_.columns:
            open_[ticker] = open_[ticker] * factor
    return open_, close, report


def clean(open_: pd.DataFrame, close: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Fix Yahoo's two common UK glitches: pence/pounds unit flips and one-day spikes."""
    open_, close = open_.copy(), close.copy()
    for ticker in close.columns:
        series = close[ticker]
        valid = series.dropna()
        if len(valid) < 3:
            continue
        ratio = valid / valid.shift(1)
        factor = pd.Series(1.0, index=valid.index)
        running = 1.0
        for stamp in reversed(valid.index[1:]):                      # latest prices are taken as correct
            factor[stamp] = running
            r = ratio[stamp]
            if 50 <= r <= 200:
                running *= 100.0
            elif 1 / 200 <= r <= 1 / 50:
                running /= 100.0
        factor[valid.index[0]] = running
        fixed = valid * factor
        r1 = fixed / fixed.shift(1)
        r2 = fixed.shift(-1) / fixed
        spike = ((r1 > 2.5) & (r2 < 1 / 2.5)) | ((r1 < 1 / 2.5) & (r2 > 2.5))
        fixed[spike] = np.nan
        close[ticker] = fixed.reindex(close.index)
        if ticker in open_.columns:
            o = open_[ticker].reindex(valid.index) * factor
            o[spike] = np.nan
            bad = (o / fixed > 3) | (o / fixed < 1 / 3)
            o[bad] = fixed[bad]
            open_[ticker] = o.reindex(open_.index)
    return open_, close


def month_ends(dates: pd.DatetimeIndex) -> list[pd.Timestamp]:
    frame = pd.Series(dates, index=dates)
    return list(frame.groupby([dates.year, dates.month]).max())


# ----------------------------------------------------------------------
# The model
# ----------------------------------------------------------------------

def ranking(close_ff: pd.DataFrame, sma: pd.DataFrame, counts: pd.DataFrame, last_seen: pd.DataFrame,
            position: int) -> pd.DataFrame:
    """Every share eligible at this date, best momentum first (uses only prices up to this date)."""
    if position < LOOKBACK:
        return pd.DataFrame(columns=["score", "price", "average_200d"])
    stamp = close_ff.index[position]
    now, then, recent = close_ff.iloc[position], close_ff.iloc[position - LOOKBACK], close_ff.iloc[position - SKIP]
    fresh = (stamp - last_seen.iloc[position]) <= pd.Timedelta(days=7)
    enough = counts.iloc[position] >= LOOKBACK + 10
    score = recent / then - 1
    frame = pd.DataFrame({"score": score, "price": now, "average_200d": sma.iloc[position]})
    ok = fresh & enough & frame["score"].notna() & np.isfinite(frame["score"]) & (frame["price"] > frame["average_200d"])
    return frame[ok].sort_values("score", ascending=False)


def market_ok(market_close: pd.Series, stamp: pd.Timestamp) -> bool:
    history = market_close.loc[:stamp].dropna()
    if len(history) < TREND_DAYS:
        return True
    return bool(history.iloc[-1] > history.tail(TREND_DAYS).mean())


def run_model(open_: pd.DataFrame, close: pd.DataFrame, market_close: pd.Series, settings: Settings) -> Result:
    dates = close.index
    close_ff = close.ffill()
    sma = close_ff.rolling(TREND_DAYS, min_periods=TREND_DAYS).mean()
    counts = close.notna().cumsum()
    seen = pd.DataFrame(np.where(close.notna(), dates.values[:, None], np.datetime64("NaT")),
                        index=dates, columns=close.columns).ffill()
    position_of = {stamp: number for number, stamp in enumerate(dates)}
    # decision at a month-end close -> trade at the next morning's open
    schedule = {position_of[s] + 1: position_of[s] for s in month_ends(dates)
                if s >= pd.Timestamp(settings.start) and position_of[s] + 1 < len(dates)}
    result = Result(values=pd.Series(dtype=float))
    if not schedule:
        return result
    cash, units, cost_basis, bought_on = settings.pot, {}, {}, {}
    closes = close_ff.to_numpy()
    column = {t: n for n, t in enumerate(close.columns)}

    def price_for(ticker: str, day: int) -> float | None:
        value = open_.iat[day, open_.columns.get_loc(ticker)] if ticker in open_.columns else np.nan
        if not (math.isfinite(value) and value > 0):
            value = closes[day - 1, column[ticker]]
        return float(value) if math.isfinite(value) and value > 0 else None

    first = min(schedule)
    values = []
    for day in range(first, len(dates)):
        if day in schedule:
            p = schedule[day]
            stamp = dates[p]
            ranked = ranking(close_ff, sma, counts, seen, p)
            in_market = market_ok(market_close, stamp) if settings.market_filter else True
            if in_market:
                keep_zone = set(ranked.index[:settings.keep])
                target = [t for t in units if t in keep_zone]
                for ticker in ranked.index:
                    if len(target) >= settings.top:
                        break
                    if ticker not in target:
                        target.append(ticker)
            else:
                target = []
                result.months_in_cash += 1
            result.rebalances += 1
            for ticker in [t for t in units if t not in target]:          # sells first
                price = price_for(ticker, day)
                if price is None:
                    continue
                proceeds = units[ticker] * price * (1 - DEALING)
                result.costs += units[ticker] * price * DEALING
                cash += proceeds
                result.trades.append({
                    "Share": ticker, "Bought": bought_on[ticker], "Sold": dates[day].date(),
                    "Cost (£)": round(cost_basis[ticker], 2), "Proceeds (£)": round(proceeds, 2),
                    "Profit / loss (£)": round(proceeds - cost_basis[ticker], 2),
                    "Per £100 after costs (£)": round((proceeds / cost_basis[ticker] - 1) * 100, 2),
                    "Why sold": "Market filter: FTSE 100 below its 200-day average" if not in_market
                    else "Dropped out of the top 20 or below its 200-day average"})
                del units[ticker], cost_basis[ticker], bought_on[ticker]
            new = [t for t in target if t not in units]
            buyable = [(t, price_for(t, day)) for t in new]
            buyable = [(t, price) for t, price in buyable if price is not None]
            if buyable and cash > 1:
                each = cash / len(buyable)
                for ticker, price in buyable:
                    units[ticker] = each / (price * (1 + DEALING + STAMP))
                    cost_basis[ticker] = each
                    bought_on[ticker] = dates[day].date()
                    result.costs += each - each / (1 + DEALING + STAMP)
                    cash -= each
        values.append(cash + sum(n * closes[day, column[t]] for t, n in units.items()))
    result.values = pd.Series(values, index=dates[first:])
    result.holdings = {t: n * float(closes[-1, column[t]]) for t, n in units.items()}
    return result


def equal_weight_everything(close: pd.DataFrame, start: pd.Timestamp, end: pd.Timestamp, pot: float) -> pd.Series:
    """Hold every share on the list equally, rebalanced monthly, before costs (the survivorship-bias check)."""
    ends = [s for s in month_ends(close.index) if start <= s <= end]
    monthly = close.ffill().loc[ends]
    returns = (monthly / monthly.shift(1) - 1).where(close.notna().loc[ends] & close.notna().loc[ends].shift(1, fill_value=False))
    growth = (1 + returns.mean(axis=1).fillna(0)).cumprod() * pot
    return growth


def buy_and_hold(open_: pd.Series, close: pd.Series, start: pd.Timestamp, pot: float) -> pd.Series:
    close = close.ffill().loc[start:]
    opening = open_.loc[start:].dropna()
    if close.empty or opening.empty:
        return pd.Series(dtype=float)
    units = pot * (1 - DEALING) / float(opening.iloc[0])
    return close * units


# ----------------------------------------------------------------------
# Measuring
# ----------------------------------------------------------------------

def measures(name: str, values: pd.Series, result: Result | None = None) -> dict:
    values = values.dropna()
    if values.empty:
        return {"Strategy": name}
    monthly = values.groupby([values.index.year, values.index.month]).last()
    years = (values.index[-1] - values.index[0]).days / 365.25
    growth = (values.iloc[-1] / values.iloc[0]) ** (1 / years) - 1 if years > 0 else np.nan
    peak = values.cummax()
    row = {"Strategy": name, "Start": values.index[0].date(), "End": values.index[-1].date(),
           "£1,000 became (£)": round(float(values.iloc[-1] / values.iloc[0] * 1000), 2),
           "Average yearly growth": growth, "Worst fall from peak": float((values / peak - 1).min()),
           "Months that rose": float((monthly.pct_change().dropna() > 0).mean())}
    yearly = calendar_years(values)
    row["Best year"], row["Worst year"] = float(yearly.max()), float(yearly.min())
    if result is not None:
        closed = pd.DataFrame(result.trades)
        row["Shares bought"] = len(result.trades) + len(result.holdings)
        row["Costs paid (£)"] = round(result.costs, 2)
        row["Months in cash"] = result.months_in_cash
        if not closed.empty:
            row["Closed trades that made money"] = float((closed["Profit / loss (£)"] > 0).mean())
            row["Avg per £100 after costs (£)"] = round(float(closed["Per £100 after costs (£)"].mean()), 2)
    return row


def calendar_years(values: pd.Series) -> pd.Series:
    values = values.dropna()
    ends = values.groupby(values.index.year).last()
    starts = pd.concat([pd.Series([values.iloc[0]], index=[values.index[0].year - 1]), ends]).shift(1).loc[ends.index]
    return ends / starts - 1


# ----------------------------------------------------------------------
# Workbook
# ----------------------------------------------------------------------

HEAD, NOTE, GOOD, BAD = "1F3B57", "EEF2F6", "D9EAD3", "F4CCCC"
MONEY = '"£"#,##0.00;[Red]-"£"#,##0.00'


def _note(ws, row: int, text: str, width: int = 12) -> None:
    ws.cell(row, 1, text).font = Font(name="Arial")
    ws.cell(row, 1).fill = PatternFill("solid", fgColor=NOTE)
    ws.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=width)
    ws.row_dimensions[row].height = 15 * max(1, math.ceil(len(text) / 140))


def _table(ws, top: int, frame: pd.DataFrame, pct=(), money=()) -> int:
    for c, name in enumerate(frame.columns, 1):
        cell = ws.cell(top, c, name)
        cell.font, cell.fill = Font(name="Arial", bold=True, color="FFFFFF"), PatternFill("solid", fgColor=HEAD)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[top].height = 42
    for r, (_, row) in enumerate(frame.iterrows(), top + 1):
        for c, name in enumerate(frame.columns, 1):
            value = row[name]
            if isinstance(value, float) and not math.isfinite(value):
                value = None
            cell = ws.cell(r, c, value)
            cell.font = Font(name="Arial")
            if name in pct and isinstance(value, (int, float)):
                cell.number_format = "0.0%"
            if name in money and isinstance(value, (int, float)):
                cell.number_format = MONEY
    return top + len(frame) + 1


def build_workbook(summary: pd.DataFrame, yearly: pd.DataFrame, curves: pd.DataFrame, trades: pd.DataFrame,
                   holdings: pd.DataFrame, picks: pd.DataFrame, picks_note: str, generated: str, names: dict,
                   dividend_text: str = "") -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.cell(1, 1, "Claude's model: monthly momentum on FTSE 350 shares").font = Font(name="Arial", bold=True, size=13)
    _note(ws, 2, f"Generated {generated}. Each strategy starts with £1,000. Rules: once a month rank shares by their price change "
                 "from 12 months ago to 1 month ago; hold the best 10 that are above their 200-day average; keep a holding "
                 "while it stays in the top 20; sell everything when the FTSE 100 is below its 200-day average. Trades at the "
                 "next morning's open; costs 0.10% per buy and sell plus 0.5% stamp duty on buys. Prices include dividends.")
    _note(ws, 3, "Read this first: the share list is today's FTSE 350, so companies that failed or left the index are missing. "
                 "That flatters every strategy that picks from the list. 'Every share on the list' has the same bias, so compare "
                 "the model against that row, not only against the FTSE 100. Cash earns no interest here (it would in reality). "
                 + dividend_text)
    pct = ["Average yearly growth", "Worst fall from peak", "Months that rose", "Best year", "Worst year",
           "Closed trades that made money"]
    end = _table(ws, 5, summary, pct, ["£1,000 became (£)", "Costs paid (£)", "Avg per £100 after costs (£)"])
    for c, width in enumerate([44, 11, 11, 13, 11, 11, 10, 10, 10, 10, 10, 10, 12, 12], 1):
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.freeze_panes = "B6"

    year_ws = wb.create_sheet("Each year")
    _note(year_ws, 1, "Growth in each calendar year (the first and last years are part-years).", 8)
    _table(year_ws, 3, yearly, pct=list(yearly.columns[1:]))
    for c in range(1, len(yearly.columns) + 1):
        year_ws.column_dimensions[get_column_letter(c)].width = 10 if c == 1 else 22

    curve_ws = wb.create_sheet("Value over time")
    _table(curve_ws, 1, curves, money=list(curves.columns[1:]))
    chart = LineChart()
    chart.title, chart.y_axis.title, chart.height, chart.width = "Value of £1,000 (month ends)", "£", 11, 26
    chart.add_data(Reference(curve_ws, min_col=2, max_col=len(curves.columns), min_row=1, max_row=len(curves) + 1),
                   titles_from_data=True)
    chart.set_categories(Reference(curve_ws, min_col=1, min_row=2, max_row=len(curves) + 1))
    curve_ws.add_chart(chart, "H2")
    for c in range(1, len(curves.columns) + 1):
        curve_ws.column_dimensions[get_column_letter(c)].width = 12 if c == 1 else 18

    trade_ws = wb.create_sheet("Trades")
    shown = trades.assign(Company=trades["Share"].map(names)) if not trades.empty else trades
    if not shown.empty:
        shown = shown[["Company", "Share", "Bought", "Sold", "Cost (£)", "Proceeds (£)", "Profit / loss (£)",
                       "Per £100 after costs (£)", "Why sold"]]
    _table(trade_ws, 1, shown, money=["Cost (£)", "Proceeds (£)", "Profit / loss (£)", "Per £100 after costs (£)"])
    for c, width in enumerate([30, 10, 12, 12, 12, 12, 14, 14, 50], 1):
        trade_ws.column_dimensions[get_column_letter(c)].width = width
    trade_ws.freeze_panes = "A2"

    now_ws = wb.create_sheet("Holdings and picks now")
    _note(now_ws, 1, "What the model holds at the end of the backtest (main version), and how it would rank shares today. "
                     "For paper-trading only; this is not advice.", 8)
    row = _table(now_ws, 3, holdings, money=["Value now (£)"]) + 1
    _note(now_ws, row, picks_note, 8)
    _table(now_ws, row + 2, picks, pct=["12-1 month price change"], money=[])
    for c, width in enumerate([30, 10, 16, 14, 14, 14], 1):
        now_ws.column_dimensions[get_column_letter(c)].width = width
    return wb


# ----------------------------------------------------------------------

def run(args) -> int:
    universe = sc.load_universe(Path(args.universe))
    if args.exclude_trusts:
        universe = [c for c in universe if c.kind != "trust"]
    if args.limit:
        universe = universe[:args.limit]
    names = {c.yahoo_ticker: c.company for c in universe}
    tickers = [c.yahoo_ticker for c in universe]
    data_start = (pd.Timestamp(args.start) - pd.DateOffset(months=15)).strftime("%Y-%m-%d")
    print(f"Downloading daily prices for {len(tickers)} shares, the FTSE 100 and {BENCHMARK_TICKER} from {data_start}...")
    open_, close, dividends = download(tickers + [MARKET_TICKER, BENCHMARK_TICKER], data_start)
    if close.empty or MARKET_TICKER not in close.columns:
        print("No data downloaded (is Yahoo reachable?).")
        return 1
    market = close[MARKET_TICKER]
    shares = [t for t in tickers if t in close.columns and close[t].notna().sum() > 0]
    bench = [BENCHMARK_TICKER] if BENCHMARK_TICKER in close.columns else []
    open_, close = clean(open_[shares + bench], close[shares + bench])
    open_, close, dividend_report = with_dividends(open_, close, dividends)
    bench_open, bench_close = (open_.pop(BENCHMARK_TICKER), close.pop(BENCHMARK_TICKER)) if bench else (None, None)
    bench_paid = int((dividends.get(BENCHMARK_TICKER, pd.Series(dtype=float)) > 0).sum())
    dividend_text = (f"Dividends: Yahoo had dividend records for {dividend_report['with_dividends']} of {len(shares) + len(bench)} "
                     f"shares and funds ({dividend_report['corrected']} corrected from pounds to pence); the FTSE 100 tracker "
                     f"had {bench_paid} payments.")
    print("  " + dividend_text)
    close = close.loc[market.dropna().index.min():]
    open_ = open_.reindex(close.index)
    print(f"  {len(shares)} shares with data; {close.index[0].date()} to {close.index[-1].date()}")

    main = Settings(start=args.start, pot=1000.0)
    unfiltered = Settings(start=args.start, pot=1000.0, market_filter=False)
    results = {"Claude's model (with market filter)": run_model(open_, close, market, main),
               "Claude's model without the market filter": run_model(open_, close, market, unfiltered)}
    first = min(r.values.index[0] for r in results.values())
    last = close.index[-1]
    curves = {name: r.values for name, r in results.items()}
    if bench_close is not None and bench_close.notna().any():
        label = "FTSE 100 (tracker fund, with dividends)" if bench_paid else "FTSE 100 (tracker fund, price only: no dividend data)"
        curves[label] = buy_and_hold(bench_open, bench_close, first, 1000.0)
    curves["Every share on the list, equally (before costs)"] = equal_weight_everything(close, first, last, 1000.0)

    summary = pd.DataFrame([measures(name, values, results.get(name)) for name, values in curves.items()])
    yearly = pd.DataFrame({name: calendar_years(values) for name, values in curves.items()})
    yearly.index.name = "Year"
    yearly = yearly.reset_index()
    month_points = sorted({s for s in month_ends(close.index) if s >= first})
    curve_frame = pd.DataFrame({name: values.reindex(values.index.union(month_points)).ffill().reindex(month_points)
                                for name, values in curves.items()})
    curve_frame.index = [s.strftime("%Y-%m") for s in curve_frame.index]
    curve_frame.index.name = "Month"
    curve_frame = curve_frame.round(2).reset_index()

    model = results["Claude's model (with market filter)"]
    trades = pd.DataFrame(model.trades)
    holdings = pd.DataFrame([{"Company": names.get(t, t), "Share": t, "Value now (£)": round(v, 2)}
                             for t, v in sorted(model.holdings.items(), key=lambda kv: -kv[1])])
    if holdings.empty:
        holdings = pd.DataFrame({"Company": ["(holding cash)"], "Share": [""], "Value now (£)": [None]})
    close_ff = close.ffill()
    sma = close_ff.rolling(TREND_DAYS, min_periods=TREND_DAYS).mean()
    counts = close.notna().cumsum()
    seen = pd.DataFrame(np.where(close.notna(), close.index.values[:, None], np.datetime64("NaT")),
                        index=close.index, columns=close.columns).ffill()
    today = ranking(close_ff, sma, counts, seen, len(close) - 1).head(20)
    picks = pd.DataFrame({"Company": [names.get(t, t) for t in today.index], "Share": today.index,
                          "12-1 month price change": today["score"].round(4).values,
                          "Price": today["price"].round(2).values, "200-day average": today["average_200d"].round(2).values})
    filter_state = "ABOVE (model may hold shares)" if market_ok(market, close.index[-1]) else "BELOW (model would hold cash)"
    picks_note = (f"Ranking as at {close.index[-1].date()} (the model itself only trades on the last trading day of each month). "
                  f"FTSE 100 versus its 200-day average: {filter_state}.")
    generated = datetime.now(sc.UK_TZ).strftime("%Y-%m-%d %H:%M UK time")
    book = build_workbook(summary, yearly, curve_frame, trades, holdings, picks, picks_note, generated, names, dividend_text)
    saved = sc.save_workbook(book, Path(args.out))
    pd.set_option("display.width", 220)
    print(summary.to_string(index=False))
    print(f"\nWorkbook saved: {saved}")
    return 0


def main() -> None:
    parser = argparse.ArgumentParser(description="Backtest Claude's monthly momentum model on FTSE 350 shares")
    parser.add_argument("--universe", default=str(sc.UNIVERSE_PATH))
    parser.add_argument("--start", default="2011-01-01", help="First month-end the model may trade (default 2011-01-01)")
    parser.add_argument("--exclude-trusts", action="store_true")
    parser.add_argument("--limit", type=int, default=None)
    parser.add_argument("--out", default="ftse_momentum_backtest.xlsx")
    sys.exit(run(parser.parse_args()))


if __name__ == "__main__":
    main()
