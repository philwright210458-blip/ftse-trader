"""Daily-chart backtest of the MA-zone pullback (ftse_ma_zone.py) on FTSE 100 shares.

    python ftse_ma_daily_backtest.py --index "FTSE 100" --out results/ftse_ma_zone_daily.xlsx

* About 10 years of Yahoo daily candles (split-adjusted; dividends left out,
  which slightly understates buy results).
* Signal on a day's candle; entry at the next day's open; target reset to 2x the
  stop distance from that actual entry. Each later day: stop counted first if
  both stop and target are inside the day's range; a gap through the stop fills
  at the open. Closed at the close after HOLD_DAYS if neither is hit.
* One open trade per share at a time. Every trade is £100. Costs: 10p dealing
  per trade plus 50p stamp duty (buys of real shares).
* Yardstick: the same entries, stops and targets taken on every uptrend day
  (rule 1 only), to show whether the dip-and-bounce and the slope add anything.

Caution: the share list is today's FTSE 100, so companies that left the index
are missing. That flatters every buy strategy here, the yardstick included.
"""

from __future__ import annotations

import argparse
import math
import sys
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook

import ftse_backtest as bt
import ftse_ma_zone as mz
import ftse_momentum as mm
import ftse_trade_scanner as sc

HOLD_DAYS = 20
DEALING_GBP, STAMP_GBP = 0.10, 0.50


def download(tickers: list[str], period: str) -> dict[str, pd.DataFrame]:
    raw = sc.fetch_many(tickers, "1d", period)
    result = {}
    for ticker, frame in raw.items():
        frame = frame[["open", "high", "low", "close"]].astype(float)
        frame.index = pd.DatetimeIndex(frame.index).tz_localize(None).normalize() if pd.DatetimeIndex(frame.index).tz \
            else pd.DatetimeIndex(frame.index).normalize()
        frame = frame[~frame.index.duplicated(keep="last")].sort_index()
        close = frame[["close"]]
        o, c = mm.clean(frame[["open"]].rename(columns={"open": "close"}), close)
        h, _ = mm.clean(frame[["high"]].rename(columns={"high": "close"}), close)
        l, _ = mm.clean(frame[["low"]].rename(columns={"low": "close"}), close)
        fixed = pd.DataFrame({"open": o["close"], "high": h["close"], "low": l["close"], "close": c["close"]}).dropna()
        fixed = fixed[(fixed["high"] >= fixed[["open", "close"]].max(axis=1))
                      & (fixed["low"] <= fixed[["open", "close"]].min(axis=1)) & (fixed["low"] > 0)]
        if len(fixed) > 300:
            result[ticker] = fixed
    return result


def simulate(frame: pd.DataFrame, signals: pd.DataFrame, hold: int = HOLD_DAYS) -> list[dict]:
    """Daily trades for one share, one open trade at a time."""
    trades, busy_until = [], -1
    position = {stamp: n for n, stamp in enumerate(frame.index)}
    o, h, l, c = (frame[k].to_numpy() for k in ("open", "high", "low", "close"))
    for stamp, row in signals.iterrows():
        n = position[stamp]
        if n <= busy_until or n + 1 >= len(frame):
            continue
        fill, stop = o[n + 1], float(row["stop"])
        if not (math.isfinite(fill) and fill > stop):
            continue                                            # opened at or below the stop: skipped
        risk = fill - stop
        target = fill + mz.REWARD_RISK * risk
        outcome, exit_price, exit_n = "timeout", None, None
        last = min(n + hold, len(frame) - 1)
        for k in range(n + 1, last + 1):
            if l[k] <= stop:                                    # stop first; a gap down fills at the open
                outcome, exit_price, exit_n = "stop", min(o[k], stop), k
                break
            if h[k] >= target:
                outcome, exit_price, exit_n = "target", max(o[k], target), k
                break
        if exit_price is None:
            if n + hold >= len(frame):
                break                                           # still open at the end of the data
            exit_price, exit_n = c[last], last
        gross = (exit_price - fill) / fill * 100
        trades.append({"signal": stamp.date(), "entry_day": frame.index[n + 1].date(), "exit_day": frame.index[exit_n].date(),
                       "days": exit_n - n, "fill": round(fill, 4), "stop": round(stop, 4), "target": round(target, 4),
                       "exit": round(exit_price, 4), "outcome": outcome, "slope": round(float(row["slope"]), 4),
                       "gross_gbp": round(gross, 2), "costs_gbp": DEALING_GBP + STAMP_GBP,
                       "net_gbp": round(gross - DEALING_GBP - STAMP_GBP, 2)})
        busy_until = exit_n
    return trades


def summary(name: str, trades: pd.DataFrame, split: pd.Timestamp) -> dict:
    if trades.empty:
        return {"Strategy": name, "Completed trades": 0}
    entry = pd.to_datetime(trades["entry_day"])
    yearly = trades.groupby(entry.dt.year)["net_gbp"].mean()
    return {"Strategy": name, "Completed trades": len(trades), "Shares traded": trades["ticker"].nunique(),
            "Target hit": (trades["outcome"] == "target").mean(), "Stop hit": (trades["outcome"] == "stop").mean(),
            "Closed after 20 days": (trades["outcome"] == "timeout").mean(),
            "Made money after costs": (trades["net_gbp"] > 0).mean(),
            "Avg before costs per £100 (£)": trades["gross_gbp"].mean(), "Avg after costs per £100 (£)": trades["net_gbp"].mean(),
            "Total after costs (£)": trades["net_gbp"].sum(), "Avg days held": trades["days"].mean(),
            "First half: avg after costs (£)": trades.loc[entry < split, "net_gbp"].mean(),
            "Second half: avg after costs (£)": trades.loc[entry >= split, "net_gbp"].mean(),
            "Years with a profit": f"{int((yearly > 0).sum())} of {len(yearly)}"}


def run(args) -> int:
    universe = [c for c in sc.load_universe(Path(args.universe)) if c.kind != "trust"]
    if args.index:
        universe = [c for c in universe if c.index == args.index]
    if args.limit:
        universe = universe[:args.limit]
    names = {c.yahoo_ticker: c.company for c in universe}
    print(f"Downloading {args.period} of daily candles for {len(universe)} shares...")
    data = download([c.yahoo_ticker for c in universe], args.period)
    if not data:
        print("No data downloaded (is Yahoo reachable?).")
        return 1
    print(f"  {len(data)} shares with enough history")
    zone, every = [], []
    for ticker, frame in data.items():
        ind = mz.indicators(frame)
        for kind, mask, store in (("zone", mz.signal_mask(ind), zone), ("trend", mz.trend_mask(ind), every)):
            signals = mz.levels(ind, mask.fillna(False))
            for trade in simulate(frame, signals):
                store.append({**trade, "ticker": ticker, "company": names.get(ticker, ticker)})
    zone_frame, every_frame = pd.DataFrame(zone), pd.DataFrame(every)
    if zone_frame.empty:
        print("No signals found.")
        return 1
    zone_frame["slope group"] = mz.slope_groups(zone_frame["slope"])
    all_days = sorted(pd.to_datetime(zone_frame["entry_day"]))
    split = all_days[len(all_days) // 2]
    rows = [summary("MA zone: every signal", zone_frame, split)]
    for group in ("Flat", "Medium (about 45 degrees)", "Steep"):
        rows.append(summary(f"MA zone: {group} slope", zone_frame[zone_frame["slope group"] == group], split))
    rows.append(summary("Yardstick: buy on any uptrend day (rule 1 only)", every_frame, split))
    table = pd.DataFrame(rows)
    cuts = zone_frame.groupby("slope group")["slope"].agg(["min", "max"])
    period = f"{min(f.index[0] for f in data.values()).date()} to {max(f.index[-1] for f in data.values()).date()}"

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    from openpyxl.styles import Font
    ws.cell(1, 1, "MA-zone pullback (idea 22) on daily candles").font = Font(name="Arial", bold=True, size=13)
    bt.note(ws, 2, f"Generated {datetime.now(sc.UK_TZ):%Y-%m-%d %H:%M} UK time. {len(data)} {args.index or 'FTSE'} shares, daily "
                   f"candles {period}. Rules: 9 EMA above a rising 20 SMA; the day's low dips to the 9 EMA or into the zone "
                   "down to the 20 SMA; the day closes up, in the top half of its range. Buy at the next open; stop just below "
                   "the dip or the 20 SMA; target 2x the stop distance; closed after 20 trading days if neither is hit. "
                   "£100 a trade; costs 10p dealing + 50p stamp duty.")
    bt.note(ws, 3, "Slope groups: the signals split into three equal groups by how fast the 20 SMA was rising (per day, in "
                   "typical daily ranges): " + "; ".join(f"{g}: {r['min']:.3f} to {r['max']:.3f}" for g, r in cuts.iterrows())
                   + ". The yardstick takes the same kind of trade on every uptrend day, with no dip or bounce needed.")
    bt.note(ws, 4, "Caution: today's FTSE 100 list, so companies that dropped out are missing; this flatters every buy "
                   "strategy here, the yardstick included. Dividends are not added. First / second half = trades split "
                   "at the middle entry date.")
    money = [c for c in table.columns if "(£)" in c]
    pct = ["Target hit", "Stop hit", "Closed after 20 days", "Made money after costs"]
    bt.write_table(ws, 6, table.round(4), money, pct)
    for col, width in zip("ABCDEFGHIJKLMNO", [44, 10, 9, 9, 9, 10, 11, 12, 12, 12, 9, 13, 13, 12]):
        ws.column_dimensions[col].width = width

    year_ws = wb.create_sheet("Each year")
    entry = pd.to_datetime(zone_frame["entry_day"]).dt.year
    yearly = pd.DataFrame({
        "Trades": zone_frame.groupby(entry).size(),
        "MA zone: avg after costs per £100 (£)": zone_frame.groupby(entry)["net_gbp"].mean().round(2),
        "MA zone: total (£)": zone_frame.groupby(entry)["net_gbp"].sum().round(2),
        "Yardstick: avg after costs per £100 (£)": every_frame.groupby(pd.to_datetime(every_frame["entry_day"]).dt.year)["net_gbp"].mean().round(2)})
    yearly.index.name = "Year"
    bt.write_table(year_ws, 1, yearly.reset_index(), [c for c in yearly.columns if "(£)" in c])
    for col in "ABCDE":
        year_ws.column_dimensions[col].width = 22

    trade_ws = wb.create_sheet("Trades (MA zone)")
    shown = zone_frame.sort_values("entry_day")[["entry_day", "company", "ticker", "slope group", "fill", "stop", "target",
                                                 "outcome", "exit_day", "days", "gross_gbp", "costs_gbp", "net_gbp"]]
    shown.columns = ["Entry day", "Company", "Ticker", "Slope group", "Entry", "Stop", "Target", "Result", "Exit day",
                     "Days held", "Before costs (£)", "Costs (£)", "After costs (£)"]
    bt.write_table(trade_ws, 1, shown, ["Before costs (£)", "Costs (£)", "After costs (£)"])
    trade_ws.freeze_panes = "A2"
    saved = sc.save_workbook(wb, Path(args.out))
    pd.set_option("display.width", 250)
    print(table[["Strategy", "Completed trades", "Target hit", "Avg before costs per £100 (£)",
                 "Avg after costs per £100 (£)", "First half: avg after costs (£)", "Second half: avg after costs (£)"]].round(3)
          .to_string(index=False))
    print(f"Workbook saved: {saved}")
    return 0


def main() -> None:
    parser = argparse.ArgumentParser(description="Daily-chart backtest of the MA-zone pullback")
    parser.add_argument("--universe", default=str(sc.UNIVERSE_PATH))
    parser.add_argument("--index", default="FTSE 100", help="'FTSE 100', 'FTSE 250', or '' for both")
    parser.add_argument("--period", default="10y")
    parser.add_argument("--limit", type=int, default=None)
    parser.add_argument("--out", default="results/ftse_ma_zone_daily.xlsx")
    sys.exit(run(parser.parse_args()))


if __name__ == "__main__":
    main()
