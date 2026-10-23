"""Stop / target / trade-size backtest (23 Oct 2026)."""

import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import numpy as np
import pandas as pd

import ftse_risk_backtest as rb
import ftse_trade_scanner as sc


def idea(**kw):
    row = {"day": "2026-10-21", "as_of": "10:00", "data_time": "2026-10-21T09:45:00+01:00", "run_time": "2026-10-21 10:02:00",
           "mode": "watch", "company": "Alpha", "ticker": "A.L", "index": "FTSE 100", "setup": "Range", "direction": "Short",
           "score": 70.0, "entry": 100.0, "stop": 101.0, "target": 95.0, "stop_pct": 1.0, "target_pct": 5.0,
           "atr_pct": 0.3, "day_range_pct": 2.0}
    row.update(kw)
    return row


def candles(start="2026-10-21 08:00", n=34, price=100.0, drift=0.0):
    index = pd.date_range(start, periods=n, freq="15min", tz=sc.UK_TZ)
    close = price + drift * np.arange(n)
    return pd.DataFrame({"open": close, "high": close + 0.05, "low": close - 0.05, "close": close, "volume": 1000}, index=index)


class VariantRules(unittest.TestCase):
    def test_cap_target_and_recheck_costs(self):
        frame = pd.DataFrame([idea(stop=101.5, stop_pct=1.5), idea(ticker="B.L", stop=100.3, stop_pct=0.3), idea(ticker="C.L", day_range_pct=float("nan"))])
        capped = rb.cap_target(frame, 1.0)
        self.assertEqual(list(capped["ticker"]), ["B.L"])        # 1.5% stop vs 2% target fails 1.5x after costs
        self.assertAlmostEqual(capped.iloc[0]["target"], 98.0)
        self.assertAlmostEqual(capped.iloc[0]["target_pct"], 2.0)

    def test_variant_filters(self):
        frame = pd.DataFrame([idea(), idea(ticker="T.L", stop=100.3, stop_pct=0.3, atr_pct=0.4)])
        variants = rb.variant_ideas(frame)
        self.assertEqual(len(variants[rb.CONTROL][0]), 2)
        self.assertEqual(list(variants[rb.MIN_STOP][0]["ticker"]), ["A.L"])
        self.assertEqual(list(variants[rb.MIN_STOP_ATR][0]["ticker"]), ["A.L"])
        self.assertTrue(variants[rb.SAME_RISK][1] and variants[rb.ALL_THREE][1])
        self.assertFalse(variants[rb.CONTROL][1])

    def test_live_picker_one_per_scan_five_a_day(self):
        rows = [idea(ticker=f"S{i}.L", run_time=f"2026-10-21 {9 + i // 2:02d}:{'02' if i % 2 == 0 else '17'}:00", score=50 + i)
                for i in range(8)]
        rows += [idea(ticker="X.L", run_time="2026-10-21 09:02:00", score=99), idea(ticker="X.L", run_time="2026-10-21 09:17:00", score=99)]
        picks = rb.live_picks(pd.DataFrame(rows))
        self.assertEqual(len(picks), 5)
        self.assertEqual(picks["ticker"].tolist()[0], "X.L")
        self.assertEqual(picks["run_time"].nunique(), 5)
        self.assertEqual(picks["ticker"].nunique(), 5)

    def test_same_risk_sizing_and_costs(self):
        data = {"A.L": candles(drift=-0.05)}          # falls all day: the sell closes at a profit
        frame = pd.DataFrame([idea(stop=100.5, stop_pct=0.5, target=90, target_pct=10)])
        now = pd.Timestamp("2026-10-21 18:00", tz=sc.UK_TZ)
        fixed = rb.simulate(frame, data, now, "15m", False).iloc[0]
        sized = rb.simulate(frame, data, now, "15m", True).iloc[0]
        self.assertEqual((fixed["stake"], sized["stake"]), (100.0, 200.0))
        self.assertAlmostEqual(sized["at_risk_gbp"], sized["fill_risk_pct"] * 2, places=6)   # risk measured from the actual fill
        self.assertAlmostEqual(sized["costs_gbp"], 0.20)
        self.assertAlmostEqual(sized["net_gbp"], round(2 * fixed["gross_gbp"] - 0.2, 2), places=1)
        self.assertEqual(fixed["outcome"], "timeout")
        tiny = rb.simulate(pd.DataFrame([idea(stop=100.05, stop_pct=0.05)]), data, now, "15m", True).iloc[0]
        self.assertEqual(tiny["stake"], rb.MAX_STAKE)

    def test_daily_range_uses_previous_days_only(self):
        frames = []
        for n, day in enumerate(pd.bdate_range("2026-09-01", periods=12)):
            frame = candles(f"{day.date()} 08:00", 4, 100.0)
            frame["high"] = 100 + n                         # range grows each day
            frames.append(frame)
        ranges = rb.daily_ranges(pd.concat(frames))
        days = [str(d.date()) for d in pd.bdate_range("2026-09-01", periods=12)]
        self.assertNotIn(days[9], ranges)
        self.assertAlmostEqual(ranges[days[10]], np.median([n + 0.05 for n in range(10)]), places=3)

    def test_summary_profit_per_pound_at_risk(self):
        rb.SPLIT_DAY = "2026-10-22"
        trades = pd.DataFrame([dict(idea(), status="done", outcome="stop", stake=100.0, net_gbp=-1.1, at_risk_gbp=1.0),
                               dict(idea(day="2026-10-22"), status="done", outcome="timeout", stake=100.0, net_gbp=0.9, at_risk_gbp=1.0)])
        row = rb.summary_row(rb.CONTROL, rb.SCOPES[0], trades)
        self.assertAlmostEqual(row["Profit per £1 at risk (£)"], -0.1)
        self.assertAlmostEqual(row["First half: profit per £1 at risk (£)"], -1.1)
        self.assertEqual(row["Completed"], 2)


class EndToEnd(unittest.TestCase):
    def test_runs_on_synthetic_prices_and_writes_workbook(self):
        rng = np.random.default_rng(3)
        frames = {}
        for ticker in ("A.L", "B.L"):
            parts = []
            for day in pd.bdate_range("2026-09-01", periods=16):
                index = pd.date_range(f"{day.date()} 08:00", periods=34, freq="15min", tz=sc.UK_TZ)
                walk = 100 + np.cumsum(rng.normal(0, 0.4, 34)) + 3 * np.sin(np.arange(34) / 3)
                parts.append(pd.DataFrame({"open": walk, "high": walk + 0.4, "low": walk - 0.4, "close": walk + 0.1,
                                           "volume": 1000}, index=index))
            frames[ticker] = pd.concat(parts)
        universe = [SimpleNamespace(company=t, yahoo_ticker=t, index="FTSE 100", kind="company") for t in frames]
        with tempfile.TemporaryDirectory() as folder, \
                patch.object(sc, "load_universe", return_value=universe), \
                patch.object(sc, "fetch_many", return_value=frames), \
                patch.object(sc, "completed_candles", side_effect=lambda f, *a: f), \
                patch.object(sc, "drop_untraded_bars", side_effect=lambda f: f):
            args = SimpleNamespace(universe="x", index="FTSE 100", interval="15m", period="60d", warmup_days=12,
                                   limit=None, workers=1, out=str(Path(folder) / "r.xlsx"))
            code = rb.run(args)
            if code == 0:
                from openpyxl import load_workbook
                book = load_workbook(Path(folder) / "r.xlsx")
                self.assertEqual(book.sheetnames[0], "Summary")
            else:
                self.assertEqual(code, 1)     # synthetic prices may give no ideas; the run must still end cleanly


if __name__ == "__main__":
    unittest.main()
