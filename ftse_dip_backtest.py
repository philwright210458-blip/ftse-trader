"""Phil's dip method (plan step 3a, 31 Oct 2026): buy a slide in a big share, sell within a week at +3%.

Phil's own method from the early 2000s: a handful of big FTSE 100 shares he knew well, bought only after a fall
he had seen recover before, sold within a week for about 3%. Rules set in advance (before seeing any results):

  A. Phil's rules: BP, HSBC, Rio Tinto, Shell, SSE (standing in for Scottish Power), Vodafone. Signal when a share
     closes higher than the day before, after closing 3%+ below its highest close of the last 20 days (the slide
     has turned). Buy at the next day's open. Sell at +3% (intraday, or the open if it gaps past), otherwise at the
     close of the 5th trading day (entry day counts as day 1). No stop. One trade per share, up to 3 open at once;
     deepest slide first when several signal on the same day.
  B. Don't wait for the turn: buy at the next open after the first close 3%+ below the high.
  C. Skip crashes: no trade if the slide went 10%+ below the high ("falling knife").
  D. Skip downtrends: no trade if the share closed below its 200-day average on the signal day.
  E. Skip dividend dips: no trade if an ex-dividend date fell inside the slide.
  F. Phil's rules on every FTSE 100 share (not just the favourites), still up to 3 open at once.

* Yahoo daily prices from 2005 (split-adjusted, not dividend-adjusted, so they are the prices actually traded).
  Dividends are added when a trade was held over the ex-dividend date (bought before it, sold on or after it).
* Costs per £100 trade: spread bet 10p + overnight funding (6.5% a year, about 1.8p a night); real shares 10p +
  50p stamp duty. Phil traded about £3,000 a time: multiply the per-£100 figures by 30.
* Random buying (the chance test): the same shares bought on random days with the same exits. A variant "beats
  chance" only if random groups of the same size do as well less than 5% of the time, in both halves.

Caution: today's company list. Phil's favourites are, by hindsight, shares that survived (Scottish Power, which
left the market in 2007, is missing). That flatters every buy-the-dip test, including random buying. Paper research only.

    python ftse_dip_backtest.py --out results/ftse_dip_method.xlsx
"""

from __future__ import annotations

import argparse
import json
import math
import sys
import time
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook

import ftse_backtest as bt
import ftse_momentum as mm
import ftse_trade_scanner as sc

FAVOURITES = ["BP.L", "HSBA.L", "RIO.L", "SHEL.L", "SSE.L", "VOD.L"]
HIGH_DAYS, DIP_PCT, TARGET_PCT, HOLD_DAYS = 20, 3.0, 3.0, 5
CRASH_PCT, TREND_DAYS = 10.0, 200
MAX_OPEN = 3
SPREAD_P, STAMP_P, FUNDING_YEAR = 0.10, 0.50, 0.065
RANDOM_DRAWS, LUCK_LIMIT = 2000, 0.05

VARIANTS = {
    "A. Phil's rules (favourite 6 shares)": {"wait_for_turn": True},
    "B. Buy without waiting for the turn": {"wait_for_turn": False},
    "C. Skip crashes (10%+ below the high)": {"wait_for_turn": True, "skip_crash": True},
    "D. Skip shares below their 200-day average": {"wait_for_turn": True, "skip_downtrend": True},
    "E. Skip dips with an ex-dividend date inside": {"wait_for_turn": True, "skip_dividend_dip": True},
    "F. Phil's rules on all FTSE 100 shares": {"wait_for_turn": True, "all_shares": True},
}


# ---------------------------------------------------------------- data

def tidy_dividends(div: pd.Series, close: pd.Series) -> pd.Series:
    """Yahoo sometimes gives UK dividends in pounds instead of pence (or the reverse). Fix by size against the price."""
    div = div.reindex(close.index).fillna(0.0).astype(float)
    ratio = div / close
    div[ratio > 0.25] /= 100.0
    div[(ratio > 0) & (ratio < 0.0008)] *= 100.0
    return div


def download(tickers: list[str], start: str) -> dict[str, pd.DataFrame]:
    import yfinance as yf
    result = {}
    for n, ticker in enumerate(tickers, 1):
        frame = None
        for attempt in range(3):
            try:
                frame = yf.Ticker(ticker).history(period="max", interval="1d", auto_adjust=False, actions=True)
                break
            except Exception as error:                               # network hiccup: try again
                print(f"  {ticker}: {error} (try {attempt + 1})")
                time.sleep(3)
        if frame is None or frame.empty:
            print(f"  {ticker}: no data")
            continue
        frame.index = pd.DatetimeIndex(frame.index).tz_localize(None).normalize()
        frame = frame[~frame.index.duplicated(keep="last")].sort_index()
        frame = frame.loc[pd.Timestamp(start):]
        prices = frame[["Open", "High", "Low", "Close"]].astype(float)
        close = prices[["Close"]].rename(columns={"Close": "x"})
        fixed = {}
        for col in ("Open", "High", "Low"):
            o, c = mm.clean(prices[[col]].rename(columns={col: "x"}), close)
            fixed[col.lower()] = o["x"]
        fixed["close"] = c["x"]
        clean = pd.DataFrame(fixed).dropna()
        clean = clean[(clean["high"] >= clean[["open", "close"]].max(axis=1))
                      & (clean["low"] <= clean[["open", "close"]].min(axis=1)) & (clean["low"] > 0)]
        div = frame["Dividends"] if "Dividends" in frame.columns else pd.Series(0.0, index=frame.index)
        clean["dividend"] = tidy_dividends(div, clean["close"])
        if len(clean) > 300:
            result[ticker] = clean
        if n % 20 == 0:
            print(f"  downloaded {n} of {len(tickers)}")
    return result


# ---------------------------------------------------------------- one share

def exit_trade(frame: pd.DataFrame, t: int) -> dict | None:
    """Buy at the open of day t+1; sell at +3% or the close of day t+HOLD_DAYS. None if the data ends first."""
    o, h, c = frame["open"].to_numpy(), frame["high"].to_numpy(), frame["close"].to_numpy()
    div = frame["dividend"].to_numpy()
    first, last = t + 1, t + HOLD_DAYS
    if last >= len(frame):
        return None
    fill = o[first]
    if not (math.isfinite(fill) and fill > 0):
        return None
    target = fill * (1 + TARGET_PCT / 100)
    price, k, outcome = c[last], last, "time"
    for day in range(first, last + 1):
        if h[day] >= target:
            price, k, outcome = max(o[day], target), day, "target"
            break
    dividends = float(div[first + 1:k + 1].sum())                   # owned before the ex-date's open
    gross = (price - fill + dividends) / fill * 100
    nights = int((frame.index[k] - frame.index[first]).days)
    return {"entry_day": frame.index[first], "exit_day": frame.index[k], "days": k - first + 1, "fill": float(fill),
            "exit": float(price), "outcome": outcome, "dividend": round(dividends, 4), "gross": gross,
            "net_spread_bet": gross - SPREAD_P - FUNDING_YEAR / 365 * 100 * nights,
            "net_shares": gross - SPREAD_P - STAMP_P}


def signals(frame: pd.DataFrame, wait_for_turn=True, skip_crash=False, skip_downtrend=False,
            skip_dividend_dip=False) -> list[dict]:
    c = frame["close"]
    high = c.rolling(HIGH_DAYS, min_periods=HIGH_DAYS).max()
    below = c / high - 1                                           # 0 at the high, -0.03 = 3% below
    dip = below <= -DIP_PCT / 100
    trend = c.rolling(TREND_DAYS, min_periods=TREND_DAYS).mean()
    ex_div = frame["dividend"].to_numpy() > 0
    cv, bv, dv, tv = c.to_numpy(), below.to_numpy(), dip.to_numpy(), trend.to_numpy()
    out, used = [], False                                          # one signal per slide
    for t in range(max(HIGH_DAYS, 2), len(frame) - 1):
        if not dv[t - 1]:
            used = False                                           # back above the 3% line: a new slide can start
        if wait_for_turn:
            if used or not (dv[t - 1] and cv[t] > cv[t - 1]):
                continue
            used = True
            depth = bv[t - 1]
        else:
            if not (dv[t] and not dv[t - 1]):
                continue
            depth = bv[t]
        # the slide: from the day of the 20-day high to the signal day
        window = cv[t - HIGH_DAYS + 1:t + 1]
        peak = t - HIGH_DAYS + 1 + int(np.argmax(window))
        deepest = float(np.nanmin(bv[peak:t + 1]))
        div_in_dip = bool(ex_div[peak + 1:t + 1].any())
        if skip_crash and deepest <= -CRASH_PCT / 100:
            continue
        if skip_downtrend and not (math.isfinite(tv[t]) and cv[t] >= tv[t]):
            continue
        if skip_dividend_dip and div_in_dip:
            continue
        out.append({"t": t, "depth": float(depth), "deepest": deepest, "div_in_dip": div_in_dip})
    return out


def candidates(data: dict[str, pd.DataFrame], rules: dict) -> pd.DataFrame:
    rows = []
    for ticker, frame in data.items():
        for s in signals(frame, **{k: v for k, v in rules.items() if k != "all_shares"}):
            trade = exit_trade(frame, s["t"])
            if trade:
                rows.append({"ticker": ticker, "signal_day": frame.index[s["t"]], **s, **trade})
    return pd.DataFrame(rows)


def choose(cands: pd.DataFrame, max_open: int = MAX_OPEN) -> pd.DataFrame:
    """Take signals in date order: one trade per share at a time, at most max_open open; deepest slide first."""
    if cands.empty:
        return cands
    taken, busy_until, open_exits = [], {}, []
    for entry, day in cands.sort_values(["entry_day", "depth", "ticker"]).groupby("entry_day", sort=True):
        open_exits = [e for e in open_exits if e >= entry]            # still held on the entry morning
        for _, row in day.iterrows():
            if len(open_exits) >= max_open:
                break
            if busy_until.get(row["ticker"], pd.Timestamp.min) >= entry:
                continue
            taken.append(row)
            busy_until[row["ticker"]] = row["exit_day"]
            open_exits.append(row["exit_day"])
    return pd.DataFrame(taken).reset_index(drop=True)


def random_pool(data: dict[str, pd.DataFrame]) -> pd.DataFrame:
    """A trade bought at every day's open, with the same exits: what buying at random would earn."""
    rows = []
    for ticker, frame in data.items():
        for t in range(HIGH_DAYS, len(frame) - 1):
            trade = exit_trade(frame, t)
            if trade:
                rows.append({"ticker": ticker, "entry_day": trade["entry_day"], "gross": trade["gross"],
                             "net_spread_bet": trade["net_spread_bet"]})
    return pd.DataFrame(rows)


def luck(trades: pd.DataFrame, pool: pd.DataFrame, seed: int = 7) -> float:
    """Share of random groups (same size, same shares, same years) that did as well before costs."""
    if trades.empty or pool.empty:
        return float("nan")
    rng = np.random.default_rng(seed)
    years = set(pd.to_datetime(trades["entry_day"]).dt.year)
    pool = pool[pool["ticker"].isin(set(trades["ticker"])) & pd.to_datetime(pool["entry_day"]).dt.year.isin(years)]
    values = pool["gross"].to_numpy()
    if len(values) < len(trades):
        return float("nan")
    target = trades["gross"].mean()
    draws = np.array([rng.choice(values, size=len(trades), replace=False).mean() for _ in range(RANDOM_DRAWS)])
    return float((draws >= target).mean())


# ---------------------------------------------------------------- summary

def summary(name: str, trades: pd.DataFrame, pool: pd.DataFrame, split: pd.Timestamp) -> dict:
    if trades.empty:
        return {"Rules": name, "Trades": 0}
    entry = pd.to_datetime(trades["entry_day"])
    first, second = trades[entry < split], trades[entry >= split]
    months = max(1.0, (entry.max() - entry.min()).days / 30.44)
    yearly = trades.groupby(entry.dt.year)["net_spread_bet"].sum()
    used = pool[pool["ticker"].isin(set(trades["ticker"]))]
    return {
        "Rules": name, "Trades": len(trades), "Trades a year": len(trades) / max(1, entry.dt.year.nunique()),
        "Hit +3%": float((trades["outcome"] == "target").mean()), "Made money (spread bet)": float((trades["net_spread_bet"] > 0).mean()),
        "Avg days held": float(trades["days"].mean()),
        "Before costs per £100 (£)": float(trades["gross"].mean()),
        "Random buying per £100 (£)": float(used["gross"].mean()) if len(used) else float("nan"),
        "After costs: spread bet (£)": float(trades["net_spread_bet"].mean()),
        "After costs: real shares (£)": float(trades["net_shares"].mean()),
        "Spread bet, first half (£)": float(first["net_spread_bet"].mean()) if len(first) else float("nan"),
        "Spread bet, second half (£)": float(second["net_spread_bet"].mean()) if len(second) else float("nan"),
        "Random as good: whole period": luck(trades, pool),
        "Random as good: first half": luck(first, pool),
        "Random as good: second half": luck(second, pool),
        "Worst trade (£)": float(trades["net_spread_bet"].min()),
        "At £3,000 a trade: £ a month (spread bet)": float(trades["net_spread_bet"].sum() * 30 / months),
        "Years in profit (spread bet)": f"{int((yearly > 0).sum())} of {len(yearly)}",
        "Dividend dips": float(trades["div_in_dip"].mean()),
    }


def verdict_line(row: dict, base: dict | None) -> str:
    name = row["Rules"]
    if not row.get("Trades"):
        return f"{name}: no trades."
    both = row["Spread bet, first half (£)"] > 0 and row["Spread bet, second half (£)"] > 0
    chance = all(row[f"Random as good: {h}"] < LUCK_LIMIT for h in ("first half", "second half"))
    money = (f"made money after spread-bet costs in both halves" if both else
             "did not make money after costs in both halves")
    edge = ("beat random buying in both halves" if chance else
            f"did not clearly beat random buying (random groups as good {row['Random as good: whole period']:.0%} of the time)")
    text = (f"{name}: {money}; {edge}. {row['Trades']} trades, {row['After costs: spread bet (£)']:+.2f} per £100 "
            f"(spread bet), {row['After costs: real shares (£)']:+.2f} on real shares.")
    if base and name != base["Rules"]:
        better = (row["Spread bet, first half (£)"] > base["Spread bet, first half (£)"]
                  and row["Spread bet, second half (£)"] > base["Spread bet, second half (£)"])
        text += " Better than A in both halves." if better else " Not better than A in both halves."
    return text


def run(args) -> int:
    universe = [c for c in sc.load_universe(Path(args.universe)) if c.kind != "trust" and c.index == "FTSE 100"]
    names = {c.yahoo_ticker: c.company for c in universe}
    tickers = sorted(set(FAVOURITES) | ({c.yahoo_ticker for c in universe} if not args.favourites_only else set()))
    if args.limit:
        tickers = FAVOURITES + [t for t in tickers if t not in FAVOURITES][:args.limit]
    print(f"Downloading daily prices and dividends for {len(tickers)} shares from {args.start}...")
    data = download(tickers, args.start)
    favourites = {t: f for t, f in data.items() if t in FAVOURITES}
    if not favourites:
        print("No data for the favourite shares (is Yahoo reachable?).")
        return 1
    pool = random_pool(data)
    results, all_trades = {}, []
    for name, rules in VARIANTS.items():
        shares = data if rules.get("all_shares") else favourites
        trades = choose(candidates(shares, rules))
        if not trades.empty:
            trades["variant"] = name
            trades["company"] = trades["ticker"].map(names).fillna(trades["ticker"])
            all_trades.append(trades)
        results[name] = trades
        print(f"  scored: {name} ({len(trades)} trades)")
    base_trades = results[next(iter(VARIANTS))]
    if base_trades.empty:
        print("No trades under Phil's rules.")
        return 1
    entries = sorted(pd.to_datetime(base_trades["entry_day"]))
    split = entries[len(entries) // 2]
    rows = [summary(name, trades, pool, split) for name, trades in results.items()]
    table = pd.DataFrame(rows)
    span = f"{min(f.index[0] for f in favourites.values()).date()} to {max(f.index[-1] for f in favourites.values()).date()}"
    verdict = [f"Daily prices {span}; favourites found: {', '.join(names.get(t, t) for t in favourites)}; "
               f"{len(data)} FTSE 100 shares for F. Split at {split.date()}."]
    verdict += [verdict_line(r, rows[0]) for r in rows]

    wb = Workbook()
    ws = wb.active
    ws.title = "Verdict"
    ws.cell(1, 1, "Phil's dip method: buy a 3% slide, sell at +3% within a week").font = bt.Font(name="Arial", bold=True, size=13)
    for n, line in enumerate(verdict, 3):
        bt.note(ws, n, line)
    bt.note(ws, len(verdict) + 4, "Per £100 trade. Spread bet: 10p + about 1.8p a night; real shares: 10p + 50p stamp duty. "
                                  "Random buying = same shares on any day with the same exits. Caution: today's company list "
                                  "(survivors only), which flatters buy-the-dip results, random buying included.")
    ws.column_dimensions["A"].width = 150
    var_ws = wb.create_sheet("Variants")
    money = [c for c in table.columns if "(£)" in c or "£ a month" in c]
    pct = ["Hit +3%", "Made money (spread bet)", "Random as good: whole period", "Random as good: first half",
           "Random as good: second half", "Dividend dips"]
    bt.write_table(var_ws, 1, table.round(4), money, pct)
    var_ws.column_dimensions["A"].width = 44
    if all_trades:
        trades = pd.concat(all_trades, ignore_index=True)
        yearly = (trades.assign(year=pd.to_datetime(trades["entry_day"]).dt.year)
                  .pivot_table(index="year", columns="variant", values="net_spread_bet", aggfunc="sum").round(2))
        year_ws = wb.create_sheet("Each year")
        bt.write_table(year_ws, 1, yearly.reset_index().rename(columns={"year": "Year"}), list(yearly.columns))
        shown = trades[["variant", "signal_day", "entry_day", "exit_day", "company", "ticker", "depth", "deepest",
                        "div_in_dip", "fill", "exit", "outcome", "dividend", "days", "gross", "net_spread_bet",
                        "net_shares"]].copy()
        for col in ("signal_day", "entry_day", "exit_day"):
            shown[col] = pd.to_datetime(shown[col]).dt.date
        shown["depth"] = (shown["depth"] * 100).round(2)
        shown["deepest"] = (shown["deepest"] * 100).round(2)
        shown.columns = ["Rules", "Signal day", "Entry day", "Exit day", "Company", "Ticker", "Below high at signal (%)",
                         "Deepest (%)", "Ex-dividend in slide", "Entry", "Exit", "Result", "Dividend", "Days",
                         "Before costs (£)", "Spread bet (£)", "Real shares (£)"]
        trade_ws = wb.create_sheet("Trades")
        bt.write_table(trade_ws, 1, shown.round(4), ["Before costs (£)", "Spread bet (£)", "Real shares (£)"])
        trade_ws.freeze_panes = "A2"
        csv_path = Path(args.out).with_name(Path(args.out).stem + "_trades.csv")
        shown.to_csv(csv_path, index=False)
    saved = sc.save_workbook(wb, Path(args.out))
    Path(args.out).with_suffix(".json").write_text(json.dumps(
        {"generated": datetime.now(sc.UK_TZ).strftime("%Y-%m-%d %H:%M"), "verdict": verdict,
         "variants": table.where(table.notna(), None).to_dict("records")}, indent=1, default=str))
    for line in verdict:
        print(line)
    print(f"Workbook saved: {saved}")
    return 0


def main() -> None:
    parser = argparse.ArgumentParser(description="Phil's dip method backtest (plan step 3a)")
    parser.add_argument("--universe", default=str(sc.UNIVERSE_PATH))
    parser.add_argument("--start", default="2005-01-01")
    parser.add_argument("--limit", type=int, default=None, help="Only this many extra shares for F (quick test)")
    parser.add_argument("--favourites-only", action="store_true")
    parser.add_argument("--out", default="results/ftse_dip_method.xlsx")
    sys.exit(run(parser.parse_args()))


if __name__ == "__main__":
    main()
