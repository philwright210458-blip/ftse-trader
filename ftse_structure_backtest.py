"""Backtest of the daily structure pullback (ftse_structure.py, idea 24) on FTSE 100 shares, per £250 trade.

    python ftse_structure_backtest.py --index "FTSE 100" --out results/ftse_structure_pullback_ftse100.xlsx

* About 10 years of Yahoo daily candles (split-adjusted; dividends not added).
* Entry at the next day's open after the signal. Each day after: stop counted
  first if both stop and target are inside the day's range; a gap through the
  stop or target fills at the open. Sold at the close after 60 trading days.
* One open trade per share at a time.
* Two targets on the same entries: the measured move (the report's rule) and,
  for comparison, a fixed 2x the stop distance.
* Yardstick: what simply holding the same share for the same number of days
  earned on average over the whole period (its general drift). A strategy whose
  entries add nothing earns about the yardstick.
* Costs per £250: IG share dealing £1.50 (25p price gap + £1.25 stamp duty); no
  stamp duty 25p; IG spread bet 25p + about 4.5p per night.

Caution: today's FTSE 100 list, so companies that dropped out are missing.
"""

from __future__ import annotations

import argparse
import math
import sys
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font

import ftse_backtest as bt
import ftse_ma_daily_backtest as db
import ftse_structure as st
import ftse_trade_scanner as sc

STAKE = 250.0
SHARE_COSTS, NO_STAMP_COSTS, SPREAD = 1.50, 0.25, 0.25
FUNDING_PER_NIGHT = STAKE * 0.065 / 365


def trade(frame: pd.DataFrame, n: int, stop: float, target_of, hold: int = st.MAX_HOLD) -> dict | None:
    """One trade from signal day n. target_of(fill) gives the target once the fill is known."""
    o, h, l, c = (frame[k].to_numpy() for k in ("open", "high", "low", "close"))
    if n + 1 >= len(frame):
        return None
    fill = o[n + 1]
    target = target_of(fill)
    if not (math.isfinite(fill) and stop < fill < target):
        return None
    last = n + hold
    for k in range(n + 1, min(last, len(frame) - 1) + 1):
        if l[k] <= stop:
            return finish(frame, n, fill, stop, target, min(o[k], stop), k, "stop")
        if h[k] >= target:
            return finish(frame, n, fill, stop, target, max(o[k], target), k, "target")
    if last >= len(frame):
        return None                                              # still open at the end of the data
    return finish(frame, n, fill, stop, target, c[last], last, "60 days")


def finish(frame, n, fill, stop, target, price, k, outcome) -> dict:
    gross = (price - fill) / fill * STAKE
    nights = int((frame.index[k] - frame.index[n + 1]).days)
    return {"entry_day": frame.index[n + 1].date(), "exit_day": frame.index[k].date(), "days": k - n,
            "nights": nights, "fill": round(float(fill), 4), "stop": round(float(stop), 4),
            "target": round(float(target), 4), "exit": round(float(price), 4), "outcome": outcome,
            "target_pct": round((target - fill) / fill * 100, 2), "risk_pct": round((fill - stop) / fill * 100, 2),
            "gross": round(gross, 2), "net_shares": round(gross - SHARE_COSTS, 2),
            "net_no_stamp": round(gross - NO_STAMP_COSTS, 2),
            "net_spread_bet": round(gross - SPREAD - FUNDING_PER_NIGHT * nights, 2)}


def drift(frame: pd.DataFrame, days: int) -> float:
    """Average £ per £250 from buying this share at any day's open and selling `days` trading days later."""
    o, c = frame["open"].to_numpy(), frame["close"].to_numpy()
    if len(frame) <= days + 1:
        return float("nan")
    returns = c[days:] / o[1:len(o) - days + 1] - 1
    return float(np.nanmean(returns) * STAKE)


def run_share(frame: pd.DataFrame) -> tuple[list[dict], list[dict]]:
    measured, fixed = [], []
    busy_m = busy_f = -1
    for s in st.signals(frame):
        n = s["t"]
        if n > busy_m:
            t = trade(frame, n, s["stop"], lambda fill, s=s: s["target"])
            if t:
                measured.append({**t, "depth": round(s["depth"], 3)})
                busy_m = n + t["days"]
        if n > busy_f:
            t = trade(frame, n, s["stop"], lambda fill, s=s: fill + 2 * (fill - s["stop"]))
            if t:
                fixed.append({**t, "depth": round(s["depth"], 3)})
                busy_f = n + t["days"]
    return measured, fixed


def summary(name: str, trades: pd.DataFrame, split) -> dict:
    if trades.empty:
        return {"Strategy": name, "Trades": 0}
    entry = pd.to_datetime(trades["entry_day"])
    yearly = trades.groupby(entry.dt.year)["net_shares"].sum()
    return {"Strategy": name, "Trades": len(trades), "Trades a year": len(trades) / max(1, entry.dt.year.nunique()),
            "Avg days held": trades["days"].mean(), "Target hit": (trades["outcome"] == "target").mean(),
            "Stop hit": (trades["outcome"] == "stop").mean(), "Avg target distance": trades["target_pct"].mean() / 100,
            "Made money (IG shares)": (trades["net_shares"] > 0).mean(),
            "Before costs per £250 (£)": trades["gross"].mean(),
            "Yardstick: holding the same share as long (£)": trades["drift"].mean(),
            "After costs: IG shares (£)": trades["net_shares"].mean(),
            "After costs: no stamp duty (£)": trades["net_no_stamp"].mean(),
            "After costs: IG spread bet (£)": trades["net_spread_bet"].mean(),
            "IG shares, first half (£)": trades.loc[entry < split, "net_shares"].mean(),
            "IG shares, second half (£)": trades.loc[entry >= split, "net_shares"].mean(),
            "Years in profit (IG shares)": f"{int((yearly > 0).sum())} of {len(yearly)}"}


def run(args) -> int:
    universe = [c for c in sc.load_universe(Path(args.universe)) if c.kind != "trust"]
    if args.index:
        universe = [c for c in universe if c.index == args.index]
    if args.limit:
        universe = universe[:args.limit]
    names = {c.yahoo_ticker: c.company for c in universe}
    print(f"Downloading {args.period} of daily candles for {len(universe)} shares...")
    data = db.download([c.yahoo_ticker for c in universe], args.period)
    if not data:
        print("No data downloaded (is Yahoo reachable?).")
        return 1
    measured, fixed = [], []
    for ticker, frame in data.items():
        m, f = run_share(frame)
        cache = {}
        for rows, store in ((m, measured), (f, fixed)):
            for row in rows:
                days = row["days"]
                if days not in cache:
                    cache[days] = drift(frame, days)
                store.append({**row, "drift": round(cache[days], 2), "ticker": ticker, "company": names.get(ticker, ticker)})
    mframe, fframe = pd.DataFrame(measured), pd.DataFrame(fixed)
    if mframe.empty:
        print("No trades found.")
        return 1
    entries = sorted(pd.to_datetime(mframe["entry_day"]))
    split = entries[len(entries) // 2]
    rows = [summary("Measured-move target (the report's rule)", mframe, split),
            summary("  pullback 38.2-50%", mframe[mframe["depth"] < 0.5], split),
            summary("  pullback 50-61.8%", mframe[mframe["depth"] >= 0.5], split),
            summary("Fixed target 2x the stop (comparison)", fframe, split)]
    table = pd.DataFrame(rows)
    period = f"{min(f.index[0] for f in data.values()).date()} to {max(f.index[-1] for f in data.values()).date()}"

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.cell(1, 1, "Daily structure pullback (idea 24): per £250 trade").font = Font(name="Arial", bold=True, size=13)
    bt.note(ws, 2, f"Generated {datetime.now(sc.UK_TZ):%Y-%m-%d %H:%M} UK time. {len(data)} {args.index or 'FTSE'} shares, daily "
                   f"candles {period}. Impulse: swing low to swing high of at least 3 typical daily ranges, in an uptrend "
                   "(higher high and higher low), mostly up days with large bodies closing near their highs. Pullback: "
                   "38.2-61.8% of the impulse. Entry: a strong up day closing above the previous 3 days' highs; buy at the "
                   "next open. Stop below the whole pullback. Target: measured move (impulse height added to the pullback "
                   "low), or 60 trading days.")
    bt.note(ws, 3, "Costs per £250: IG share dealing £1.50 (25p price gap + £1.25 stamp duty); no stamp duty 25p; IG spread "
                   "bet 25p + about 4.5p per night held. Yardstick = what holding the same share for the same number of days "
                   "earned on average over the whole period, before costs: the strategy's entries add value only if "
                   "'Before costs' beats it.")
    bt.note(ws, 4, "Caution: today's FTSE 100 list (companies that dropped out are missing), no dividends. First / second "
                   "half = trades split at the middle entry date.")
    money = [c for c in table.columns if "(£)" in c]
    bt.write_table(ws, 6, table.round(4), money, ["Target hit", "Stop hit", "Avg target distance", "Made money (IG shares)"])
    for col, width in zip("ABCDEFGHIJKLMNOPQ", [40, 8, 8, 8, 8, 8, 9, 10, 11, 13, 12, 12, 12, 12, 12, 12]):
        ws.column_dimensions[col].width = width

    year_ws = wb.create_sheet("Each year")
    years = pd.to_datetime(mframe["entry_day"]).dt.year
    yearly = pd.DataFrame({"Trades": mframe.groupby(years).size(),
                           "Total, IG shares (£)": mframe.groupby(years)["net_shares"].sum().round(2),
                           "Avg per trade, IG shares (£)": mframe.groupby(years)["net_shares"].mean().round(2),
                           "Avg yardstick (£)": mframe.groupby(years)["drift"].mean().round(2)})
    yearly.index.name = "Year"
    bt.write_table(year_ws, 1, yearly.reset_index(), [c for c in yearly.columns if "(£)" in c])
    for col in "ABCDE":
        year_ws.column_dimensions[col].width = 22

    trade_ws = wb.create_sheet("Trades")
    shown = mframe.sort_values("entry_day")[["entry_day", "company", "ticker", "depth", "fill", "stop", "target", "outcome",
                                             "exit_day", "days", "gross", "net_shares", "net_spread_bet", "drift"]]
    shown.columns = ["Entry day", "Company", "Ticker", "Pullback depth", "Entry", "Stop", "Target", "Result", "Exit day",
                     "Days held", "Before costs (£)", "After costs: IG shares (£)", "After costs: spread bet (£)",
                     "Yardstick (£)"]
    bt.write_table(trade_ws, 1, shown, [c for c in shown.columns if "(£)" in c])
    trade_ws.freeze_panes = "A2"
    saved = sc.save_workbook(wb, Path(args.out))
    pd.set_option("display.width", 300)
    print(table[["Strategy", "Trades", "Avg days held", "Target hit", "Before costs per £250 (£)",
                 "Yardstick: holding the same share as long (£)", "After costs: IG shares (£)",
                 "After costs: IG spread bet (£)"]].round(2).to_string(index=False))
    print(f"Workbook saved: {saved}")
    return 0


def main() -> None:
    parser = argparse.ArgumentParser(description="Daily structure pullback backtest, per £250")
    parser.add_argument("--universe", default=str(sc.UNIVERSE_PATH))
    parser.add_argument("--index", default="FTSE 100")
    parser.add_argument("--period", default="10y")
    parser.add_argument("--limit", type=int, default=None)
    parser.add_argument("--out", default="results/ftse_structure_pullback_ftse100.xlsx")
    sys.exit(run(parser.parse_args()))


if __name__ == "__main__":
    main()
