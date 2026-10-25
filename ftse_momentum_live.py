"""Paper trader for Claude's momentum model (ftse_momentum.py), run forward from October 2026.

Runs every weekday evening after the outcome tracker (FTSE Outcomes workflow).
It never looks ahead and it catches up if a day is missed:

* Decide: on the last trading day of each month, after the close, it ranks the
  shares exactly as the backtest does (only prices up to that close) and writes
  the planned sells and buys to scan_data/momentum_live/orders.csv.
* Fill: once the next trading day's opening prices exist, the planned orders
  are filled at those opening prices. Sells first; the cash is then split
  equally between the new buys. The first pot is £1,000, i.e. £100 per share.
* Value: every run values the paper portfolio at the latest closing prices,
  credits dividends on shares held over their ex-dividend date, and compares
  it with £1,000 put in a FTSE 100 tracker fund on the same day.

Costs: 0.10% dealing on each buy and sell, 0.5% stamp duty on buys.
Results go to ftse_momentum_live.xlsx. Paper trading only: no real orders.
"""

from __future__ import annotations

import argparse
import math
import sys
import warnings
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.utils import get_column_letter

import ftse_momentum as mm
import ftse_trade_scanner as sc
from ftse_sessions import calendar

FIRST_DECISION = "2026-10-30"          # the last trading day of October 2026
POT = 1000.0
DECIDE_AFTER = "16:45"                 # UK time: month-end decisions only once the closing auction is over
DATA_DIR = Path(__file__).parent / "scan_data" / "momentum_live"
ORDERS_FILE = "orders.csv"
WORKBOOK = Path(__file__).parent / "ftse_momentum_live.xlsx"
WAITING = ("planned", "waiting for opening prices")
ORDER_COLUMNS = ["decided_for", "recorded_at", "action", "ticker", "company", "rank", "score", "reason", "status",
                 "fill_date", "fill_price", "units", "amount_gbp", "costs_gbp"]


# ----------------------------------------------------------------------
# Calendar
# ----------------------------------------------------------------------

def month_end_sessions(start, end) -> list[pd.Timestamp]:
    sessions = calendar().sessions_in_range(pd.Timestamp(start).normalize(), pd.Timestamp(end).normalize())
    if len(sessions) == 0:
        return []
    frame = pd.Series(sessions, index=sessions)
    ends = list(frame.groupby([sessions.year, sessions.month]).max())
    last_of_range = sessions[-1]
    # the range may stop part-way through a month: only keep a month's end if it really is the last session
    return [d for d in ends if d != last_of_range or calendar().next_session(d).month != d.month]


def due_month_ends(now: pd.Timestamp) -> list[pd.Timestamp]:
    """Month-end sessions from the first decision up to now whose closing prices are final."""
    today = now.tz_convert(sc.UK_TZ) if now.tzinfo else now.tz_localize(sc.UK_TZ)
    horizon = today.tz_localize(None).normalize() + pd.Timedelta(days=40)
    due = []
    for day in month_end_sessions(FIRST_DECISION, horizon):
        if day.date() < today.date() or (day.date() == today.date() and today.strftime("%H:%M") >= DECIDE_AFTER):
            due.append(day)
    return due


def next_session(day: pd.Timestamp) -> pd.Timestamp:
    return calendar().next_session(pd.Timestamp(day).normalize())


# ----------------------------------------------------------------------
# Data
# ----------------------------------------------------------------------

class Market:
    """Prices on London trading days: raw (for fills and values) and dividend-adjusted (for ranking)."""

    def __init__(self, open_raw: pd.DataFrame, close_raw: pd.DataFrame, dividends: pd.DataFrame, market: pd.Series):
        first, last = close_raw.index.min(), close_raw.index.max()
        days = calendar().sessions_in_range(first, last)
        open_raw, close_raw = open_raw.reindex(days), close_raw.reindex(days)
        open_raw, close_raw = mm.clean(open_raw, close_raw)
        self.days = days
        self.open, self.close = open_raw, close_raw
        self.close_ff = close_raw.ffill()
        self.paid, _ = mm.corrected_dividends(close_raw, dividends.reindex(days).fillna(0.0))
        _, adjusted, _ = mm.with_dividends(open_raw, close_raw, dividends.reindex(days).fillna(0.0))
        self.adjusted_ff = adjusted.ffill()
        self.sma = self.adjusted_ff.rolling(mm.TREND_DAYS, min_periods=mm.TREND_DAYS).mean()
        self.counts = adjusted.notna().cumsum()
        self.seen = pd.DataFrame(np.where(adjusted.notna(), days.values[:, None], np.datetime64("NaT")),
                                 index=days, columns=adjusted.columns).ffill()
        self.market = market.reindex(days).ffill()

    def has(self, day) -> bool:
        day = pd.Timestamp(day).normalize()
        return day in self.days and bool(self.close.loc[day].notna().any())

    def ranking(self, day) -> pd.DataFrame:
        ranked = mm.ranking(self.adjusted_ff, self.sma, self.counts, self.seen, self.days.get_loc(pd.Timestamp(day)))
        return ranked.drop(index=mm.BENCHMARK_TICKER, errors="ignore")      # the tracker fund is not a pick

    def open_price(self, ticker: str, day) -> float | None:
        """Only the observed open of this exact session may execute an order."""
        day = pd.Timestamp(day)
        if ticker not in self.open.columns or day not in self.open.index:
            return None
        value = self.open.at[day, ticker]
        return float(value) if pd.notna(value) and math.isfinite(value) and value > 0 else None


# ----------------------------------------------------------------------
# The paper portfolio
# ----------------------------------------------------------------------

def read_orders(folder: Path) -> pd.DataFrame:
    path = folder / ORDERS_FILE
    if not path.exists() or path.stat().st_size == 0:
        return pd.DataFrame(columns=ORDER_COLUMNS)
    return pd.read_csv(path, dtype={"decided_for": str, "fill_date": str, "ticker": str})


def write_orders(folder: Path, orders: pd.DataFrame) -> None:
    folder.mkdir(parents=True, exist_ok=True)
    temporary = folder / (ORDERS_FILE + ".tmp")
    orders[ORDER_COLUMNS].to_csv(temporary, index=False)
    temporary.replace(folder / ORDERS_FILE)


def positions(orders: pd.DataFrame) -> list[dict]:
    """Every holding, open or closed, rebuilt from filled orders (one holding per share at a time)."""
    filled = orders[orders["status"] == "filled"].sort_values(["fill_date", "action"], ascending=[True, False])
    open_now, done = {}, []
    for _, row in filled.iterrows():
        if row["action"] == "Buy":
            open_now[row["ticker"]] = {"ticker": row["ticker"], "company": row["company"], "bought": row["fill_date"],
                                       "buy_price": float(row["fill_price"]), "units": float(row["units"]),
                                       "cost": float(row["amount_gbp"]), "sold": None, "proceeds": None, "why": ""}
        elif row["action"] == "Sell" and row["ticker"] in open_now:
            held = open_now.pop(row["ticker"])
            held.update(sold=row["fill_date"], proceeds=float(row["amount_gbp"]), why=row["reason"])
            done.append(held)
    return done + list(open_now.values())


def cash_now(orders: pd.DataFrame) -> float:
    filled = orders[orders["status"] == "filled"]
    spent = filled.loc[filled["action"] == "Buy", "amount_gbp"].astype(float).sum()
    received = filled.loc[filled["action"] == "Sell", "amount_gbp"].astype(float).sum()
    return POT - spent + received


def dividends_for(position: dict, market: Market, until) -> float:
    """Dividends earned: ex-dividend date after the buy and no later than the sale (sold at that morning's open)."""
    ticker = position["ticker"]
    if ticker not in market.paid.columns:
        return 0.0
    paid = market.paid[ticker]
    end = pd.Timestamp(position["sold"]) if position["sold"] else pd.Timestamp(until)
    window = paid[(paid.index > pd.Timestamp(position["bought"])) & (paid.index <= min(end, pd.Timestamp(until)))]
    return float(window.sum() * position["units"])


def decide(day: pd.Timestamp, orders: pd.DataFrame, market: Market, names: dict, now: pd.Timestamp,
           settings: mm.Settings) -> pd.DataFrame:
    held = [p["ticker"] for p in positions(orders) if not p["sold"]]
    ranked = market.ranking(day)
    in_market = mm.market_ok(market.market.dropna(), day)
    if in_market:
        keep_zone = set(ranked.index[:settings.keep])
        target = [t for t in held if t in keep_zone]
        for ticker in ranked.index:
            if len(target) >= settings.top:
                break
            if ticker not in target:
                target.append(ticker)
    else:
        target = []
    rank_of = {t: n for n, t in enumerate(ranked.index, 1)}
    stamp = now.tz_convert(sc.UK_TZ).strftime("%Y-%m-%d %H:%M")
    rows = []

    def row(action, ticker, reason):
        score = ranked["score"].get(ticker) if ticker in ranked.index else None
        rows.append({"decided_for": str(day.date()), "recorded_at": stamp, "action": action, "ticker": ticker,
                     "company": names.get(ticker, ticker), "rank": rank_of.get(ticker),
                     "score": round(float(score), 4) if score is not None else None, "reason": reason,
                     "status": "planned"})

    for ticker in held:
        if ticker not in target:
            reason = ("Market filter: FTSE 100 below its 200-day average" if not in_market
                      else "Below its 200-day average" if ticker not in ranked.index
                      else f"Dropped to rank {rank_of[ticker]} (outside the top {settings.keep})")
            row("Sell", ticker, reason)
    for ticker in target:
        if ticker not in held:
            row("Buy", ticker, f"Rank {rank_of[ticker]} by 12-1 month price change")
    if not rows:
        row("No change", "", "Holdings unchanged" if in_market else "Market filter: staying in cash")
    new = pd.DataFrame(rows).reindex(columns=ORDER_COLUMNS).astype(object)
    if orders.empty:
        return new
    with warnings.catch_warnings():          # the not-yet-filled columns are empty by design
        warnings.simplefilter("ignore", FutureWarning)
        return pd.concat([orders.astype(object), new], ignore_index=True)


def fill(day: pd.Timestamp, orders: pd.DataFrame, market: Market) -> pd.DataFrame:
    """Fill the orders decided on `day` at the next session's opening prices."""
    orders = orders.copy()
    for column in ("decided_for", "recorded_at", "action", "ticker", "company", "reason", "status", "fill_date"):
        orders[column] = orders[column].astype(object)
    trade_day = next_session(day)
    mine = orders.index[(orders["decided_for"] == str(day.date())) & orders["status"].isin(WAITING)]
    holdings = {p["ticker"]: p for p in positions(orders) if not p["sold"]}
    required = [i for i in mine if orders.at[i, "action"] in ("Buy", "Sell")]
    if any(market.open_price(orders.at[i, "ticker"], trade_day) is None for i in required):
        # Wait for the whole rebalance. This reserves cash for every intended buy
        # and avoids spending proceeds from a sale with no observed opening price.
        orders.loc[required, "status"] = "waiting for opening prices"
        return orders
    for i in mine:
        if orders.at[i, "action"] == "No change":
            orders.loc[i, ["status", "fill_date"]] = ["done", str(trade_day.date())]
    for i in [i for i in mine if orders.at[i, "action"] == "Sell"]:
        ticker = orders.at[i, "ticker"]
        price = market.open_price(ticker, trade_day)
        if price is None or ticker not in holdings:
            orders.at[i, "status"] = "not filled (no holding)"
            continue
        units = holdings[ticker]["units"]
        gross = units * price
        orders.loc[i, ["status", "fill_date", "fill_price", "units", "amount_gbp", "costs_gbp"]] = [
            "filled", str(trade_day.date()), round(price, 4), units, round(gross * (1 - mm.DEALING), 2),
            round(gross * mm.DEALING, 2)]
    buys = [(i, market.open_price(orders.at[i, "ticker"], trade_day)) for i in mine if orders.at[i, "action"] == "Buy"]
    cash = cash_now(orders)
    if buys and cash > 1:
        each = math.floor(cash / len(buys) * 100) / 100
        for i, price in buys:
            units = each / (price * (1 + mm.DEALING + mm.STAMP))
            orders.loc[i, ["status", "fill_date", "fill_price", "units", "amount_gbp", "costs_gbp"]] = [
                "filled", str(trade_day.date()), round(price, 4), units, each,
                round(each - units * price, 2)]
    elif buys:
        orders.loc[[i for i, _ in buys], "status"] = "not filled (insufficient cash)"
    return orders


def step(orders: pd.DataFrame, market: Market, names: dict, now: pd.Timestamp,
         settings: mm.Settings | None = None) -> pd.DataFrame:
    """Make every decision and fill that is due, in order, stopping at the first that has to wait for data."""
    settings = settings or mm.Settings()
    decided = set(orders["decided_for"].astype(str))
    today = now.tz_convert(sc.UK_TZ).tz_localize(None).normalize()
    for day in due_month_ends(now):
        if str(day.date()) not in decided:
            if not market.has(day):
                break                                           # that month-end's prices are not in yet
            orders = decide(day, orders, market, names, now, settings)
        waiting = (orders["decided_for"] == str(day.date())) & orders["status"].isin(WAITING)
        if waiting.any():
            trade_day = next_session(day)
            if trade_day > today or not market.has(trade_day):
                break                                           # fills wait for the next morning's prices
            orders = fill(day, orders, market)
            if ((orders["decided_for"] == str(day.date())) & orders["status"].isin(WAITING)).any():
                break  # unresolved rebalance must not leak cash/holdings into later decisions
    return orders


# ----------------------------------------------------------------------
# Valuing and reporting
# ----------------------------------------------------------------------

def value_history(orders: pd.DataFrame, market: Market) -> pd.Series:
    filled = orders[orders["status"] == "filled"]
    if filled.empty:
        return pd.Series(dtype=float)
    first = pd.Timestamp(filled["fill_date"].min())
    all_positions = positions(orders)
    values = {}
    for day in market.days[market.days >= first]:
        cash = POT
        total = 0.0
        for p in all_positions:
            bought = pd.Timestamp(p["bought"])
            if bought > day:
                continue
            cash -= p["cost"]
            if p["sold"] and pd.Timestamp(p["sold"]) <= day:
                cash += p["proceeds"]
            else:
                price = market.close_ff.at[day, p["ticker"]]
                total += p["units"] * (price if math.isfinite(price) else p["buy_price"])
            cash += dividends_for(p, market, day)
        values[day] = cash + total
    return pd.Series(values)


def tracker_history(market: Market, start, open_: pd.Series, close: pd.Series, paid: pd.Series) -> pd.Series:
    start = pd.Timestamp(start)
    price = open_.get(start, np.nan)
    if close is None or not (isinstance(price, float) and math.isfinite(price) and price > 0):
        return pd.Series(dtype=float)
    units = POT * (1 - mm.DEALING) / price
    later = close.ffill().loc[start:]
    income = paid.loc[start:].where(paid.loc[start:].index > start, 0.0).cumsum() * units
    return later * units + income.reindex(later.index).fillna(0.0)


def build_workbook(orders: pd.DataFrame, market: Market, names: dict, tracker: pd.Series, now: pd.Timestamp,
                   latest_rank: pd.DataFrame) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    from openpyxl.styles import Font
    ws.cell(1, 1, "Claude's model: paper trading (monthly momentum)").font = Font(name="Arial", bold=True, size=13)
    last_day = market.days[-1]
    values = value_history(orders, market)
    all_positions = positions(orders)
    upcoming = [d for d in month_end_sessions(last_day, last_day + pd.Timedelta(days=45)) if d >= last_day]
    filter_on = mm.market_ok(market.market.dropna(), last_day)
    rows = [("Prices up to", str(last_day.date())), ("Updated", now.tz_convert(sc.UK_TZ).strftime("%Y-%m-%d %H:%M UK time"))]
    if values.empty:
        rows += [("Status", f"Waiting for the first decision after the close on {FIRST_DECISION}; first paper trades at "
                            f"the open on {next_session(pd.Timestamp(FIRST_DECISION)).date()}. Starting pot £1,000 "
                            "(10 shares at £100 each).")]
    else:
        open_pos = [p for p in all_positions if not p["sold"]]
        closed = [p for p in all_positions if p["sold"]]
        dividends = sum(dividends_for(p, market, last_day) for p in all_positions)
        costs = orders.loc[orders["status"] == "filled", "costs_gbp"].astype(float).sum()
        rows += [("Started", str(values.index[0].date())), ("Value now (£)", round(float(values.iloc[-1]), 2)),
                 ("Change since start (£)", round(float(values.iloc[-1] - POT), 2)),
                 ("Change since start (%)", round(float(values.iloc[-1] / POT - 1) * 100, 2))]
        if not tracker.empty:
            rows += [("£1,000 in a FTSE 100 tracker on the same day is now (£)", round(float(tracker.iloc[-1]), 2))]
        rows += [("Cash (£)", round(cash_now(orders) + sum(dividends_for(p, market, last_day) for p in all_positions), 2)),
                 ("Shares held", len(open_pos)), ("Closed trades", len(closed)),
                 ("Dividends received (£)", round(dividends, 2)), ("Dealing costs and stamp duty paid (£)", round(costs, 2))]
        if closed:
            per100 = [((p["proceeds"] + dividends_for(p, market, last_day)) / p["cost"] - 1) * 100 for p in closed]
            rows += [("Closed trades: average per £100 after costs (£)", round(float(np.mean(per100)), 2)),
                     ("Closed trades that made money", f"{np.mean([x > 0 for x in per100]):.0%}")]
    rows += [("Next decision", f"after the close on {upcoming[0].date()}" if upcoming else "next month end"),
             ("FTSE 100 versus its 200-day average", "above: the model may hold shares" if filter_on
              else "below: the model would move to cash")]
    for r, (label, value) in enumerate(rows, 3):
        ws.cell(r, 1, label).font = Font(name="Arial", bold=True)
        ws.cell(r, 2, value).font = Font(name="Arial")
    note_row = len(rows) + 4
    ws.cell(note_row, 1, "Rules: on the last trading day of each month, rank shares by price change from 12 months ago to "
                         "1 month ago; hold the best 10 above their 200-day average; keep a holding while it stays in the "
                         "top 20; hold cash when the FTSE 100 is below its 200-day average. Orders fill at the next "
                         "morning's open. Costs: 0.10% per buy and sell, 0.5% stamp duty on buys. Paper trading only.")
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 58, 60

    hold_ws = wb.create_sheet("Holdings")
    header = ["Company", "Share", "Bought", "Buy price", "Cost (£)", "Value now (£)", "Dividends (£)",
              "Profit / loss (£)", "Per £100 (£)"]
    hold_ws.append(header)
    for p in [p for p in all_positions if not p["sold"]]:
        price = market.close_ff.at[last_day, p["ticker"]]
        value = p["units"] * price
        dividend = dividends_for(p, market, last_day)
        hold_ws.append([p["company"], p["ticker"], p["bought"], round(p["buy_price"], 2), round(p["cost"], 2),
                        round(value, 2), round(dividend, 2), round(value + dividend - p["cost"], 2),
                        round((value + dividend) / p["cost"] * 100 - 100, 2)])

    closed_ws = wb.create_sheet("Closed trades")
    closed_ws.append(["Company", "Share", "Bought", "Sold", "Cost (£)", "Proceeds (£)", "Dividends (£)",
                      "Profit / loss (£)", "Per £100 (£)", "Why sold"])
    for p in [p for p in all_positions if p["sold"]]:
        dividend = dividends_for(p, market, last_day)
        closed_ws.append([p["company"], p["ticker"], p["bought"], p["sold"], round(p["cost"], 2), round(p["proceeds"], 2),
                          round(dividend, 2), round(p["proceeds"] + dividend - p["cost"], 2),
                          round((p["proceeds"] + dividend) / p["cost"] * 100 - 100, 2), p["why"]])

    log_ws = wb.create_sheet("Decisions")
    shown = orders.rename(columns={"decided_for": "Decided at close of", "recorded_at": "Recorded", "action": "Action",
                                   "ticker": "Share", "company": "Company", "rank": "Rank", "score": "12-1 month change",
                                   "reason": "Reason", "status": "Status", "fill_date": "Filled on",
                                   "fill_price": "Fill price", "units": "Units", "amount_gbp": "Amount (£)",
                                   "costs_gbp": "Costs (£)"})
    log_ws.append(list(shown.columns))
    for _, row in shown.iterrows():
        log_ws.append([None if (isinstance(v, float) and not math.isfinite(v)) else v for v in row.tolist()])

    value_ws = wb.create_sheet("Value over time")
    value_ws.append(["Date", "Claude's model (£)", "FTSE 100 tracker (£)"])
    for day, value in values.items():
        value_ws.append([str(day.date()), round(float(value), 2),
                         round(float(tracker.get(day, np.nan)), 2) if day in tracker.index else None])
    if len(values) > 1:
        chart = LineChart()
        chart.title, chart.y_axis.title, chart.height, chart.width = "Value of £1,000", "£", 10, 22
        chart.add_data(Reference(value_ws, min_col=2, max_col=3, min_row=1, max_row=len(values) + 1), titles_from_data=True)
        chart.set_categories(Reference(value_ws, min_col=1, min_row=2, max_row=len(values) + 1))
        value_ws.add_chart(chart, "E2")

    rank_ws = wb.create_sheet("Ranking now")
    rank_ws.append([f"How the model would rank shares on {last_day.date()} (it only trades at month ends)"])
    rank_ws.append(["Rank", "Company", "Share", "12-1 month change", "Price", "200-day average"])
    for n, (ticker, row) in enumerate(latest_rank.head(20).iterrows(), 1):
        rank_ws.append([n, names.get(ticker, ticker), ticker, f"{row['score']:.1%}", round(float(row["price"]), 2),
                        round(float(row["average_200d"]), 2)])
    for sheet in wb.worksheets[1:]:
        for c in range(1, 15):
            sheet.column_dimensions[get_column_letter(c)].width = 16
        sheet.column_dimensions["A"].width = 28
    return wb


# ----------------------------------------------------------------------

def run(args, now: pd.Timestamp | None = None) -> int:
    now = now or pd.Timestamp.now(tz=sc.UK_TZ)
    universe = [c for c in sc.load_universe(Path(args.universe)) if c.kind != "trust"]
    names = {c.yahoo_ticker: c.company for c in universe}
    tickers = [c.yahoo_ticker for c in universe]
    start = (now.tz_convert(sc.UK_TZ).tz_localize(None) - pd.Timedelta(days=args.history_days)).strftime("%Y-%m-%d")
    print(f"Momentum paper trader: downloading daily prices for {len(tickers)} shares from {start}...")
    open_, close, dividends = mm.download(tickers + [mm.MARKET_TICKER, mm.BENCHMARK_TICKER], start)
    if close.empty or mm.MARKET_TICKER not in close.columns:
        print("No data downloaded (is Yahoo reachable?). Nothing changed.")
        return 1
    market_series = close[mm.MARKET_TICKER]
    shares = [t for t in tickers if t in close.columns]
    columns = shares + ([mm.BENCHMARK_TICKER] if mm.BENCHMARK_TICKER in close.columns else [])
    market = Market(open_[columns], close[columns], dividends.reindex(columns=columns).fillna(0.0), market_series)
    folder = Path(args.data_dir)
    orders = step(read_orders(folder), market, names, now)
    write_orders(folder, orders)
    filled = orders[orders["status"] == "filled"]
    tracker = pd.Series(dtype=float)
    if not filled.empty and mm.BENCHMARK_TICKER in market.close.columns:
        tracker = tracker_history(market, filled["fill_date"].min(), market.open[mm.BENCHMARK_TICKER],
                                  market.close[mm.BENCHMARK_TICKER], market.paid[mm.BENCHMARK_TICKER])
    ranking = market.ranking(market.days[-1])
    book = build_workbook(orders, market, names, tracker, now, ranking)
    saved = sc.save_workbook(book, Path(args.out))
    latest = orders.tail(12)
    print(latest[["decided_for", "action", "company", "status", "fill_date", "amount_gbp"]].to_string(index=False)
          if not latest.empty else f"No decisions yet: the first is after the close on {FIRST_DECISION}.")
    print(f"Workbook saved: {saved}")
    return 0


def main() -> None:
    parser = argparse.ArgumentParser(description="Paper-trade Claude's momentum model forward from October 2026")
    parser.add_argument("--universe", default=str(sc.UNIVERSE_PATH))
    parser.add_argument("--data-dir", default=str(DATA_DIR))
    parser.add_argument("--out", default=str(WORKBOOK))
    parser.add_argument("--history-days", type=int, default=560, help="Calendar days of prices to download (default 560)")
    sys.exit(run(parser.parse_args()))


if __name__ == "__main__":
    main()
