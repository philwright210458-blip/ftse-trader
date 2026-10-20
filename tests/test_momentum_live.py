import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest import mock

import numpy as np
import openpyxl
import pandas as pd

import ftse_momentum as mm
import ftse_momentum_live as live
from ftse_sessions import calendar

UK = "Europe/London"
DAYS = calendar().sessions_in_range("2025-01-02", "2027-01-29")
TICKERS = [f"T{n}.L" for n in range(14)]
NAMES = {t: f"Company {t}" for t in TICKERS}


def frames(seed=1, until=None):
    rng = np.random.default_rng(seed)
    drift = np.r_[[0.5] * 6, [0.0] * 8] / 252
    steps = drift + rng.normal(0, 0.01, (len(DAYS), len(TICKERS)))
    close = pd.DataFrame(100 * np.exp(np.cumsum(steps, axis=0)), index=DAYS, columns=TICKERS)
    close[mm.BENCHMARK_TICKER] = np.linspace(800, 900, len(DAYS))
    open_ = close * 1.001
    dividends = pd.DataFrame(0.0, index=DAYS, columns=close.columns)
    market = pd.Series(np.linspace(8000, 10000, len(DAYS)), index=DAYS)
    if until is not None:
        keep = DAYS <= pd.Timestamp(until)
        open_, close, dividends, market = open_[keep], close[keep], dividends[keep], market[keep]
    return open_, close, dividends, market


def market_until(until, seed=1, tweak=None):
    open_, close, dividends, market = frames(seed, until)
    if tweak:
        tweak(open_, close, dividends)
    return live.Market(open_, close, dividends, market)


def at(text):
    return pd.Timestamp(text, tz=UK)


class LiveMomentumTests(unittest.TestCase):
    def test_month_ends_and_due_dates(self):
        ends = live.month_end_sessions("2026-10-01", "2026-12-31")
        self.assertEqual([str(d.date()) for d in ends], ["2026-10-30", "2026-11-30", "2026-12-31"])
        self.assertEqual(live.due_month_ends(at("2026-10-30 16:00")), [])
        self.assertEqual([str(d.date()) for d in live.due_month_ends(at("2026-10-30 17:05"))], ["2026-10-30"])

    def test_nothing_before_first_decision(self):
        orders = live.step(pd.DataFrame(columns=live.ORDER_COLUMNS), market_until("2026-10-20"), NAMES, at("2026-10-20 17:05"))
        self.assertTrue(orders.empty)

    def test_decide_at_month_end_then_fill_next_morning(self):
        empty = pd.DataFrame(columns=live.ORDER_COLUMNS)
        decided = live.step(empty, market_until("2026-10-30"), NAMES, at("2026-10-30 17:05"))
        buys = decided[decided["action"] == "Buy"]
        self.assertEqual(len(buys), 10)
        self.assertTrue((decided["status"] == "planned").all())
        self.assertNotIn(mm.BENCHMARK_TICKER, set(decided["ticker"]))
        self.assertTrue(set(TICKERS[:6]) <= set(buys["ticker"]))      # the strong risers are chosen

        market = market_until("2026-11-02")
        filled = live.step(decided, market, NAMES, at("2026-11-02 17:05"))
        done = filled[filled["status"] == "filled"]
        self.assertEqual(len(done), 10)
        self.assertEqual(set(done["fill_date"]), {"2026-11-02"})
        self.assertAlmostEqual(done["amount_gbp"].astype(float).sum(), 1000.0, delta=0.1)
        first = done.iloc[0]
        self.assertAlmostEqual(float(first["fill_price"]), float(market.open.at[pd.Timestamp("2026-11-02"), first["ticker"]]), 3)
        again = live.step(filled, market, NAMES, at("2026-11-02 17:40"))
        pd.testing.assert_frame_equal(again.reset_index(drop=True), filled.reset_index(drop=True))

    def test_catches_up_after_missed_runs(self):
        orders = live.step(pd.DataFrame(columns=live.ORDER_COLUMNS), market_until("2026-12-15"), NAMES, at("2026-12-15 17:05"))
        self.assertEqual(set(orders["decided_for"]), {"2026-10-30", "2026-11-30"})
        self.assertFalse((orders["status"] == "planned").any())

    def test_decision_ignores_later_prices(self):
        def wreck(open_, close, dividends):
            later = close.index > pd.Timestamp("2026-10-30")
            close.loc[later, TICKERS[:6]] *= 0.2
            open_.loc[later, TICKERS[:6]] *= 0.2
        plain = live.step(pd.DataFrame(columns=live.ORDER_COLUMNS), market_until("2026-11-20"), NAMES, at("2026-10-30 17:05"))
        changed = live.step(pd.DataFrame(columns=live.ORDER_COLUMNS), market_until("2026-11-20", tweak=wreck), NAMES,
                            at("2026-10-30 17:05"))
        self.assertEqual(list(plain["ticker"]), list(changed["ticker"]))

    def test_dividends_credited_while_held(self):
        def pay(open_, close, dividends):
            dividends.at[pd.Timestamp("2026-11-10"), TICKERS[0]] = 2.0
        market = market_until("2026-11-20", tweak=pay)
        orders = live.step(pd.DataFrame(columns=live.ORDER_COLUMNS), market, NAMES, at("2026-11-20 17:05"))
        position = [p for p in live.positions(orders) if p["ticker"] == TICKERS[0]][0]
        self.assertAlmostEqual(live.dividends_for(position, market, pd.Timestamp("2026-11-20")), 2.0 * position["units"])
        values = live.value_history(orders, market)
        self.assertEqual(str(values.index[0].date()), "2026-11-02")

    def test_run_writes_orders_and_workbook(self):
        open_, close, dividends, market = frames(until="2026-11-03")
        open_[mm.MARKET_TICKER], close[mm.MARKET_TICKER], dividends[mm.MARKET_TICKER] = market, market, 0.0
        lines = "\n".join(f"Company {t},{t[:-2]},{t},FTSE 250,company" for t in TICKERS)
        with tempfile.TemporaryDirectory() as folder:
            universe = Path(folder) / "u.csv"
            universe.write_text("company,epic,yahoo_ticker,index,type\n" + lines + "\n")
            args = SimpleNamespace(universe=str(universe), data_dir=str(Path(folder) / "live"),
                                   out=str(Path(folder) / "live.xlsx"), history_days=700)
            with mock.patch.object(mm, "download", return_value=(open_, close, dividends)):
                self.assertEqual(live.run(args, now=at("2026-11-03 17:05")), 0)
            book = openpyxl.load_workbook(args.out)
            self.assertEqual(book.sheetnames, ["Summary", "Holdings", "Closed trades", "Decisions", "Value over time",
                                               "Ranking now"])
            self.assertEqual(book["Holdings"].max_row, 11)
            orders = pd.read_csv(Path(args.data_dir) / live.ORDERS_FILE)
            self.assertEqual((orders["status"] == "filled").sum(), 10)


if __name__ == "__main__":
    unittest.main()
