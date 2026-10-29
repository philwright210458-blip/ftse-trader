"""Trade feasibility backtest, with self-search for improvements (29 Oct 2026).

Follows the Haleon v Admiral trade autopsy and Phil's research plan (claude/trade-feasibility-research.md).
FTSE 100 shares only, Yahoo's last ~60 days of 15-minute candles, the live scanner's rules (target cap,
spread-bet costs 0.10%, fill check from 28 Oct, up to 5 trades open and 10 a day, £100 a trade).

Earlier backtests mostly tested filters and found little. A filter can only remove trades, so it cannot
create an edge the signal does not have. This backtest therefore asks four questions:

  1. PRE-SET RULES (chosen before seeing any results), each replayed through the live trader:
       B0-B3  breakout still intact at the fill: the fill must still be past the zone by 0, 0.1, 0.25 or
              0.5 typical candle ranges (breakout / "reclaim" setups only)
       C      at least 75% of the planned stop room left at the fill
       D      stop not already traded through earlier that day
       E      no trades in a sideways regime (flat 9 EMA / 20 SMA plus a flip or price near the daily
              50-day and 4-hour 200 averages)
       F      zone at least 2 trading hours old
  2. SELF-SEARCH (discovery done safely): every recorded feature is split into groups; the FIRST half of the
     period is searched for the groups worth skipping or keeping, ranked by how much they raise the AVERAGE
     per trade (not total pounds, which rewards simply trading less while trades lose); the best few are then
     checked on the SECOND half, which the search never saw, against random sets of the same number of trades.
     Only candidates that beat over 99% of those random sets (5% shared across the six candidates) are
     reported as worth a live test.
  3. EXITS: the same entries with exits fixed in advance: current target and stop, half the target, hold
     to 16:30 (stop kept), stop to break-even after a gain equal to the stop distance, 2-hour time limit,
     and direction only (no stop or target: open to 16:30).
  4. CHANCE: is the scanner better than random? Each trade is compared with (a) the same trade in a random
     direction and (b) the same trade at a random time that day; 1,000 random versions of the whole set
     give the share of random versions that did as well as the scanner.

Results: results/ftse_feasibility.xlsx (sheets Verdict, Rules, Self-search, Exits, Chance, Trades) and
results/ftse_feasibility.json (for the results page). Paper research only: nothing changes live trading.

    python ftse_feasibility_backtest.py --out results/ftse_feasibility.xlsx
"""

from __future__ import annotations

import argparse
import json
import math
import os
import sys
from concurrent.futures import ProcessPoolExecutor
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook

import ftse_backtest as bt
import ftse_entry_check_backtest as ec
import ftse_model_check_backtest as mc
import ftse_outcomes as out
import ftse_risk_backtest as rb
import ftse_trade_notes as tn
import ftse_trade_scanner as sc
from ftse_picker import load_views
from ftse_sessions import is_session

COST_PCT, STAKE = 0.10, 100.0
BUFFERS = [0.0, 0.1, 0.25, 0.5]           # breakout-intact buffers, in typical candle ranges (set in advance)
ROOM_KEPT = 0.75                          # rule C
ZONE_MIN_MINUTES = 120                    # rule F
MIN_GROUP_FIRST_HALF = 20                 # self-search: smallest group considered in the first half
MIN_GROUP_TOTAL = 30                      # ... and in both halves together
MIN_GROUP_SECOND_HALF = 10
TOP_CANDIDATES = 6
LUCK_LIMIT = 0.05 / TOP_CANDIDATES        # several candidates are checked, so each must beat luck more clearly
RANDOM_RUNS = 1000
SEED = 29102026


# ----------------------------------------------------------------------
# Replaying the scanner, keeping the zone and the typical candle range
# ----------------------------------------------------------------------

def share_setups(job) -> list[dict]:
    """As ftse_model_check_backtest.share_setups, also keeping the zone's touches and the typical range."""
    company, data, cuts, interval_min = job
    step = pd.Timedelta(minutes=interval_min)
    ends = data.index + step
    found = []
    for cut in cuts:
        visible = data[ends <= cut]
        cfg = sc.INDEX_RULES[company.index]
        if len(visible) <= cfg.swing_lookback * 2 or visible.index[-1].date() != cut.date():
            continue
        zones = sc.qualify_zones(sc.cluster_zones(sc.find_swing_points(visible, cfg.swing_lookback), cfg.zone_tolerance_pct),
                                 visible, cfg)
        setups = sc.build_setups(company.company, company.yahoo_ticker, company.index, visible, zones, cfg)
        atr = sc.typical_true_range(visible)
        for setup in setups:
            setup.kind, setup.last_trade = company.kind, visible.index[-1] + step
            sc.cap_target(setup, visible)
            setup.filter_note = sc.entry_rule_note(setup, visible)
            first = min(point.ts for point in setup.zone.touches) if setup.zone.touches else None
            found.append({"cut": cut, "setup": setup, "atr": atr, "first_touch": first})
    return found


def replay(universe, intraday, interval_min: int, warmup_days: int, workers: int) -> pd.DataFrame:
    """Every idea the live scanner would have offered at every scan (same steps as the model check replay)."""
    all_days = sorted({stamp.normalize() for data in intraday.values() for stamp in data.index})
    sessions = [day for day in all_days if is_session(day)][warmup_days:]
    cuts = [cut for day in sessions for cut in rb.scan_cuts(day)]
    jobs = [(c, intraday[c.yahoo_ticker], cuts, interval_min) for c in universe if c.yahoo_ticker in intraday]
    print(f"Replaying {len(cuts)} scans over {len(sessions)} days for {len(jobs)} shares ({workers} workers)...")
    raw = []
    with ProcessPoolExecutor(max_workers=workers) as pool:
        for number, rows in enumerate(pool.map(share_setups, jobs, chunksize=1), 1):
            raw.extend(rows)
            if number % 10 == 0:
                print(f"  ...{number}/{len(jobs)} shares replayed")
    by_cut: dict = {}
    for row in raw:
        by_cut.setdefault(row["cut"], []).append(row)
    rows = []
    for cut in sorted(by_cut):
        items = by_cut[cut]
        setups = [item["setup"] for item in items]
        extra = {id(item["setup"]): item for item in items}
        newest = max(s.last_trade for s in setups)
        sc.mark_stale(setups, newest, 30)
        fresh = [s for s in setups if not s.stale]
        sc.annotate_conflicts(fresh)
        sc.apply_entry_rules(fresh)
        for s in sc.dedupe_setups(fresh):
            if s.conflict:
                continue
            info = extra.get(id(s), {})
            rows.append({"day": str(cut.date()), "as_of": cut.strftime("%H:%M"),
                         "data_time": (newest - pd.Timedelta(minutes=interval_min)).isoformat(),
                         "run_time": (cut + rb.RUN_DELAY).strftime("%Y-%m-%d %H:%M:%S"), "mode": "watch",
                         "company": s.company, "ticker": s.ticker, "index": s.index_label, "setup": s.setup_type,
                         "direction": s.direction, "score": s.quality_score, "entry": s.entry, "stop": s.stop,
                         "target": s.target, "stop_pct": abs(s.entry - s.stop) / s.entry * 100,
                         "zone_low": s.zone.low, "zone_high": s.zone.high, "zone_touches": s.zone.touch_count,
                         "first_touch": info.get("first_touch"), "atr": info.get("atr")})
    return pd.DataFrame(rows)


# ----------------------------------------------------------------------
# Signal-time features for every idea (regime, zone age)
# ----------------------------------------------------------------------

def add_signal_features(ideas: pd.DataFrame, intraday: dict, daily: dict, closes4: dict) -> pd.DataFrame:
    """regime and zone age at the signal, using only what was known then."""
    averages = {}
    for ticker, frame in intraday.items():
        close = frame["close"].astype(float)
        averages[ticker] = (close.ewm(span=9, adjust=False).mean(), close.rolling(20).mean())
    regimes, ages = [], []
    ideas = ideas.sort_values(["day", "run_time"]).reset_index(drop=True)
    run_times: dict = {}
    for (day, ticker, direction), part in ideas.groupby(["day", "ticker", "direction"]):
        run_times[(day, ticker, direction)] = [out.market_stamp(t) for t in part["run_time"]]
    for _, row in ideas.iterrows():
        ticker = row["ticker"]
        signal_candle = out.market_stamp(row["data_time"])
        ema, sma = averages.get(ticker, (None, None))
        state = ""
        if ema is not None and signal_candle in ema.index:
            n = ema.index.get_loc(signal_candle)
            e, s = float(ema.iloc[n]), float(sma.iloc[n])
            if n >= tn.SLOPE_CANDLES and math.isfinite(s):
                gap = (e - s) / s * 100
                slope = (s - float(sma.iloc[n - tn.SLOPE_CANDLES])) / s * 100
                state = "Up" if gap > tn.FLAT_GAP_PCT and slope > 0 else "Down" if gap < -tn.FLAT_GAP_PCT and slope < 0 else "Flat"
        days = daily.get(ticker)
        d_trend, d_pct = "", None
        if days is not None:
            before = days[days.index < row["day"]]
            if len(before) >= 50:
                last, avg = float(before["close"].iloc[-1]), float(before["close"].astype(float).tail(50).mean())
                d_trend, d_pct = tn.trend_word(last > avg), (last - avg) / avg * 100
        h4_trend, h4_pct = "", None
        series = closes4.get(ticker)
        if series is not None and len(series):
            done = series[series.index <= signal_candle + pd.Timedelta(minutes=15)]
            if len(done) >= tn.SMA_4H_BARS:
                avg = float(done.tail(tn.SMA_4H_BARS).mean())
                h4_trend, h4_pct = tn.trend_word(float(row["entry"]) > avg), (float(row["entry"]) - avg) / avg * 100
        signal_time = out.market_stamp(row["run_time"])
        same = run_times.get((row["day"], ticker, "Short" if row["direction"] == "Long" else "Long"), [])
        flips = int(sum(signal_time - tn.FLIP_WINDOW <= t <= signal_time for t in same))
        regimes.append(tn.regime_label(state, flips, d_pct, h4_pct, d_trend, h4_trend) if state else "")
        age = None
        first = row.get("first_touch")
        frame = intraday.get(ticker)
        if first is not None and not pd.isna(first) and frame is not None:
            age = int(((frame.index >= first) & (frame.index <= signal_candle)).sum()) * 15
        ages.append(age)
    return ideas.assign(regime=regimes, zone_age_minutes=ages)


# ----------------------------------------------------------------------
# 1. Pre-set rules as fill-time checks
# ----------------------------------------------------------------------

def fill_candle(row) -> pd.Timestamp:
    signal_close = out.market_stamp(row["data_time"]) + pd.Timedelta(minutes=15)
    return max(signal_close, out.market_stamp(row["run_time"])).ceil("15min")


def breakout_intact(buffer: float):
    def check(row, result) -> bool:
        if not ec.half_left(row, result):
            return False
        if row["setup"] != "Breakout" or not math.isfinite(float(row.get("atr") or float("nan"))):
            return True
        need = float(row["atr"]) * buffer
        fill = float(result["fill"])
        return fill > float(row["zone_high"]) + need if row["direction"] == "Long" else fill < float(row["zone_low"]) - need
    check.__name__ = f"breakout_intact_{buffer}"
    return check


def room_kept(row, result) -> bool:
    return result["fill_risk_pct"] >= ROOM_KEPT * float(row["stop_pct"])


def not_pretested(intraday: dict):
    def check(row, result) -> bool:
        if not ec.half_left(row, result):
            return False
        frame = intraday.get(row["ticker"])
        if frame is None:
            return True
        start = fill_candle(row)
        earlier = frame[(frame.index < start) & (frame.index.strftime("%Y-%m-%d") == row["day"])]
        if earlier.empty:
            return True
        stop = float(row["stop"])
        return float(earlier["low"].min()) > stop if row["direction"] == "Long" else float(earlier["high"].max()) < stop
    return check


# ----------------------------------------------------------------------
# 3. Exits and 4. chance: a plain bar-by-bar simulator from a known fill
# ----------------------------------------------------------------------

def run_exit(bars: pd.DataFrame, side: str, fill: float, stop: float | None, target: float | None,
             breakeven_after: float | None = None, max_bars: int | None = None) -> tuple[float, str]:
    """Exit price and reason for one trade. bars: the day's candles from the fill candle on. Stop counted first
    when a candle reaches both; a gap through a level fills at the open."""
    long = side == "Long"
    current_stop = stop
    for n, (_, bar) in enumerate(bars.iterrows()):
        o, h, l, c = (float(bar[k]) for k in ("open", "high", "low", "close"))
        if current_stop is not None and (l <= current_stop if long else h >= current_stop):
            return (min(o, current_stop) if long else max(o, current_stop)), "stop"
        if target is not None and (h >= target if long else l <= target):
            return (max(o, target) if long else min(o, target)), "target"
        if breakeven_after is not None and current_stop is not None:
            moved = (h - fill) if long else (fill - l)
            if moved >= breakeven_after:
                current_stop = fill if (long and fill > current_stop) or (not long and fill < current_stop) else current_stop
        if max_bars is not None and n + 1 >= max_bars:
            return c, "time limit"
    return (float(bars["close"].iloc[-1]), "16:30") if len(bars) else (fill, "no candles")


def net_gbp(side: str, fill: float, exit_price: float) -> float:
    move = (exit_price - fill) / fill if side == "Long" else (fill - exit_price) / fill
    return move * STAKE - STAKE * COST_PCT / 100


def day_bars(frame: pd.DataFrame, start: pd.Timestamp) -> pd.DataFrame:
    day = start.strftime("%Y-%m-%d")
    return frame[(frame.index >= start) & (frame.index.strftime("%Y-%m-%d") == day)]


EXITS = ["Current: stop and target", "Half the target", "Hold to 16:30 (stop kept)",
         "Stop to break-even after a gain of one stop distance", "2-hour time limit", "Direction only: open to 16:30"]


def exit_results(trade: dict, frame: pd.DataFrame) -> dict:
    start = fill_candle(trade)
    bars = day_bars(frame, start)
    side, fill, stop, target = trade["direction"], float(trade["fill"]), float(trade["stop"]), float(trade["target"])
    risk = abs(fill - stop)
    half = fill + (target - fill) / 2
    plans = {EXITS[0]: (stop, target, None, None), EXITS[1]: (stop, half, None, None), EXITS[2]: (stop, None, None, None),
             EXITS[3]: (stop, target, risk, None), EXITS[4]: (stop, target, None, 8), EXITS[5]: (None, None, None, None)}
    return {name: net_gbp(side, fill, run_exit(bars, side, fill, s, t, b, m)[0]) for name, (s, t, b, m) in plans.items()}


def mirrored(trade: dict) -> tuple[float, float]:
    """Stop and target for the same distances in the opposite direction."""
    fill, stop, target = float(trade["fill"]), float(trade["stop"]), float(trade["target"])
    return fill + (fill - stop), fill - (target - fill)


def chance_inputs(trade: dict, frame: pd.DataFrame) -> dict:
    """Result of the trade in the opposite direction, and at every possible fill time that day (same direction,
    same stop and target distances as % of the price)."""
    start = fill_candle(trade)
    side = trade["direction"]
    other = "Short" if side == "Long" else "Long"
    fill = float(trade["fill"])
    o_stop, o_target = mirrored(trade)
    bars = day_bars(frame, start)
    opposite = net_gbp(other, fill, run_exit(bars, other, fill, o_stop, o_target)[0])
    stop_frac = (float(trade["stop"]) - fill) / fill
    target_frac = (float(trade["target"]) - fill) / fill
    day = frame[frame.index.strftime("%Y-%m-%d") == trade["day"]]
    times = []
    for stamp in day.index:
        if not ("08:45" <= stamp.strftime("%H:%M") <= "15:30"):
            continue
        rest = day[day.index >= stamp]
        f = float(rest["open"].iloc[0])
        times.append(net_gbp(side, f, run_exit(rest, side, f, f * (1 + stop_frac), f * (1 + target_frac))[0]))
    return {"opposite": opposite, "random_times": times}


def chance_test(actual: np.ndarray, opposite: np.ndarray, times: list[list[float]], runs: int = RANDOM_RUNS,
                seed: int = SEED) -> dict:
    rng = np.random.default_rng(seed)
    flips = rng.random((runs, len(actual))) < 0.5
    direction_totals = np.where(flips, opposite, actual).sum(axis=1)
    usable = [t for t in times if t]
    time_totals = np.array([sum(t[rng.integers(len(t))] for t in usable) for _ in range(runs)]) if usable else np.array([])
    total = float(actual.sum())
    return {"scanner_total": total, "trades": len(actual),
            "random_direction_average_total": float(direction_totals.mean()),
            "random_direction_as_good": float((direction_totals >= total).mean()),
            "opposite_total": float(opposite.sum()),
            "random_time_average_total": float(time_totals.mean()) if len(time_totals) else None,
            "random_time_as_good": float((time_totals >= total).mean()) if len(time_totals) else None}


# ----------------------------------------------------------------------
# 2. Self-search: find on the first half, confirm on the second
# ----------------------------------------------------------------------

NUMERIC_FEATURES = ["fill_room_pct", "target_vs_daily_range", "fill_zone_margin_pct", "gap_toward_trade_pct",
                    "fill_vs_ema9_toward_trade_pct", "signal_volume_ratio", "zone_age_minutes", "zone_touches",
                    "score", "planned_stop_pct", "hour", "trends_agreeing", "flip_count_60m"]
CATEGORY_FEATURES = ["setup", "direction", "stop_pretested", "signal_intact", "regime", "ema_state",
                     "zone_formed_today", "weekday", "fill_room_band"]
LABELS = {"fill_room_pct": "stop room left at fill (%)", "target_vs_daily_range": "target / typical day",
          "fill_zone_margin_pct": "fill past the zone (%)", "gap_toward_trade_pct": "opening gap in trade's direction (%)",
          "fill_vs_ema9_toward_trade_pct": "fill v 9 EMA in trade's direction (%)", "signal_volume_ratio": "signal volume v normal",
          "zone_age_minutes": "zone age (trading minutes)", "zone_touches": "zone touches", "score": "quality score",
          "planned_stop_pct": "planned stop (%)", "hour": "hour of entry", "trends_agreeing": "trends agreeing (0-3)",
          "flip_count_60m": "flips in the hour before", "setup": "setup", "direction": "direction",
          "stop_pretested": "stop already traded through", "signal_intact": "signal still intact at fill",
          "regime": "regime", "ema_state": "9 EMA v 20 SMA", "zone_formed_today": "zone formed today",
          "weekday": "weekday", "fill_room_band": "stop room band"}


def groups_for(table: pd.DataFrame, first: pd.Series) -> list[tuple[str, str, pd.Series]]:
    """(feature, group label, mask) for every feature; numeric features cut into thirds at FIRST-half quantiles."""
    out_groups = []
    for feature in NUMERIC_FEATURES:
        if feature not in table.columns:
            continue
        values = pd.to_numeric(table[feature], errors="coerce")
        known = values[first & values.notna()]
        if known.nunique() < 3:
            continue
        low, high = known.quantile([1 / 3, 2 / 3])
        if low == high:
            continue
        for label, mask in ((f"under {low:.3g}", values < low), (f"{low:.3g} to {high:.3g}", (values >= low) & (values <= high)),
                            (f"over {high:.3g}", values > high)):
            out_groups.append((feature, label, mask.fillna(False)))
    for feature in CATEGORY_FEATURES:
        values = table.get(feature)
        if values is None:
            continue
        for value in sorted(values.dropna().astype(str).unique()):
            if value:
                out_groups.append((feature, value, values.astype(str) == value))
    return out_groups


def random_subset_share(values: np.ndarray, size: int, target: float, runs: int = RANDOM_RUNS, seed: int = SEED) -> float:
    """Share of random subsets of the same size whose average is at least `target` (how often luck does as well)."""
    if size <= 0 or size > len(values):
        return 1.0
    rng = np.random.default_rng(seed)
    averages = np.array([values[rng.choice(len(values), size, replace=False)].mean() for _ in range(runs)])
    return float((averages >= target - 1e-12).mean())


def self_search(table: pd.DataFrame, split: str) -> pd.DataFrame:
    """'Skip this group' and 'only this group' rules, ranked by how much they raise the AVERAGE per trade on the
    first half. (Ranking by total pounds would reward simply trading less while the average trade loses.)
    The best few are checked on the second half: the kept trades' average must beat the second half's average
    by more than random subsets of the same size do in all but LUCK_LIMIT of cases (5% shared across the candidates)."""
    net = table["net_gbp"].astype(float)
    first, second = table["day"] < split, table["day"] >= split
    base_first, base_second = float(net[first].mean()), float(net[second].mean())
    rows = []
    for feature, label, mask in groups_for(table, first):
        if (mask & first).sum() < MIN_GROUP_FIRST_HALF or (first & ~mask).sum() < MIN_GROUP_FIRST_HALF:
            continue
        for action in ("Skip", "Only"):
            kept = ~mask if action == "Skip" else mask
            avg_first = float(net[kept & first].mean())
            avg_second = float(net[kept & second].mean()) if (kept & second).any() else None
            rows.append({"Rule": f"{action} trades where {LABELS.get(feature, feature)} is {label}", "Feature": feature,
                         "Group": label, "Action": action,
                         "Trades kept (first half)": int((kept & first).sum()),
                         "Trades kept (second half)": int((kept & second).sum()),
                         "Average kept, first half (£)": avg_first, "All trades, first half (£)": base_first,
                         "Average kept, second half (£)": avg_second, "All trades, second half (£)": base_second,
                         "Improvement, first half (£ per trade)": avg_first - base_first,
                         "Improvement, second half (£ per trade)": None if avg_second is None else avg_second - base_second})
    found = pd.DataFrame(rows)
    if found.empty:
        return found
    found = found.sort_values("Improvement, first half (£ per trade)", ascending=False).reset_index(drop=True)
    top = found[found["Improvement, first half (£ per trade)"] > 0].drop_duplicates("Feature").head(TOP_CANDIDATES).index
    found["Candidate"] = found.index.isin(top)
    second_values = net[second].to_numpy()
    lucky, verdicts = [], []
    for _, row in found.iterrows():
        if not row["Candidate"]:
            lucky.append(None)
            verdicts.append("")
            continue
        size = int(row["Trades kept (second half)"])
        share = random_subset_share(second_values, size, row["Average kept, second half (£)"]) if size else 1.0
        lucky.append(share)
        if size < MIN_GROUP_SECOND_HALF or row["Trades kept (first half)"] + size < MIN_GROUP_TOTAL:
            verdicts.append("Not confirmed: too few trades")
        elif (row["Improvement, second half (£ per trade)"] or 0) <= 0:
            verdicts.append("Not confirmed: did not help on the second half")
        elif share >= LUCK_LIMIT:
            verdicts.append("Not confirmed: random trades of the same number did as well too often")
        else:
            profit = " and the kept trades made money" if row["Average kept, second half (£)"] > 0 else \
                ", but the kept trades still lost money"
            verdicts.append(f"Confirmed on unseen half{profit}")
    found["Random subsets as good (second half)"] = lucky
    found["Check on second half"] = verdicts
    return found


# ----------------------------------------------------------------------
# Summaries and the verdict
# ----------------------------------------------------------------------

def rule_summary(name: str, trades: pd.DataFrame, split: str, days: int) -> dict:
    done = ec.money(trades) if not trades.empty else trades
    if done is None or done.empty:
        return {"Rules": name, "Trades": 0}
    first, second = done[done["day"] < split], done[done["day"] >= split]
    return {"Rules": name, "Trades": len(done), "Trades a day": len(done) / max(1, days),
            "Skipped at the fill": int((trades["status"] == "skipped").sum()),
            "Won after costs": (done["net_gbp"] > 0).mean(), "Stop hit": (done["outcome"] == "stop").mean(),
            "Average per £100 trade (£)": done["net_gbp"].mean(), "Total (£)": done["net_gbp"].sum(),
            "First half: average (£)": first["net_gbp"].mean() if len(first) else None,
            "Second half: average (£)": second["net_gbp"].mean() if len(second) else None,
            "First half: total (£)": first["net_gbp"].sum(), "Second half: total (£)": second["net_gbp"].sum()}


def rule_verdicts(table: pd.DataFrame) -> list[str]:
    base = table.iloc[0]
    lines = []
    for _, row in table.iloc[1:].iterrows():
        if not row.get("Trades"):
            continue
        better = [row.get(f"{half}: average (£)", 0) is not None and base.get(f"{half}: average (£)") is not None
                  and row[f"{half}: average (£)"] > base[f"{half}: average (£)"] for half in ("First half", "Second half")]
        state = "helped in both halves" if all(better) else "helped in one half only" if any(better) else "did not help"
        money = "makes money" if row["Average per £100 trade (£)"] > 0 else "still loses money"
        lines.append(f"{row['Rules']}: {state}, {money} (average per trade £{row['Average per £100 trade (£)']:+.2f} "
                     f"v £{base['Average per £100 trade (£)']:+.2f}; {int(row['Trades'])} trades)")
    return lines


def feature_table(trades: pd.DataFrame, intraday: dict, daily: dict, closes4: dict, ideas: pd.DataFrame,
                  data_dir: Path) -> pd.DataFrame:
    """The trade-notes checklist for every completed trade of the live replay."""
    done = trades[trades["status"] == "done"].copy()
    dailies = {t: tn.daily_by_date(d) for t, d in daily.items()}
    by_day = {day: part for day, part in ideas.groupby("day")}
    rows = []
    for _, t in done.iterrows():
        trade = t.to_dict()
        side = 1 if trade["direction"] == "Long" else -1
        fill = float(trade["fill"])
        trade["exit_price"] = fill + side * float(trade["r_gross"]) * abs(fill - float(trade["stop"]))
        trade["key"] = f"{trade['run_time']}|{trade['ticker']}|{trade['direction']}"
        trade["signal_time"] = trade["run_time"]
        first = trade.get("first_touch")
        zone = (float(trade["zone_low"]), float(trade["zone_high"]),
                None if first is None or pd.isna(first) else pd.Timestamp(first), trade.get("zone_touches"))
        note = tn.note_trade(trade, trade, intraday[trade["ticker"]], dailies.get(trade["ticker"], tn.daily_by_date(None)),
                             closes4.get(trade["ticker"], pd.Series(dtype=float)), by_day.get(trade["day"], pd.DataFrame()),
                             load_views(data_dir, trade["day"]), zone_info=zone)
        if note is None:
            continue
        note["hour"] = int(note["fill_time"][:2])
        note["weekday"] = pd.Timestamp(trade["day"]).day_name()
        note["gap_toward_trade_pct"] = None if note["gap_pct"] is None else side * note["gap_pct"]
        note["fill_vs_ema9_toward_trade_pct"] = None if note["fill_vs_ema9_pct"] is None else side * note["fill_vs_ema9_pct"]
        note["run_time"], note["data_time"] = trade["run_time"], trade["data_time"]
        rows.append(note)
    return pd.DataFrame(rows)


def write_sheet(wb: Workbook, title: str, frame: pd.DataFrame, notes: list[str], money: list[str] | None = None,
                pct: list[str] | None = None, first_width: int = 60) -> None:
    ws = wb.create_sheet(title)
    for n, text in enumerate(notes, 1):
        bt.note(ws, n, text)
    if frame is not None and not frame.empty:
        bt.write_table(ws, len(notes) + 2, frame.round(4), money or [c for c in frame.columns if "(£)" in c], pct or [])
    ws.column_dimensions["A"].width = first_width
    for col in range(2, (len(frame.columns) if frame is not None else 1) + 1):
        ws.column_dimensions[bt.get_column_letter(col)].width = 15


def run(args) -> int:
    universe = [c for c in sc.load_universe(Path(args.universe)) if c.kind != "trust" and c.index == "FTSE 100"]
    if args.limit:
        universe = universe[:args.limit]
    tickers = [c.yahoo_ticker for c in universe]
    print(f"Downloading 15-minute, hourly and daily candles for {len(tickers)} FTSE 100 shares...")
    intraday = {t: sc.drop_untraded_bars(sc.completed_candles(f, "15m")) for t, f in sc.fetch_many(tickers, "15m", "60d").items()}
    intraday = {t: f for t, f in intraday.items() if len(f) > 100}
    hourly = {t: sc.completed_candles(f, "1h") for t, f in sc.fetch_many(tickers, "1h", "1y").items()}
    daily_raw = sc.fetch_many(tickers, "1d", "1y")
    daily = {t: tn.daily_by_date(f) for t, f in daily_raw.items()}
    closes4 = {t: tn.four_hour_closes(h) for t, h in hourly.items()}
    ideas = replay(universe, intraday, 15, args.warmup_days, args.workers)
    if ideas.empty:
        print("No ideas passed the live rules.")
        return 1
    print(f"{len(ideas)} ideas; adding regime and zone age...")
    ideas = add_signal_features(ideas, intraday, daily, closes4)
    days = sorted(ideas["day"].unique())
    split = days[len(days) // 2]
    now = pd.Timestamp.now(tz=sc.UK_TZ)

    # 1. pre-set rules
    plans = [("A. Live rules (fill check from 28 Oct)", ideas, ec.half_left)]
    plans += [(f"B{n}. + breakout still past the zone by {b:g} typical candle ranges", ideas, breakout_intact(b))
              for n, b in enumerate(BUFFERS)]
    plans += [("C. + at least 75% of the planned stop room left", ideas,
               lambda row, result: ec.half_left(row, result) and room_kept(row, result)),
              ("D. + stop not already traded through that day", ideas, not_pretested(intraday)),
              ("E. + no trades in a sideways regime", ideas[ideas["regime"] != "Sideways"], ec.half_left),
              (f"F. + zone at least {ZONE_MIN_MINUTES // 60} trading hours old",
               ideas[pd.to_numeric(ideas["zone_age_minutes"], errors="coerce") >= ZONE_MIN_MINUTES], ec.half_left)]
    rule_rows, base_trades = [], None
    for name, chosen, check in plans:
        trades = ec.live_trader(chosen, intraday, now, check) if not chosen.empty else pd.DataFrame()
        if base_trades is None:
            base_trades = trades
        rule_rows.append(rule_summary(name, trades, split, len(days)))
        print(f"  scored: {name}")
    rules = pd.DataFrame(rule_rows)

    # features of the base trades, self-search
    table = feature_table(base_trades, intraday, daily, closes4, ideas, Path(args.data_dir))
    search = self_search(table, split) if not table.empty else pd.DataFrame()

    # 3. exits and 4. chance
    exit_rows, actual, opposite, times = [], [], [], []
    done = table.to_dict("records")
    for trade in done:
        frame = intraday[trade["ticker"]]
        results = exit_results(trade, frame)
        exit_rows.append({"day": trade["day"], **results})
        chance = chance_inputs(trade, frame)
        actual.append(results[EXITS[0]])
        opposite.append(chance["opposite"])
        times.append(chance["random_times"])
    exits = pd.DataFrame(exit_rows)
    exit_table = pd.DataFrame([{"Exit": name, "Trades": len(exits), "Total (£)": exits[name].sum(),
                                "Average per £100 trade (£)": exits[name].mean(),
                                "First half: total (£)": exits.loc[exits["day"] < split, name].sum(),
                                "Second half: total (£)": exits.loc[exits["day"] >= split, name].sum(),
                                "Won after costs": (exits[name] > 0).mean()} for name in EXITS]) if len(exits) else pd.DataFrame()
    chance = chance_test(np.array(actual), np.array(opposite), times) if actual else {}

    # verdict
    verdict = [f"FTSE 100 shares, {days[0]} to {days[-1]} ({len(days)} trading days), {len(table)} completed trades "
               f"under the live rules; split at {split} into an earlier and a later half."]
    if chance:
        verdict.append(f"Chance test: the scanner made £{chance['scanner_total']:+.2f}. The same trades in a random direction "
                       f"averaged £{chance['random_direction_average_total']:+.2f}; {chance['random_direction_as_good']:.0%} of random "
                       f"versions did as well. At a random time the same day they averaged "
                       f"£{(chance['random_time_average_total'] or 0):+.2f}; {(chance['random_time_as_good'] or 0):.0%} did as well.")
        time_share = chance["random_time_as_good"]
        beats = chance["random_direction_as_good"] < 0.05 and time_share is not None and time_share < 0.05
        costs = chance["trades"] * STAKE * COST_PCT / 100
        verdict.append(f"Before costs (£{costs:.2f} in all) the scanner made £{chance['scanner_total'] + costs:+.2f}; "
                       f"trading every signal the opposite way would have made £{chance['opposite_total']:+.2f} after costs.")
        verdict.append("The scanner's choice of direction and timing does better than chance." if beats else
                       "The scanner does not clearly beat chance: filters cannot create an edge that is not there, "
                       "so the improvement has to come from the entry signal itself or from the exits.")
    verdict += rule_verdicts(rules)
    confirmed = search[search["Check on second half"].str.startswith("Confirmed")] if not search.empty else search
    if len(confirmed):
        for _, row in confirmed.iterrows():
            verdict.append(f"Self-search found and confirmed: {row['Rule']} (average per trade £{row['Average kept, first half (£)']:+.2f} "
                           f"v £{row['All trades, first half (£)']:+.2f} on the first half; £{row['Average kept, second half (£)']:+.2f} "
                           f"v £{row['All trades, second half (£)']:+.2f} on the unseen second half; "
                           f"{int(row['Trades kept (first half)'] + row['Trades kept (second half)'])} trades). {row['Check on second half']}.")
    else:
        verdict.append("Self-search: no candidate held up on the unseen second half.")
    if not exit_table.empty:
        best = exit_table.sort_values("Total (£)", ascending=False).iloc[0]
        current = exit_table.iloc[0]
        both = best["First half: total (£)"] > current["First half: total (£)"] and \
            best["Second half: total (£)"] > current["Second half: total (£)"]
        verdict.append(f"Exits: best was '{best['Exit']}' with £{best['Total (£)']:+.2f} v £{current['Total (£)']:+.2f} for the "
                       f"current exit ({'better in both halves' if both else 'not better in both halves'}).")
    verdict.append("All results are paper research on about 60 days of one market period. Nothing changes live trading "
                   "without Phil's agreement.")

    # workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Verdict"
    ws.cell(1, 1, "Trade feasibility backtest: what helps, and can it find its own improvement?").font = bt.Font(name="Arial", bold=True, size=13)
    for n, line in enumerate(verdict, 3):
        bt.note(ws, n, line)
    ws.column_dimensions["A"].width = 160
    write_sheet(wb, "Rules", rules, ["Pre-set rules, each replayed through the live trader (5 open, 10 a day, £100 a trade). "
                                     "A rule 'helps' only if the average per trade is better in BOTH halves."],
                pct=["Won after costs", "Stop hit"])
    shown_search = search[search["Candidate"]] if not search.empty else search
    write_sheet(wb, "Self-search", shown_search,
                ["Found on the first half only; checked on the second half, which the search never saw.",
                 f"A candidate counts only if it helps on the second half with at least {MIN_GROUP_SECOND_HALF} trades kept "
                 f"there and {MIN_GROUP_TOTAL} in all."])
    write_sheet(wb, "All groups", search, ["Every group tried by the self-search, best first-half gain first."], first_width=80)
    write_sheet(wb, "Exits", exit_table, ["The same entries with different exits, set in advance. Uses only each trade's own "
                                          "day's candles after the fill; stop counted first when a candle reaches both."],
                pct=["Won after costs"], first_width=55)
    chance_frame = pd.DataFrame([{"Measure": k, "Value": v} for k, v in chance.items()])
    write_sheet(wb, "Chance", chance_frame, ["Random direction: each trade flipped to the other direction at random (same "
                                             "distances). Random time: same share, day, direction and distances, a random "
                                             f"entry candle between 08:45 and 15:30. {RANDOM_RUNS} random versions."],
                money=[], first_width=40)
    if not table.empty:
        cols = [c for c in ["day", "fill_time", "company", "setup", "direction", "result", "net_gbp", "fill_room_pct",
                            "stop_pretested", "target_vs_daily_range", "signal_intact", "fill_zone_margin_pct", "regime",
                            "flip_count_60m", "zone_age_minutes", "zone_formed_today", "gap_pct", "news_day", "mfe_pct",
                            "mae_pct", "close_outcome_gbp", "data_quality_flag"] if c in table.columns]
        write_sheet(wb, "Trades", table[cols], [], money=["net_gbp", "close_outcome_gbp"], first_width=12)
    saved = sc.save_workbook(wb, Path(args.out))
    summary_json = {"generated": datetime.now(sc.UK_TZ).strftime("%Y-%m-%d %H:%M"), "first_day": days[0], "last_day": days[-1],
                    "days": len(days), "split": split, "verdict": verdict,
                    "rules": rules.replace({np.nan: None}).to_dict("records"),
                    "self_search": shown_search.replace({np.nan: None}).to_dict("records") if not shown_search.empty else [],
                    "exits": exit_table.replace({np.nan: None}).to_dict("records") if not exit_table.empty else [],
                    "chance": chance,
                    "trades": table.replace({np.nan: None})[[c for c in ("day", "net_gbp", "setup", "direction", "regime",
                                                                         "stop_pretested", "fill_room_band", "signal_intact")
                                                             if c in table.columns]].to_dict("records") if not table.empty else []}
    if not table.empty:
        table.to_csv(Path(args.out).with_name(Path(args.out).stem + "_trades.csv"), index=False)
    json_path = Path(args.out).with_suffix(".json")
    json_path.write_text(json.dumps(summary_json, indent=1, default=str))
    for line in verdict:
        print(line)
    print(f"Workbook saved: {saved}; summary: {json_path}")
    return 0


def main() -> None:
    parser = argparse.ArgumentParser(description="Trade feasibility backtest with self-search")
    parser.add_argument("--universe", default=str(sc.UNIVERSE_PATH))
    parser.add_argument("--data-dir", default=str(sc.DEFAULT_DATA_DIR))
    parser.add_argument("--warmup-days", type=int, default=12)
    parser.add_argument("--limit", type=int, default=None)
    parser.add_argument("--workers", type=int, default=os.cpu_count() or 2)
    parser.add_argument("--out", default="results/ftse_feasibility.xlsx")
    sys.exit(run(parser.parse_args()))


if __name__ == "__main__":
    main()
