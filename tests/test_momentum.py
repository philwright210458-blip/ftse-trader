import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest import mock

import numpy as np
import openpyxl
import pandas as pd

import ftse_momentum as mm

DATES = pd.bdate_range("2009-01-01", "2014-12-31")


def prices(drifts, seed=0, noise=0.01):
    rng = np.random.default_rng(seed)
    data = {}
    for number, drift in enumerate(drifts):
        steps = drift / 252 + rng.normal(0, noise, len(DATES))
        data[f"T{number}.L"] = 100 * np.exp(np.cumsum(steps))
    close = pd.DataFrame(data, index=DATES)
    return close.copy(), close     # opens = closes keeps the arithmetic easy to check


def rising_market():
    return pd.Series(np.linspace(5000, 7000, len(DATES)), index=DATES)


class MomentumTests(unittest.TestCase):
    def test_picks_the_strongest_shares(self):
        open_, close = prices([0.4] * 5 + [-0.1] * 15)
        result = mm.run_model(open_, close, rising_market(), mm.Settings(top=5, keep=10))
        self.assertEqual(set(result.holdings), {f"T{n}.L" for n in range(5)})
        self.assertGreater(result.values.iloc[-1], 1000)

    def test_market_filter_moves_to_cash(self):
        open_, close = prices([0.4] * 5 + [0.0] * 5)
        market = pd.Series(np.r_[np.linspace(5000, 7000, 900), np.linspace(7000, 4000, len(DATES) - 900)], index=DATES)
        result = mm.run_model(open_, close, market, mm.Settings(top=5, keep=10))
        self.assertEqual(result.holdings, {})
        self.assertGreater(result.months_in_cash, 3)
        self.assertTrue(any("Market filter" in t["Why sold"] for t in result.trades))

    def test_decisions_do_not_use_future_prices(self):
        open_, close = prices([0.3, 0.2, 0.1, 0.0, -0.1, -0.2], seed=3)
        base = mm.run_model(open_, close, rising_market(), mm.Settings(top=2, keep=3))
        cut = pd.Timestamp("2013-01-01")
        changed_open, changed_close = open_.copy(), close.copy()
        changed_close.loc[changed_close.index > cut] *= np.linspace(1, 3, (changed_close.index > cut).sum())[:, None]
        changed_open.loc[changed_open.index > cut] = changed_close.loc[changed_close.index > cut]
        other = mm.run_model(changed_open, changed_close, rising_market(), mm.Settings(top=2, keep=3))
        before = [t for t in base.trades if pd.Timestamp(t["Sold"]) <= cut]
        self.assertEqual(before, [t for t in other.trades if pd.Timestamp(t["Sold"]) <= cut])
        pd.testing.assert_series_equal(base.values.loc[:cut], other.values.loc[:cut])

    def test_costs_on_a_flat_round_trip(self):
        close = pd.DataFrame({"A.L": 100.0, "B.L": 100.0}, index=DATES)
        close.loc[: "2010-06-30", "A.L"] = np.linspace(50, 100, len(close.loc[: "2010-06-30"]))   # A rises first, then flat
        close.loc["2010-07-01":, "B.L"] = np.linspace(100, 300, len(close.loc["2010-07-01":]))
        close.loc[: "2010-06-30", "B.L"] = np.linspace(100, 60, len(close.loc[: "2010-06-30"]))
        result = mm.run_model(close.copy(), close, rising_market(), mm.Settings(top=1, keep=1, start="2010-06-01"))
        first = result.trades[0]
        self.assertEqual(first["Share"], "A.L")
        self.assertAlmostEqual(first["Per £100 after costs (£)"], (0.999 / 1.006 - 1) * 100, places=1)

    def test_clean_fixes_unit_flip_and_spike(self):
        index = pd.bdate_range("2020-01-01", periods=8)
        close = pd.DataFrame({"X.L": [1.0, 1.01, 1.02, 103.0, 104.0, 520.0, 105.0, 106.0]}, index=index)
        _, fixed = mm.clean(close.copy(), close)
        expected = [100.0, 101.0, 102.0, 103.0, 104.0, np.nan, 105.0, 106.0]
        np.testing.assert_allclose(fixed["X.L"].to_numpy(), expected)

    def test_calendar_years(self):
        values = pd.Series([100.0, 110.0, 121.0], index=pd.to_datetime(["2020-06-30", "2020-12-31", "2021-12-31"]))
        years = mm.calendar_years(values)
        self.assertAlmostEqual(years[2020], 0.10)
        self.assertAlmostEqual(years[2021], 0.10)

    def test_run_builds_workbook(self):
        open_, close = prices([0.4] * 4 + [0.0] * 8, seed=5)
        market = rising_market().rename(mm.MARKET_TICKER)
        open_[mm.MARKET_TICKER], close[mm.MARKET_TICKER] = market, market
        open_[mm.BENCHMARK_TICKER], close[mm.BENCHMARK_TICKER] = market / 10, market / 10
        lines = "\n".join(f"Co {n},C{n},T{n}.L,FTSE 250,company" for n in range(12))
        with tempfile.TemporaryDirectory() as folder:
            universe = Path(folder) / "u.csv"
            universe.write_text("company,epic,yahoo_ticker,index,type\n" + lines + "\n")
            out = Path(folder) / "m.xlsx"
            args = SimpleNamespace(universe=str(universe), exclude_trusts=True, limit=None, start="2011-01-01", out=str(out))
            with mock.patch.object(mm, "download", return_value=(open_, close)):
                self.assertEqual(mm.run(args), 0)
            book = openpyxl.load_workbook(out)
            self.assertEqual(book.sheetnames, ["Summary", "Each year", "Value over time", "Trades", "Holdings and picks now"])
            names = [book["Summary"].cell(r, 1).value for r in range(6, 10)]
            self.assertEqual(names[0], "Claude's model (with market filter)")
            self.assertIn("Every share on the list, equally (before costs)", names)


if __name__ == "__main__":
    unittest.main()
