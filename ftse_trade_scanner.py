"""FTSE 100 / FTSE 250 individual-share scanner.

Keeps the existing range/breakout rules and 3:1 reward:risk model while
counting independent zone tests, using completed intraday candles, and
separating the ranked shortlist from the full audit trail.
"""

import argparse
import csv
import sys
import time
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from ftse_sessions import is_session, session_open_close
from ftse_execution import interval_minutes, completed_candles, drop_untraded_bars, market_stamp

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

try:
    import yfinance as yf
except ImportError:
    print("Missing dependency. Run: pip install -r requirements.txt")
    sys.exit(1)


UK_TZ = ZoneInfo("Europe/London")
UNIVERSE_PATH = Path(__file__).parent / "ftse_universe.csv"
CONFLICT_SCORE_GAP = 15.0  # Long/Short scores closer than this = genuine conflict


@dataclass
class IndexConfig:
    label: str
    min_touches: int = 3
    confirmation_closes: int = 1
    zone_tolerance_pct: float = 0.0015
    swing_lookback: int = 4
    reward_risk: float = 3.0
    min_touch_gap_bars: int = 4
    rearm_move_pct: float = 0.003
    stop_buffer_pct: float = 0.001   # minimum buffer beyond the zone edge
    atr_stop_mult: float = 1.25      # stop buffer: this many typical candle ranges (median of last 14)
    min_risk_pct: float = 0.0015     # skip setups risking less than this % of price
    min_risk_atr: float = 0.75       # ...or less than this many typical candle ranges
    min_zone_width_atr: float = 0.25  # a zone counts as at least this wide when scoring precision
    min_zone_width_pct: float = 0.0003  # ...and never narrower than this share of price


INDEX_RULES = {
    "FTSE 100": IndexConfig(label="FTSE 100"),
    "FTSE 250": IndexConfig("FTSE 250", confirmation_closes=2,
                             zone_tolerance_pct=0.0025, rearm_move_pct=0.004,
                             min_risk_pct=0.0025),
}


@dataclass
class Constituent:
    company: str
    epic: str
    yahoo_ticker: str
    index: str
    kind: str = "company"  # "company" or "trust" (investment trust / fund)


@dataclass
class SwingPoint:
    ts: pd.Timestamp
    price: float
    kind: str


@dataclass
class Zone:
    low: float
    high: float
    touches: list[SwingPoint] = field(default_factory=list)

    @property
    def mid(self) -> float:
        return (self.low + self.high) / 2

    @property
    def touch_count(self) -> int:
        return len(self.touches)


@dataclass
class TradeSetup:
    company: str
    ticker: str
    index_label: str
    setup_type: str
    direction: str
    zone: Zone
    entry: float
    stop: float
    target: float
    note: str
    quality_score: float = 0.0
    conflict: bool = False
    conflict_note: str = ""
    kind: str = "company"
    last_trade: pd.Timestamp | None = None  # when the newest traded candle closed
    stale: bool = False
    stale_note: str = ""
    filter_note: str = ""   # why the entry rules (from 21 Oct 2026) reject this setup; empty = passes


def load_universe(path: Path = UNIVERSE_PATH) -> list[Constituent]:
    if not path.exists():
        print(f"Can't find {path}. Make sure ftse_universe.csv is in the same folder as this script.")
        sys.exit(1)
    with open(path, newline="", encoding="utf-8") as handle:
        return [Constituent(row["company"], row["epic"], row["yahoo_ticker"], row["index"],
                            (row.get("type") or "company").strip().lower())
                for row in csv.DictReader(handle)]


def fetch_data(ticker: str, interval: str, period: str) -> pd.DataFrame:
    df = yf.download(ticker, interval=interval, period=period, progress=False, auto_adjust=False)
    if df.empty:
        raise RuntimeError(f"No data returned for {ticker}")
    if isinstance(df.columns, pd.MultiIndex):
        df.columns = df.columns.get_level_values(0)
    df = df.rename(columns=str.lower)
    df.index.name = "datetime"
    return df


def fetch_many(tickers: list[str], interval: str, period: str, chunk: int = 50,
               pause: float = 1.0) -> dict[str, pd.DataFrame]:
    """Download many tickers at once (much faster than one at a time). Missing tickers are left out,
    so the caller can fall back to fetch_data for them."""
    result: dict[str, pd.DataFrame] = {}
    for first in range(0, len(tickers), chunk):
        batch = tickers[first:first + chunk]
        try:
            raw = yf.download(batch, interval=interval, period=period, progress=False, auto_adjust=False,
                              group_by="ticker", threads=True)
        except Exception as error:
            print(f"  batch download failed ({error}); those shares will be fetched one by one")
            continue
        if raw is None or raw.empty:
            continue
        for ticker in batch:
            try:
                frame = raw[ticker] if isinstance(raw.columns, pd.MultiIndex) else raw
            except KeyError:
                continue
            frame = frame.dropna(how="all")
            if frame.empty or "Close" not in frame.columns:
                continue
            frame = frame.rename(columns=str.lower)
            frame.index.name = "datetime"
            result[ticker] = frame
        if first + chunk < len(tickers):
            time.sleep(pause)
    return result


def candle_lag_minutes(df: pd.DataFrame, interval: str) -> float | None:
    """Minutes since the newest candle in the raw download *ended*.

    Negative: the newest candle is still forming, so the feed is live.
    Positive during market hours: the feed is behind by about that long."""
    minutes = interval_minutes(interval)
    if minutes is None or df.empty:
        return None
    last_start = df.index[-1]
    now = pd.Timestamp.now(tz=last_start.tz) if last_start.tzinfo else pd.Timestamp.now()
    return (now - (last_start + pd.Timedelta(minutes=minutes))).total_seconds() / 60


def market_open_now() -> bool:
    """Use actual London exchange hours, holidays and shortened sessions."""
    now = pd.Timestamp.now(tz=UK_TZ)
    if not is_session(now):
        return False
    opening, closing = session_open_close(now)
    return opening <= now < closing


def describe_data_lag(lags: list[float]) -> str:
    """One-line verdict on whether Yahoo's feed is live. Uses the median across
    tickers, because a single illiquid share can go quiet for a while."""
    if not lags or not market_open_now():
        return ""
    middle = float(pd.Series(lags).median())
    if middle < 0:
        return "Data check: newest candle still forming for most tickers, so the feed looks live"
    # A live feed that publishes only finished candles reads 0-15 min, exactly like a feed
    # that runs 15 min behind, so anything up to ~17 min proves nothing either way.
    if middle > 17:
        return (f"Data check: newest candle ended about {middle:.0f} min ago for most tickers, which is more "
                f"than one candle late, so the feed is behind. Check your broker's live price")
    return (f"Data check: newest candle ended about {middle:.0f} min ago. Inconclusive: this looks the same "
            f"for a live feed and a 15-minute-delayed one. Compare the latest price with your broker")


def typical_true_range(df: pd.DataFrame, length: int = 14) -> float:
    """Median true range of the last `length` candles.

    The median ignores one-off giants (the opening-auction candle, an
    overnight gap) that would drag an average, and with it every stop, far
    wider than the share's normal movement."""
    previous_close = df["close"].shift(1)
    true_range = pd.concat([df["high"] - df["low"],
                            (df["high"] - previous_close).abs(),
                            (df["low"] - previous_close).abs()], axis=1).max(axis=1)
    value = true_range.tail(length).median()
    return float(value) if pd.notna(value) else 0.0


def find_swing_points(df: pd.DataFrame, lookback: int) -> list[SwingPoint]:
    highs, lows = df["high"].to_numpy(), df["low"].to_numpy()
    result: list[SwingPoint] = []
    for index in range(lookback, len(df) - lookback):
        if highs[index] == highs[index - lookback:index + lookback + 1].max():
            result.append(SwingPoint(df.index[index], float(highs[index]), "high"))
        if lows[index] == lows[index - lookback:index + lookback + 1].min():
            result.append(SwingPoint(df.index[index], float(lows[index]), "low"))
    return result


def cluster_zones(points: list[SwingPoint], tolerance_pct: float) -> list[Zone]:
    """Cluster with a fixed seed band, avoiding chain-merged ranges."""
    zones: list[Zone] = []
    for point in sorted(points, key=lambda item: item.price):
        for zone in zones:
            if abs(point.price - zone.touches[0].price) <= zone.touches[0].price * tolerance_pct:
                zone.low, zone.high = min(zone.low, point.price), max(zone.high, point.price)
                zone.touches.append(point)
                break
        else:
            zones.append(Zone(point.price, point.price, [point]))
    return zones


def meaningful_touches(zone: Zone, df: pd.DataFrame, cfg: IndexConfig) -> list[SwingPoint]:
    """Count only independently re-tested levels.

    A repeat must be at least four candles later and have a meaningful move
    away from the preceding support/resistance touch before it is counted.
    """
    candidates = sorted(zone.touches, key=lambda point: point.ts)
    if not candidates:
        return []
    accepted = [candidates[0]]
    for point in candidates[1:]:
        previous = accepted[-1]
        between = df.loc[(df.index > previous.ts) & (df.index <= point.ts)]
        if len(between) < cfg.min_touch_gap_bars or between.empty:
            continue
        if previous.kind == "low":
            rearmed = float(between["high"].max()) >= zone.high * (1 + cfg.rearm_move_pct)
        else:
            rearmed = float(between["low"].min()) <= zone.low * (1 - cfg.rearm_move_pct)
        if rearmed:
            accepted.append(point)
    return accepted


def qualify_zones(zones: list[Zone], df: pd.DataFrame, cfg: IndexConfig) -> list[Zone]:
    result = []
    for zone in zones:
        touches = meaningful_touches(zone, df, cfg)
        if len(touches) >= cfg.min_touches:
            prices = [point.price for point in touches]
            result.append(Zone(min(prices), max(prices), touches))
    return result


def near_zone_boundary(price: float, zone: Zone, proximity_pct: float = 0.002) -> str | None:
    band = zone.mid * proximity_pct
    candidates = [("bottom", abs(price - zone.low)), ("top", abs(price - zone.high))]
    candidates = [candidate for candidate in candidates if candidate[1] <= band]
    return min(candidates, key=lambda candidate: candidate[1])[0] if candidates else None


def detect_breakout(df: pd.DataFrame, zone: Zone, confirmation_closes: int) -> str | None:
    recent = df.tail(confirmation_closes + 3)
    if len(recent) < confirmation_closes + 1:
        return None
    closes = recent["close"].to_numpy()
    if any(close < zone.low for close in closes[:-confirmation_closes]) and all(close > zone.low for close in closes[-confirmation_closes:]):
        return "Long"
    if any(close > zone.high for close in closes[:-confirmation_closes]) and all(close < zone.high for close in closes[-confirmation_closes:]):
        return "Short"
    return None


def _is_sane(direction: str, entry: float, stop: float, target: float) -> bool:
    return (direction == "Long" and stop < entry < target) or (direction == "Short" and target < entry < stop)


def quality_score(setup: TradeSetup, as_of: pd.Timestamp | None = None, min_width: float = 0.0) -> float:
    """Score 0–100: independent tests, proximity, precision, recency, setup."""
    zone = setup.zone
    touches = min(zone.touch_count, 6) * 7.0
    edge = zone.low if setup.direction == "Long" else zone.high
    proximity = max(0.0, 25.0 - abs(setup.entry - edge) / zone.mid / 0.002 * 25.0)
    width = max(zone.high - zone.low, min_width)  # a zero-width zone is one price tick, not extra precision
    precision = max(0.0, 15.0 - width / zone.mid / 0.002 * 15.0)
    setup_bonus = 10.0 if setup.setup_type == "Breakout" else 5.0
    recency = 0.0
    if as_of is not None and zone.touches:
        age_hours = max(0.0, (as_of - max(point.ts for point in zone.touches)).total_seconds() / 3600)
        recency = max(0.0, 8.0 - age_hours / 24)
    return round(min(100.0, touches + proximity + precision + setup_bonus + recency), 1)


def build_setups(company: str, ticker: str, label: str, df: pd.DataFrame,
                 zones: list[Zone], cfg: IndexConfig) -> list[TradeSetup]:
    setups: list[TradeSetup] = []
    price = float(df["close"].iloc[-1])
    typical = typical_true_range(df)
    min_risk = max(cfg.min_risk_pct * price, cfg.min_risk_atr * typical)
    min_width = max(cfg.min_zone_width_atr * typical, cfg.min_zone_width_pct * price)
    for zone in zones:
        buffer = max(cfg.stop_buffer_pct * zone.mid, cfg.atr_stop_mult * typical)
        long_stop, short_stop = zone.low - buffer, zone.high + buffer
        candidates: list[tuple[str, str, str, float]] = []
        edge = near_zone_boundary(price, zone)
        if edge == "bottom":
            candidates.append(("Range", "Long", "Price at zone bottom", long_stop))
        elif edge == "top":
            candidates.append(("Range", "Short", "Price at zone top", short_stop))
        direction = detect_breakout(df, zone, cfg.confirmation_closes)
        if direction == "Long":
            candidates.append(("Breakout", "Long", "Reclaimed above zone low after breakdown", long_stop))
        elif direction == "Short":
            candidates.append(("Breakout", "Short", "Reclaimed below zone high after breakout", short_stop))
        for setup_type, side, text, stop in candidates:
            risk = abs(price - stop)
            target = price + cfg.reward_risk * risk if side == "Long" else price - cfg.reward_risk * risk
            if risk < min_risk or not _is_sane(side, price, stop, target):
                continue
            setup = TradeSetup(company, ticker, label, setup_type, side, zone, price, stop, target,
                               f"{text} (meaningful touches={zone.touch_count})")
            setup.quality_score = quality_score(setup, df.index[-1], min_width)
            setups.append(setup)
    return setups


def dedupe_setups(setups: list[TradeSetup]) -> list[TradeSetup]:
    """Best setup per company and direction; one that passes the entry rules beats one that doesn't."""
    best: dict[tuple[str, str], TradeSetup] = {}

    def rank(item: TradeSetup):
        return (not item.conflict, item.quality_score, -(item.zone.high - item.zone.low))

    for setup in setups:
        key = (setup.company, setup.direction)
        current = best.get(key)
        if current is None or rank(setup) > rank(current):
            best[key] = setup
    return list(best.values())


# ----------------------------------------------------------------------
# Entry rules (switched on 21 Oct 2026, after backtests on 20 Oct showed them neutral)
# ATR here = typical_true_range: the median true range of the last 14 candles.
# ----------------------------------------------------------------------

RECLAIM_DISPLACEMENT_ATR = 0.5     # reclaim: close at least this far back across the zone edge
REJECTION_MIN_RANGE_ATR = 0.75     # range: the signal candle's high-low range is at least this
REJECTION_WICK_SHARE = 0.4         # range: wick on the level's side is at least this share of the candle
REJECTION_CLOSE_LOCATION = 0.6     # range: close in the top 40% of the candle for buys (bottom 40% for sells)
# Added 22 Oct 2026 (Phil: only trade at a trigger point that could make a profit):
MIN_NET_REWARD_RISK = 1.5          # target's gain after costs must be at least 1.5x the stop's loss after costs
COSTS_PCT = {"Long": 0.60, "Short": 0.10}   # per £100: dealing 0.10% + stamp duty 0.5% on buys


# Target cap (Phil, 24 Oct 2026; ideas register 27): a same-day trade's target is brought in to the
# smaller of TARGET_CAP_PCT and the share's typical daily move (median high-to-low move of the
# previous TARGET_CAP_DAYS sessions, as % of the day's open). Applied before the entry rules, so the
# "worth it after costs" check uses the capped target.
TARGET_CAP_PCT = 2.0
TARGET_CAP_DAYS = 10


def typical_daily_move_pct(df: pd.DataFrame, days: int = TARGET_CAP_DAYS) -> float | None:
    """Median (high - low) / open, in %, over the `days` sessions before the latest candle's day."""
    if df.empty:
        return None
    by_day = df.groupby(df.index.date).agg(high=("high", "max"), low=("low", "min"), open=("open", "first"))
    previous = by_day[by_day.index < df.index[-1].date()].tail(days)
    if len(previous) < days:
        return None
    moves = (previous["high"] - previous["low"]) / previous["open"] * 100
    return float(moves.median())


def cap_target(setup: TradeSetup, df: pd.DataFrame, max_pct: float = TARGET_CAP_PCT,
               days: int = TARGET_CAP_DAYS) -> bool:
    """Bring the target in to the smaller of max_pct and the share's typical daily move. True if it moved.

    max_pct <= 0 turns the cap off. Without enough history for the typical move, max_pct alone applies."""
    if max_pct <= 0 or setup.entry <= 0:
        return False
    typical = typical_daily_move_pct(df, days)
    cap = min(max_pct, typical) if typical else max_pct
    distance = abs(setup.target - setup.entry) / setup.entry * 100
    if distance <= cap:
        return False
    old = setup.target
    setup.target = setup.entry * (1 + cap / 100) if setup.direction == "Long" else setup.entry * (1 - cap / 100)
    reason = f"typical daily move {typical:.1f}%" if typical and typical < max_pct else f"{max_pct:g}% maximum"
    setup.note = f"{setup.note} Target capped at {cap:.2f}% ({reason}; was {old:.2f}).".strip()
    return True


def entry_rule_note(setup: TradeSetup, df: pd.DataFrame) -> str:
    """Empty if the setup's signal candle passes its entry rule, otherwise the reason it fails."""
    if len(df) < 15:
        return "Not enough candles to check the entry rule."
    atr = typical_true_range(df)
    bar = df.iloc[-1]
    o, h, l, c = (float(bar[k]) for k in ("open", "high", "low", "close"))
    long = setup.direction == "Long"
    if atr <= 0:
        return "No typical candle range to measure against."
    cost_note = worth_trading_note(setup)
    if setup.setup_type == "Breakout":
        need = RECLAIM_DISPLACEMENT_ATR * atr
        ok = c >= setup.zone.low + need if long else c <= setup.zone.high - need
        if not ok:
            return f"Reclaim close less than {RECLAIM_DISPLACEMENT_ATR} ATR back across the zone edge."
        if not reclaim_is_fresh(setup, df):
            return "Reclaim happened on an earlier candle, not the latest one."
        return cost_note
    span = h - l
    if span <= 0 or span < REJECTION_MIN_RANGE_ATR * atr:
        return f"No rejection candle: candle smaller than {REJECTION_MIN_RANGE_ATR} ATR."
    if long:
        wick, location, touched = (min(o, c) - l) / span, (c - l) / span, l <= setup.zone.high
    else:
        wick, location, touched = (h - max(o, c)) / span, (h - c) / span, h >= setup.zone.low
    if not touched:
        return "No rejection candle: the candle did not reach the zone."
    if wick < REJECTION_WICK_SHARE or location < REJECTION_CLOSE_LOCATION:
        return "No rejection candle: wick too short or close not far enough from the level."
    return cost_note


def reclaim_is_fresh(setup: TradeSetup, df: pd.DataFrame) -> bool:
    """The reclaim must complete on the latest candle: the close just before the confirming close(s)
    was still on the far side of the zone edge."""
    confirm = INDEX_RULES.get(setup.index_label, INDEX_RULES["FTSE 100"]).confirmation_closes
    if len(df) < confirm + 1:
        return False
    before = float(df["close"].iloc[-confirm - 1])
    return before < setup.zone.low if setup.direction == "Long" else before > setup.zone.high


def worth_trading_note(setup: TradeSetup) -> str:
    """Empty if reaching the target would gain at least MIN_NET_REWARD_RISK times what hitting the stop
    would lose, both after costs (per £100: 60p on buys, 10p on sells)."""
    entry = setup.entry
    risk = abs(entry - setup.stop) / entry * 100
    reward = abs(setup.target - entry) / entry * 100
    cost = COSTS_PCT.get(setup.direction, 0.6)
    net_reward, net_risk = reward - cost, risk + cost
    if net_risk <= 0 or net_reward < MIN_NET_REWARD_RISK * net_risk:
        return (f"Not worth it after costs: target would make about £{net_reward:.2f} per £100 against a possible "
                f"£{net_risk:.2f} loss.")
    return ""


def apply_entry_rules(setups: list[TradeSetup]) -> int:
    """Exclude (as conflicts) setups that fail the entry rules, and every setup on a share seen both ways.

    Call after annotate_conflicts. Returns how many setups were excluded by these rules."""
    directions: dict[str, set[str]] = {}
    for setup in setups:
        directions.setdefault(setup.company, set()).add(setup.direction)
    excluded = 0
    for setup in setups:
        if len(directions[setup.company]) > 1:
            note = "Level seen both ways (a buy and a sell) in this scan; skipped."
        else:
            note = setup.filter_note
        if note:
            excluded += int(not setup.conflict)
            setup.conflict, setup.conflict_note = True, note
    return excluded


def annotate_conflicts(setups: list[TradeSetup]) -> list[TradeSetup]:
    """Flag companies that show both Long and Short ideas.

    A genuine conflict (both sides score within CONFLICT_SCORE_GAP) excludes
    both. If one side is clearly stronger, only the weaker side is excluded,
    so a weak opposite-direction zone can't veto a strong setup."""
    by_company: dict[str, list[TradeSetup]] = {}
    for setup in setups:
        setup.conflict, setup.conflict_note = False, ""
        by_company.setdefault(setup.company, []).append(setup)
    for items in by_company.values():
        best: dict[str, TradeSetup] = {}
        for setup in items:
            if setup.direction not in best or setup.quality_score > best[setup.direction].quality_score:
                best[setup.direction] = setup
        if len(best) < 2:
            continue
        long_best, short_best = best["Long"], best["Short"]
        if abs(long_best.quality_score - short_best.quality_score) < CONFLICT_SCORE_GAP:
            for setup in items:
                setup.conflict = True
                setup.conflict_note = f"Long and Short score within {CONFLICT_SCORE_GAP:.0f} points; excluded from Top Trades."
        else:
            winner = long_best if long_best.quality_score > short_best.quality_score else short_best
            for setup in items:
                if setup.direction != winner.direction:
                    setup.conflict = True
                    setup.conflict_note = f"Weaker than the opposite {winner.direction} setup; excluded from Top Trades."
    return setups


def mark_stale(setups: list[TradeSetup], newest_close: pd.Timestamp, max_minutes: float) -> int:
    """Flag setups whose newest *traded* candle is well behind the freshest candle in
    this scan. A thin share that hasn't traded for hours still shows its old last
    price as 'current', and can look like it is sitting exactly at a zone edge."""
    count = 0
    for setup in setups:
        setup.stale, setup.stale_note = False, ""
        if setup.last_trade is None:
            continue
        behind = (newest_close - setup.last_trade).total_seconds() / 60
        if behind > max_minutes:
            setup.stale = True
            setup.stale_note = (f"Stale: last traded candle closed {behind:.0f} min before the newest candle in this scan; "
                                f"excluded from Top Trades.")
            count += 1
    return count


# ----------------------------------------------------------------------
# History: one small CSV pair that every run appends to
#   scan_data/runs.csv    one row per scan (which candle it saw, which mode)
#   scan_data/setups.csv  one row per setup flagged in each scan
# The Excel workbook is rebuilt from these after every run, so there is a
# single file to open instead of one per run.
# ----------------------------------------------------------------------

SCRIPT_DIR = Path(__file__).parent
DEFAULT_DATA_DIR = SCRIPT_DIR / "scan_data"
DEFAULT_WORKBOOK = SCRIPT_DIR / "ftse_scan_latest.xlsx"

RUN_COLUMNS = ["day", "as_of", "data_time", "run_time", "mode", "scanned", "failed", "flagged"]
SETUP_COLUMNS = ["day", "as_of", "data_time", "run_time", "mode", "company", "ticker", "index",
                 "setup", "direction", "score", "entry", "stop", "target", "risk_pct",
                 "zone_low", "zone_high", "touches", "conflict"]
NUMERIC_COLUMNS = ["score", "entry", "stop", "target", "risk_pct", "zone_low", "zone_high", "touches"]


@dataclass
class RunInfo:
    data_time: pd.Timestamp   # start of the last completed candle seen (UK time)
    as_of: str                # HH:MM when that candle closed, e.g. "10:15"
    day: str                  # UK date of that candle
    run_time: str
    mode: str                 # "full" or "watch"
    scanned: int
    failed: int
    data_check: str = ""


def make_run_info(data_time: pd.Timestamp, interval: str, mode: str, scanned: int, failed: int) -> RunInfo:
    stamp = data_time.tz_convert(UK_TZ) if data_time.tzinfo else data_time.tz_localize(UK_TZ)
    closes = stamp + pd.Timedelta(minutes=interval_minutes(interval) or 0)
    return RunInfo(stamp, closes.strftime("%H:%M"), stamp.strftime("%Y-%m-%d"),
                   datetime.now(UK_TZ).strftime("%Y-%m-%d %H:%M:%S"), mode, scanned, failed)


def _append_rows(path: Path, columns: list[str], rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    new_file = not path.exists() or path.stat().st_size == 0
    with open(path, "a", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns)
        if new_file:
            writer.writeheader()
        writer.writerows(rows)


def _read_table(path: Path, columns: list[str]) -> pd.DataFrame:
    if not path.exists() or path.stat().st_size == 0:
        return pd.DataFrame(columns=columns)
    table = pd.read_csv(path, dtype=str, keep_default_na=False)
    for column in NUMERIC_COLUMNS:
        if column in table.columns:
            table[column] = pd.to_numeric(table[column], errors="coerce")
    if "conflict" in table.columns:
        table["conflict"] = table["conflict"].str.lower().eq("true")
    return table


def record_run(data_dir: Path, info: RunInfo, selected: list[TradeSetup]) -> None:
    _append_rows(data_dir / "runs.csv", RUN_COLUMNS, [{
        "day": info.day, "as_of": info.as_of, "data_time": info.data_time.isoformat(),
        "run_time": info.run_time, "mode": info.mode, "scanned": info.scanned,
        "failed": info.failed, "flagged": len(selected)}])
    rows = [{
        "day": info.day, "as_of": info.as_of, "data_time": info.data_time.isoformat(),
        "run_time": info.run_time, "mode": info.mode, "company": s.company, "ticker": s.ticker,
        "index": s.index_label, "setup": s.setup_type, "direction": s.direction,
        "score": s.quality_score, "entry": round(s.entry, 2), "stop": round(s.stop, 2),
        "target": round(s.target, 2), "risk_pct": round(abs(s.entry - s.stop) / s.entry * 100, 3),
        "zone_low": round(s.zone.low, 2), "zone_high": round(s.zone.high, 2),
        "touches": s.zone.touch_count, "conflict": s.conflict} for s in selected]
    if rows:
        _append_rows(data_dir / "setups.csv", SETUP_COLUMNS, rows)


def is_repeat_candle(runs: pd.DataFrame, info: RunInfo) -> bool:
    """True if the latest recorded scan already saw this exact candle."""
    if runs.empty:
        return False
    last = runs.iloc[-1]
    return last["day"] == info.day and last["as_of"] == info.as_of


def latest_full_watchlist(setups: pd.DataFrame, runs: pd.DataFrame, day: str, limit: int) -> list[str] | None:
    """Tickers from today's most recent full scan, best score first.
    None means there has been no full scan today."""
    full_runs = runs[(runs["day"] == day) & (runs["mode"] == "full")]
    if full_runs.empty:
        return None
    latest = sorted(full_runs["as_of"])[-1]
    rows = setups[(setups["day"] == day) & (setups["as_of"] == latest)
                  & (setups["mode"] == "full") & (~setups["conflict"])]
    rows = rows.sort_values("score", ascending=False)
    return list(dict.fromkeys(rows["ticker"]))[:limit]


def persistence(setups: pd.DataFrame, runs: pd.DataFrame, day: str) -> dict[tuple[str, str], tuple[str, str]]:
    """(ticker, direction) -> ("seen/possible", first as_of) for one day.

    'possible' counts every scan from the first sighting onward, so 4/5 means
    it was flagged in 4 of the 5 scans since it first appeared."""
    todays_scans = sorted(set(runs.loc[runs["day"] == day, "as_of"]))
    flagged = setups[(setups["day"] == day) & (~setups["conflict"])]
    result = {}
    for (ticker, direction), group in flagged.groupby(["ticker", "direction"]):
        seen = sorted(set(group["as_of"]))
        possible = [t for t in todays_scans if t >= seen[0]]
        result[(ticker, direction)] = (f"{len(seen)}/{len(possible)}", seen[0])
    return result


# ----------------------------------------------------------------------
# Workbook (rebuilt from history after every run)
# ----------------------------------------------------------------------

HEADERS = ["Rank", "Company", "Ticker", "Type", "Index", "Setup", "Direction", "Quality Score", "Conflict",
           "Entry", "Stop", "Target", "Reward:Risk", "Risk %", "Zone Low", "Zone High",
           "Meaningful Touches", "Last trade", "Note"]
TOP_HEADERS = HEADERS[:-1] + ["Seen today", "First seen (as of)", "Note"]
LONG_FILL, SHORT_FILL, CONFLICT_FILL, STALE_FILL = "C6EFCE", "FFC7CE", "FFF2CC", "E7E6E6"


def report_timestamp(now: datetime | None = None) -> str:
    stamp = now.astimezone(UK_TZ) if now else datetime.now(UK_TZ)
    return f"Scan generated {stamp.strftime('%Y-%m-%d %H:%M')} UK time"


def _style_header(ws, headers: list[str], widths: dict[int, int] | None = None) -> None:
    ws.append(headers)
    for cell in ws[ws.max_row]:
        cell.font, cell.fill = Font(name="Arial", bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center")
    for number, header in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(number)].width = (widths or {}).get(number, max(11, min(48, len(header) + 4)))
    ws.freeze_panes = "A3"  # a note row is inserted above the header later


def _fill_row(ws, fill: str) -> None:
    for cell in ws[ws.max_row]:
        cell.font, cell.fill = Font(name="Arial"), PatternFill("solid", fgColor=fill)


def _direction_fill(direction: str, conflict: bool = False) -> str:
    return CONFLICT_FILL if conflict else (LONG_FILL if direction == "Long" else SHORT_FILL)


def _kind_label(kind: str | None) -> str:
    return "Trust" if kind == "trust" else "Company"


def _clock(stamp: pd.Timestamp | None) -> str:
    if stamp is None:
        return ""
    return (stamp.tz_convert(UK_TZ) if stamp.tzinfo else stamp).strftime("%H:%M")


def _append_setup(ws, setup: TradeSetup, rank: int | str = "", extra: tuple = ()) -> None:
    rr = abs(setup.target - setup.entry) / abs(setup.entry - setup.stop) if setup.entry != setup.stop else 0
    row = [rank, setup.company, setup.ticker, _kind_label(setup.kind), setup.index_label, setup.setup_type, setup.direction,
           setup.quality_score, "Yes" if setup.conflict else "No", round(setup.entry, 2), round(setup.stop, 2),
           round(setup.target, 2), round(rr, 2), round(abs(setup.entry - setup.stop) / setup.entry * 100, 3),
           round(setup.zone.low, 2), round(setup.zone.high, 2), setup.zone.touch_count, _clock(setup.last_trade),
           f"{setup.note} {setup.conflict_note} {setup.stale_note}".strip()]
    ws.append(row[:-1] + list(extra) + row[-1:])
    _fill_row(ws, STALE_FILL if setup.stale else _direction_fill(setup.direction, setup.conflict))


def _add_note(ws, text: str) -> None:
    ws.insert_rows(1)
    ws.cell(1, 1, text)
    ws.cell(1, 1).font = Font(name="Arial", italic=True, size=9)


MOCK_HEADERS = ["#", "Signal time", "Company", "Ticker", "Setup", "Direction", "Entry (filled)", "Stop", "Target",
                "Status so far", "At", "Latest or exit price", "Before costs per £100 (£)", "Costs (£)",
                "After costs per £100 (£)"]


def _write_mock_sheet(ws, rows: list[dict], info: "RunInfo") -> None:
    """Today's live mock trades with their progress so far, so they stay visible after they leave Top Trades."""
    _style_header(ws, MOCK_HEADERS, {1: 5, 2: 11, 3: 30, 4: 10})
    money = '"£"#,##0.00;[Red]-"£"#,##0.00'
    for row in rows:
        p = row.get("progress", {})
        ws.append([int(row["trade_no"]), str(row["signal_time"])[11:16], row["company"], row["ticker"],
                   "Reclaim" if row["setup"] == "Breakout" else row["setup"],
                   "Buy" if row["direction"] == "Long" else "Sell", p.get("fill"), float(row["stop"]),
                   float(row["target"]), p.get("status", ""), p.get("time", ""), p.get("price"), p.get("gross"),
                   p.get("costs"), p.get("net")])
        for column in (13, 14, 15):
            ws.cell(ws.max_row, column).number_format = money
        status = p.get("status", "")
        colour = "D9EAD3" if status == "Target hit" else "F4CCCC" if status == "Stopped out" else None
        if colour:
            for cell in ws[ws.max_row]:
                cell.fill = PatternFill("solid", fgColor=colour)
    if not rows:
        ws.append(["", "", "No mock trades yet today."])
    _add_note(ws, f"Mock trades made today (up to 5 open at once and 10 a day, one per share, FTSE 100 shares only, only ideas that pass every rule), with progress from completed 15-minute candles up to {info.as_of} UK. "
                  "Provisional: Yahoo prices can be up to 15 minutes behind. The final result is worked out at 17:05 "
                  "in ftse_outcomes.xlsx (Live Mock Trades sheet). Costs per £100: 10p dealing, plus 50p stamp duty on buys.")


def build_workbook(info: RunInfo, all_setups: list[TradeSetup], selected: list[TradeSetup],
                   ranked: list[TradeSetup], setups_history: pd.DataFrame, runs: pd.DataFrame,
                   kinds: dict[str, str] | None = None, mock_rows: list[dict] | None = None) -> Workbook:
    """One sheet: today's live mock trades. (Top Trades, Today, All Setups and History were removed on 23 Oct 2026;
    every flagged setup is still saved in scan_data/setups.csv.)"""
    workbook = Workbook()
    mocks = workbook.active
    mocks.title = "Today's Mock Trades"
    _write_mock_sheet(mocks, mock_rows or [], info)
    if info.mode == "close":
        _add_note(mocks, f"{report_timestamp()} | final check after the market close (16:30): open trades are closed at the last price")
    else:
        _add_note(mocks, f"{report_timestamp()} | {info.mode} scan, data to {info.as_of} UK | {len(ranked)} idea(s) found this scan"
                         + (f" | {info.data_check}" if info.data_check else ""))
    mocks.freeze_panes = "A4"                          # two note rows, then the column headings
    return workbook


def todays_mock_trades(data_dir: Path, info: "RunInfo", interval: str) -> list[dict]:
    """Today's live mock trades, each with its progress so far (one small price download per trade)."""
    from ftse_live_mock import LIVE_FILENAME, mock_progress, read_live
    live = read_live(Path(data_dir) / LIVE_FILENAME)
    rows = live[live["day"] == info.day].to_dict("records") if not live.empty else []
    for row in rows:
        try:
            data = drop_untraded_bars(completed_candles(fetch_data(row["ticker"], interval, "5d"), interval))
        except Exception:
            data = None
        row["progress"] = mock_progress(row, data, interval_minutes(interval) or 15, now=market_stamp(info.run_time))
    return rows


def refresh_mock_workbook(data_dir: Path, workbook_path: Path, interval: str, now: datetime | None = None) -> int:
    """Rewrite the Today's Mock Trades sheet after the close, so trades still open at the last scan show how
    the day ended. Uses today's UK date; does nothing on a day with no scans."""
    now = now or datetime.now(UK_TZ)
    day = now.strftime("%Y-%m-%d")
    runs = _read_table(Path(data_dir) / "runs.csv", RUN_COLUMNS)
    if runs.empty or day not in set(runs["day"]):
        print(f"No scans recorded on {day}; scan workbook left as it is.")
        return 0
    info = RunInfo(pd.Timestamp(f"{day} 16:15", tz=UK_TZ), "16:30", day, now.strftime("%Y-%m-%d %H:%M:%S"), "close", 0, 0)
    rows = todays_mock_trades(Path(data_dir), info, interval)
    save_workbook(build_workbook(info, [], [], [], pd.DataFrame(), runs, None, rows), workbook_path)
    print(f"Today's Mock Trades updated to the close: {len(rows)} trade(s).")
    return 0


def save_workbook(workbook: Workbook, path: Path) -> Path:
    """Save, falling back to a new name if the file is open in Excel (Windows locks it)."""
    try:
        workbook.save(path)
        return path
    except PermissionError:
        alternative = path.with_name(f"{path.stem}_{datetime.now(UK_TZ).strftime('%H%M%S')}{path.suffix}")
        workbook.save(alternative)
        print(f"\n'{path.name}' is open in Excel, so this scan was saved as '{alternative.name}' instead.\n"
              f"Close the workbook so the next scan can update it.")
        return alternative


# ----------------------------------------------------------------------
# Scanning
# ----------------------------------------------------------------------

def run_scan(universe: list[Constituent], interval: str, period: str, delay: float, batch: bool = False,
             target_cap_pct: float = TARGET_CAP_PCT, target_cap_days: int = TARGET_CAP_DAYS):
    all_setups: list[TradeSetup] = []
    failures: list[str] = []
    latest: pd.Timestamp | None = None
    lags: list[float] = []
    prefetched = fetch_many([c.yahoo_ticker for c in universe], interval, period) if batch else {}
    if batch:
        print(f"  downloaded {len(prefetched)}/{len(universe)} shares in batches; fetching the rest one by one")
    for number, company in enumerate(universe, 1):
        try:
            cfg = INDEX_RULES[company.index]
            raw = prefetched.get(company.yahoo_ticker)
            if raw is None:
                raw = fetch_data(company.yahoo_ticker, interval, period)
                time.sleep(delay)
            lag = candle_lag_minutes(raw, interval)
            if lag is not None:
                lags.append(lag)
            data = drop_untraded_bars(completed_candles(raw, interval))
            if len(data) <= cfg.swing_lookback * 2:
                raise RuntimeError("Not enough completed candles")
            if latest is None or data.index[-1] > latest:
                latest = data.index[-1]
            zones = qualify_zones(cluster_zones(find_swing_points(data, cfg.swing_lookback), cfg.zone_tolerance_pct), data, cfg)
            setups = build_setups(company.company, company.yahoo_ticker, company.index, data, zones, cfg)
            closes = data.index[-1] + pd.Timedelta(minutes=interval_minutes(interval) or 0)
            for setup in setups:
                setup.kind, setup.last_trade = company.kind, closes
                cap_target(setup, data, target_cap_pct, target_cap_days)
                setup.filter_note = entry_rule_note(setup, data)
            all_setups.extend(setups)
            if setups:
                print(f"  [{number}/{len(universe)}] {company.company} ({company.yahoo_ticker}) -> {len(setups)} setup(s)")
        except Exception as error:
            failures.append(f"{company.company} ({company.yahoo_ticker}): {error}")
        if number % 25 == 0:
            print(f"  ...{number}/{len(universe)} scanned, {len(all_setups)} setup(s) found so far")
    return all_setups, failures, latest, lags


def main() -> None:
    parser = argparse.ArgumentParser(description="FTSE 100/250 individual share scanner")
    parser.add_argument("--interval", default="15m", help="Candle interval, e.g. 15m, 1h (default: 15m)")
    parser.add_argument("--period", default="60d", help="History window, e.g. 60d (max for 15m on Yahoo)")
    parser.add_argument("--watch", action="store_true",
                        help="Re-check only the shortlist from today's latest full scan (fast). "
                             "Records nothing if no new candle has formed since the last scan.")
    parser.add_argument("--watch-top", type=int, default=15,
                        help="How many shortlisted tickers to watch (default: 15; 0 = every share, no full scan needed)")
    parser.add_argument("--out", default=str(DEFAULT_WORKBOOK), help="Workbook to update (default: ftse_scan_latest.xlsx next to this script)")
    parser.add_argument("--data-dir", default=str(DEFAULT_DATA_DIR), help="Folder for the run history CSVs (default: scan_data next to this script)")
    parser.add_argument("--limit", type=int, default=None, help="Only scan the first N companies (full scans; for testing)")
    parser.add_argument("--index", choices=["FTSE 100", "FTSE 250"], default=None, help="Scan only one index (full scans)")
    parser.add_argument("--universe", default=str(UNIVERSE_PATH), help="Path to the constituent CSV")
    parser.add_argument("--delay", type=float, default=0.3, help="Seconds to pause between tickers (default: 0.3)")
    parser.add_argument("--max-stale", type=int, default=30,
                        help="Exclude setups whose last traded candle is more than this many minutes behind the "
                             "freshest candle in the scan; 0 turns the guard off (default: 30)")
    parser.add_argument("--exclude-trusts", action="store_true",
                        help="Leave out investment trusts and funds (the 'type' column in ftse_universe.csv)")
    parser.add_argument("--live-max-per-day", type=int, default=10,
                        help="Live mock trader: most paper trades entered in a day, in the order they trigger (default: 10; 0 turns it off)")
    parser.add_argument("--live-max-open", type=int, default=5,
                        help="Live mock trader: most trades open at the same time; a trade that stops out or hits its "
                             "target frees its slot (default: 5; 0 = no limit on open trades)")
    parser.add_argument("--live-max-per-scan", type=int, default=0,
                        help="Live mock trader: most new paper trades from any one scan (default: 0 = no limit; every idea "
                             "that passes the rules is taken while open and daily slots remain)")
    parser.add_argument("--live-index", default="FTSE 100",
                        help="Live mock trader: only trade shares in this index (default: FTSE 100, whose Yahoo data is "
                             "more complete); use '' for every share. The scan itself still covers every share")
    parser.add_argument("--one-by-one", action="store_true",
                        help="Download shares one at a time instead of in batches (slower; for troubleshooting)")
    parser.add_argument("--no-entry-rules", action="store_true",
                        help="Turn off the entry rules added on 21 Oct 2026 (seen-both-ways, rejection candle, reclaim distance)")
    parser.add_argument("--target-cap-pct", type=float, default=TARGET_CAP_PCT,
                        help="Cap each target at the smaller of this %% and the share's typical daily move "
                             "(default: 2; 0 turns the cap off)")
    parser.add_argument("--target-cap-days", type=int, default=TARGET_CAP_DAYS,
                        help="Days used for the share's typical daily move (default: 10)")
    parser.add_argument("--picker", choices=["filter", "off"], default="filter",
                        help="Use the morning picker's views (scan_data/picker/<day>.csv): skip ideas against them and "
                             "trade agreeing ideas first (default: filter; off ignores them)")
    parser.add_argument("--refresh-mock", action="store_true",
                        help="After the close: re-check today's mock trades to the end of the day and rewrite the "
                             "workbook (no scan, nothing recorded). Run by the 17:05 outcomes workflow")
    parser.add_argument("--top", type=int, default=10, help="Show top N non-conflicting setups by quality score; 0 means no cap (default: 10)")
    args = parser.parse_args()

    data_dir, workbook_path = Path(args.data_dir), Path(args.out)
    if args.refresh_mock:
        sys.exit(refresh_mock_workbook(data_dir, workbook_path, args.interval))
    runs = _read_table(data_dir / "runs.csv", RUN_COLUMNS)
    setups_history = _read_table(data_dir / "setups.csv", SETUP_COLUMNS)
    universe = load_universe(Path(args.universe))
    kinds = {company.yahoo_ticker: company.kind for company in universe}
    if args.exclude_trusts:
        if not any(company.kind == "trust" for company in universe):
            print("Note: no 'type' column with trusts in the universe file, so --exclude-trusts has no effect.")
        universe = [company for company in universe if company.kind != "trust"]
    mode = "watch" if args.watch else "full"

    if args.watch and args.watch_top <= 0:
        print(f"Watching all {len(universe)} companies...")
    elif args.watch:
        today = datetime.now(UK_TZ).strftime("%Y-%m-%d")
        tickers = latest_full_watchlist(setups_history, runs, today, args.watch_top)
        if tickers is None:
            print("No full scan found for today. Run a full scan first:  python ftse_trade_scanner.py")
            sys.exit(1)
        if not tickers:
            print("Today's full scan flagged nothing to watch.")
            return
        universe = [company for company in universe if company.yahoo_ticker in set(tickers)]
        print(f"Watching {len(universe)} shortlisted companies from today's latest full scan...")
    else:
        if args.index:
            universe = [company for company in universe if company.index == args.index]
        if args.limit:
            universe = universe[:args.limit]
        print(f"Scanning {len(universe)} companies ({'/'.join(sorted(set(company.index for company in universe)))})...")

    all_setups, failures, latest, lags = run_scan(universe, args.interval, args.period, args.delay,
                                                   batch=not args.one_by_one, target_cap_pct=args.target_cap_pct,
                                                   target_cap_days=args.target_cap_days)
    if latest is None:
        print("No usable data came back (market closed, or Yahoo unavailable). Nothing recorded.")
        sys.exit(1)

    info = make_run_info(latest, args.interval, mode, len(universe), len(failures))
    info.data_check = describe_data_lag(lags)
    if args.watch and args.watch_top <= 0 and info.day != datetime.now(UK_TZ).strftime("%Y-%m-%d"):
        print(f"\nNo candle from today yet (newest data is from {info.day}). Nothing recorded.")
        return
    if args.watch and is_repeat_candle(runs, info):
        print(f"\nNo new candle since the last scan (data still to {info.as_of}). Nothing recorded.")
        return

    stale_count = 0
    minutes = interval_minutes(args.interval)
    if minutes and args.max_stale > 0:
        stale_count = mark_stale(all_setups, info.data_time + pd.Timedelta(minutes=minutes), args.max_stale)
    fresh = [setup for setup in all_setups if not setup.stale]
    annotate_conflicts(fresh)
    rule_excluded = 0 if args.no_entry_rules else apply_entry_rules(fresh)
    picker_skipped, picker_note = 0, ""
    if args.picker != "off":
        from ftse_picker import apply_views, load_views
        views = load_views(data_dir, info.day)
        if views is None:
            picker_note = f"No morning picker file for {info.day}; scanning without it."
        else:
            picker_skipped = apply_views(fresh, views)
            picker_note = (f"Morning picker: {len(views)} share view(s); {picker_skipped} idea(s) skipped as against it.")
    selected = dedupe_setups(fresh)
    ranked = sorted((s for s in selected if not s.conflict), key=lambda s: s.quality_score, reverse=True)
    if args.top > 0:
        ranked = ranked[:args.top]

    record_run(data_dir, info, selected)
    live_rows = []
    if args.live_max_per_day > 0:
        from ftse_live_mock import record_live_mock
        live_candidates = [s for s in selected if not s.conflict
                           and (not args.live_index or s.index_label == args.live_index)]
        from ftse_live_mock import still_open
        open_now = still_open(todays_mock_trades(data_dir, info, args.interval)) if args.live_max_open > 0 else 0
        live_rows = record_live_mock(data_dir, info, live_candidates, args.live_max_per_day,
                                     max_per_scan=args.live_max_per_scan, max_open=args.live_max_open, open_now=open_now)
    runs = _read_table(data_dir / "runs.csv", RUN_COLUMNS)
    setups_history = _read_table(data_dir / "setups.csv", SETUP_COLUMNS)
    mock_rows = todays_mock_trades(data_dir, info, args.interval)
    saved = save_workbook(build_workbook(info, all_setups, selected, ranked, setups_history, runs, kinds, mock_rows),
                          workbook_path)

    print(f"\nScanned {len(universe)} companies ({mode}), data to {info.as_of} UK.")
    if stale_count:
        print(f"{stale_count} setup(s) excluded as stale (last trade more than {args.max_stale} min behind the freshest candle).")
    print(f"{len(all_setups) - stale_count} fresh setup(s) -> {len(selected)} after de-duplication; "
          f"{sum(s.conflict for s in selected)} conflicting/weaker excluded from Top Trades.")
    if rule_excluded:
        print(f"{rule_excluded} setup(s) excluded by the entry rules (seen both ways, no rejection candle, or reclaim too shallow).")
    if picker_note:
        print(picker_note)
    if info.data_check:
        print(info.data_check)
    if failures:
        print(f"{len(failures)} ticker(s) failed to fetch:")
        for failure in failures[:15]:
            print(f"  - {failure}")
        if len(failures) > 15:
            print(f"  ...and {len(failures) - 15} more")
    for row in live_rows:
        print(f"Live mock trade #{row['trade_no']} today: {row['direction']} {row['company']} ({row['ticker']}), "
              f"entry about {row['entry']}, stop {row['stop']}, target {row['target']}. Filled at the next candle's open.")
    print(f"\nWorkbook updated: {saved}")


if __name__ == "__main__":
    main()
