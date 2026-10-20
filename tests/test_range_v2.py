import unittest
from unittest import mock

import pandas as pd

import ftse_range_v2 as v2
import ftse_trade_scanner as sc
from ftse_sessions import calendar, session_open_close

COMPANY = sc.Constituent("Alpha plc", "AAA", "AAA.L", "FTSE 100", "company")


def candles(closes, start_price=None):
    """15-minute candles following a path of closes across real London sessions."""
    sessions = calendar().sessions_in_range("2026-08-03", "2026-09-30")
    stamps = []
    for day in sessions:
        opening, closing = session_open_close(day)
        stamps.extend(pd.date_range(opening, closing, freq="15min", inclusive="left"))
        if len(stamps) >= len(closes):
            break
    rows, previous = [], start_price if start_price is not None else closes[0]
    for stamp, close in zip(stamps, closes):
        rows.append((stamp, previous, max(previous, close) + 0.05, min(previous, close) - 0.05, close, 1000))
        previous = close
    return pd.DataFrame(rows, columns=["datetime", "open", "high", "low", "close", "volume"]).set_index("datetime")


FLAT = [100.0, 100.2] * 30
RISE = [100.0 + 0.5 * k for k in range(1, 9)]                    # 100.5 ... 104.0
BOX_UP = [103.7, 103.4, 103.1, 103.4, 103.7, 104.0] * 3          # 18 candles between about 103 and 104
FALL = [104.2 - 0.5 * k for k in range(1, 9)]                    # 103.7 ... 100.2
BOX_DOWN = [100.5, 100.8, 101.1, 100.8, 100.5, 100.2] * 3

ALL_DAY = mock.patch.multiple(v2, FIRST_SIGNAL="00:00", LAST_SIGNAL="23:59", QUIET_FROM="99:98", QUIET_TO="99:99")


class RangeV2Tests(unittest.TestCase):
    def check_levels(self, row):
        long = row["direction"] == "Long"
        self.assertTrue(row["stop"] < row["entry"] < row["target"] if long else row["target"] < row["entry"] < row["stop"])
        risk, reward = abs(row["entry"] - row["stop"]), abs(row["target"] - row["entry"])
        self.assertGreaterEqual(reward / risk, v2.MIN_REWARD_RISK - 1e-6)

    @ALL_DAY
    def test_confirmed_breakout_goes_with_the_break(self):
        data = candles(FLAT + RISE + BOX_UP + [104.3, 104.2, 104.6, 104.8])
        rows = v2.signals(COMPANY, data)
        breakouts = [r for r in rows if r["setup"] == "Breakout"]
        self.assertEqual(len(breakouts), 1)
        self.assertEqual(breakouts[0]["direction"], "Long")
        self.assertEqual(breakouts[0]["entry"], 104.6)             # the confirming candle, not the first break
        self.check_levels(breakouts[0])

    @ALL_DAY
    def test_unconfirmed_break_gives_no_breakout(self):
        data = candles(FLAT + RISE + BOX_UP + [104.3, 104.25, 104.2, 104.15, 104.1])
        self.assertEqual([r for r in v2.signals(COMPANY, data) if r["setup"] == "Breakout"], [])

    @ALL_DAY
    def test_reclaim_after_false_breakdown_buys_towards_start_of_move(self):
        data = candles([104.0, 104.2] * 30 + FALL + BOX_DOWN + [99.8, 100.4, 100.6])
        rows = [r for r in v2.signals(COMPANY, data) if r["setup"] == "Reclaim"]
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["direction"], "Long")
        self.assertGreater(rows[0]["target"], 103.5)                # back to where the fall began
        self.check_levels(rows[0])

    @ALL_DAY
    def test_no_box_without_big_move(self):
        sideways = [round(x - 3.6, 2) for x in BOX_UP]                 # same swings, around 100, no big move first
        data = candles(FLAT + sideways * 3 + [100.7, 100.6, 101.0])
        self.assertEqual(v2.signals(COMPANY, data), [])

    @ALL_DAY
    def test_signals_use_only_past_candles(self):
        data = candles(FLAT + RISE + BOX_UP + [103.4, 103.1, 103.4, 103.7, 104.3, 104.2, 104.6, 104.8])
        full = v2.signals(COMPANY, data)
        self.assertTrue(full)
        for row in full:
            cut = data[data.index <= pd.Timestamp(row["data_time"])]
            again = [r for r in v2.signals(COMPANY, cut) if r["data_time"] == row["data_time"]]
            self.assertEqual(again, [row])

    def test_quiet_hours_and_late_signals_dropped(self):
        data = candles(FLAT + RISE + BOX_UP + [104.3, 104.2, 104.6, 104.8])
        for row in v2.signals(COMPANY, data):
            self.assertTrue(v2.FIRST_SIGNAL <= row["as_of"] <= v2.LAST_SIGNAL)
            self.assertFalse(v2.QUIET_FROM <= row["as_of"] <= v2.QUIET_TO)

    def test_touch_count_needs_price_to_move_away(self):
        highs = [104.0, 104.0, 104.0, 103.5, 104.0]
        lows = [103.5, 103.6, 103.0, 103.0, 103.4]
        closes = [103.9, 103.9, 103.5, 103.2, 103.9]
        self.assertEqual(v2.count_touches(highs, lows, closes, 104.0, 103.0), (2, 1))


class RangeBacktestRunTests(unittest.TestCase):
    def test_run_builds_workbook_for_shares_and_index(self):
        import tempfile
        from pathlib import Path
        from types import SimpleNamespace

        import openpyxl

        import ftse_range_backtest as rb

        episode = (FLAT + RISE + BOX_UP + [103.4, 103.1, 103.4, 103.7, 104.3, 104.2, 104.6, 104.8]
                   + [105.0 + 0.1 * k for k in range(20)])
        path = []
        for k in range(8):
            path.extend([round(p + k * 0.01, 2) for p in episode])
        share = candles(path)
        index = candles(path)
        index["volume"] = 1
        with tempfile.TemporaryDirectory() as folder:
            universe = Path(folder) / "u.csv"
            universe.write_text("company,epic,yahoo_ticker,index,type\nAlpha plc,AAA,AAA.L,FTSE 100,company\n")
            out_path = Path(folder) / "v2.xlsx"
            args = SimpleNamespace(universe=str(universe), exclude_trusts=True, limit=None, interval="15m", period="60d",
                                   delay=0, warmup_days=0, horizon=3, out=str(out_path))
            with mock.patch.object(rb, "download", return_value={"AAA.L": share, rb.INDEX_TICKER: index}):
                self.assertEqual(rb.run(args), 0)
            book = openpyxl.load_workbook(out_path)
            self.assertEqual(book.sheetnames, ["Summary", "By setup", "Daily", "Trades (shares)", "Trades (index)"])
            markets = {book["Summary"].cell(r, 1).value for r in range(8, book["Summary"].max_row + 1)}
            self.assertEqual(markets, {"FTSE shares", "FTSE 100 index"})

    def test_index_costs_add_funding_per_night(self):
        import ftse_range_backtest as rb
        trades = pd.DataFrame({"status": ["done", "done"], "run_time": ["2026-10-12 10:02:00", "2026-10-12 10:02:00"],
                               "exit_time": ["2026-10-12T15:00:00+01:00", "2026-10-14T11:00:00+01:00"],
                               "gross_gbp": [1.0, 1.0]})
        costed = rb.index_costs(trades)
        self.assertEqual(list(costed["costs_gbp"]), [0.02, 0.06])
        self.assertEqual(list(costed["net_gbp"]), [0.98, 0.94])


if __name__ == "__main__":
    unittest.main()
