"""Build the FTSE Trader Today web page (GitHub Pages) from the files the scans already save.

Run by the "FTSE Trader Today page" workflow after every Watch scan and the 17:05 Outcomes run, so the
page is never more than one scan behind. It reads only committed files; it never downloads prices.

    python ftse_site.py --out site          # writes site/index.html

Sources:
  * ftse_scan_latest.xlsx, sheet "Today's Mock Trades": today's live mock trades and their progress
  * scan_data/picker/<day>.csv: the morning picker's views
  * scan_data/daily_moves/<day>.csv: each FTSE 100 share's move and the picker's scorecard (after 17:05)
  * scan_data/picker_trades.csv: the picker paper trader, every Buy/Sell view as a £100 trade
  * ftse_outcomes.xlsx, sheet "Live Mock Trades": the scanner's running total
  * scan_data/structure_live/trades.csv: the daily structure pullback model
"""

from __future__ import annotations

import argparse
import html
import math
import re
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).parent
UK = "Europe/London"


def esc(value) -> str:
    return html.escape("" if value is None or (isinstance(value, float) and math.isnan(value)) else str(value))


def money(value, signed: bool = True) -> str:
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return "–"
    value = float(value)
    sign = "+" if signed and value > 0 else "-" if value < 0 else ""
    return f"{sign}£{abs(value):,.2f}"


def tone(value) -> str:
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return ""
    return "up" if value > 0.004 else "down" if value < -0.004 else ""


# ----------------------------------------------------------------------
# Reading the saved files
# ----------------------------------------------------------------------

def scan_today(path: Path) -> dict:
    """Today's mock trades from the scan workbook, plus when the scan ran."""
    result = {"generated": "", "day": "", "status_line": "", "trades": []}
    if not path.exists():
        return result
    ws = load_workbook(path, read_only=True, data_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    first = str(rows[0][0] or "") if rows else ""
    found = re.search(r"(\d{4}-\d{2}-\d{2}) (\d{2}:\d{2})", first)
    if found:
        result["day"], result["generated"] = found.group(1), f"{found.group(1)} {found.group(2)}"
    result["status_line"] = first.split("|", 1)[1].strip() if "|" in first else ""
    header_at = next((i for i, r in enumerate(rows) if r and r[0] == "#"), None)
    if header_at is None:
        return result
    header = list(rows[header_at])
    for r in rows[header_at + 1:]:
        if not r or r[0] in (None, ""):
            continue
        result["trades"].append(dict(zip(header, r)))
    return result


def read_csv(path: Path) -> pd.DataFrame:
    if not path.exists() or path.stat().st_size == 0:
        return pd.DataFrame()
    try:
        return pd.read_csv(path, dtype=str, keep_default_na=False)
    except Exception:
        return pd.DataFrame()


def scanner_total(path: Path) -> dict | None:
    """The scanner's running total from the outcomes workbook (totals row under the first header)."""
    if not path.exists():
        return None
    ws = load_workbook(path, read_only=True, data_only=True)["Live Mock Trades"]
    rows = list(ws.iter_rows(min_row=1, max_row=12, values_only=True))
    for i, r in enumerate(rows[:-1]):
        if r and r[0] == "Net after all fees (£)":
            values = rows[i + 1]
            return {"net": values[0] if isinstance(values[0], (int, float)) else None, "completed": values[1],
                    "won": values[2], "lost": values[3], "waiting": values[5]}
    return None


def latest_moves_before(folder: Path, day: str) -> tuple[str, pd.DataFrame]:
    files = sorted(p for p in folder.glob("*.csv") if p.stem < day) if folder.exists() else []
    return (files[-1].stem, read_csv(files[-1])) if files else ("", pd.DataFrame())


# ----------------------------------------------------------------------
# Page
# ----------------------------------------------------------------------

STYLE = """
:root {
  /* Layout: one column of panels, summary first; tables scroll inside their panel on a phone. */
  --bg: #f3f5f2; --panel: #ffffff; --ink: #18211c; --muted: #5c6b62; --line: #dbe2dc;
  --accent: #1f6f5c; --up: #17803d; --down: #b4322e; --up-bg: #e3f3e6; --down-bg: #fbe7e5; --wait-bg: #fff3d6;
  --display: "Fraunces", Georgia, serif; --body: "IBM Plex Sans", system-ui, sans-serif; --mono: "IBM Plex Mono", ui-monospace, monospace;
}
@media (prefers-color-scheme: dark) { :root:not([data-theme="light"]) {
  --bg: #111613; --panel: #1a211d; --ink: #e8eee9; --muted: #9aa9a0; --line: #2c3730;
  --accent: #6cc4a8; --up: #5fd17f; --down: #ff7b72; --up-bg: #183222; --down-bg: #3a1c1b; --wait-bg: #3a3218; color-scheme: dark; } }
:root[data-theme="dark"] {
  --bg: #111613; --panel: #1a211d; --ink: #e8eee9; --muted: #9aa9a0; --line: #2c3730;
  --accent: #6cc4a8; --up: #5fd17f; --down: #ff7b72; --up-bg: #183222; --down-bg: #3a1c1b; --wait-bg: #3a3218; color-scheme: dark; }
* { box-sizing: border-box; }
body { margin: 0; background: var(--bg); color: var(--ink); font: 15px/1.5 var(--body); }
.wrap { max-width: 1040px; margin: 0 auto; padding-block: 20px 40px;
  padding-left: max(16px, env(safe-area-inset-left, 0px)); padding-right: max(16px, env(safe-area-inset-right, 0px));
  display: grid; gap: 18px; }
header { display: flex; flex-wrap: wrap; align-items: baseline; justify-content: space-between; gap: 8px 16px; }
h1 { font: 600 30px/1.1 var(--display); margin: 0; letter-spacing: -0.01em; }
h2 { font: 600 19px/1.2 var(--display); margin: 0; }
.stamp { color: var(--muted); font-size: 13px; }
.pill { display: inline-block; padding: 2px 10px; border-radius: 999px; font-size: 12px; font-weight: 600;
  letter-spacing: 0.04em; text-transform: uppercase; background: var(--line); color: var(--ink); }
.pill.live { background: var(--up-bg); color: var(--up); }
.tiles { display: grid; grid-template-columns: repeat(auto-fit, minmax(150px, 1fr)); gap: 12px; }
.tile { background: var(--panel); border: 1px solid var(--line); border-radius: 10px; padding: 12px 14px; }
.tile .label { color: var(--muted); font-size: 12px; text-transform: uppercase; letter-spacing: 0.05em; }
.tile .value { font: 600 26px/1.2 var(--display); font-variant-numeric: tabular-nums; }
.tile .sub { color: var(--muted); font-size: 13px; }
.panel { background: var(--panel); border: 1px solid var(--line); border-radius: 10px; padding: 14px 16px; display: grid; gap: 10px; min-width: 0; }
.panel p { margin: 0; color: var(--muted); font-size: 14px; }
.scroll { overflow-x: auto; }
table { width: 100%; border-collapse: collapse; font-size: 14px; }
th { text-align: left; color: var(--muted); font-weight: 500; font-size: 12px; text-transform: uppercase; letter-spacing: 0.04em;
  border-bottom: 1px solid var(--line); padding: 6px 8px; white-space: nowrap; }
td { border-bottom: 1px solid var(--line); padding: 7px 8px; vertical-align: top; }
tr:last-child td { border-bottom: 0; }
td.num, th.num { text-align: right; font-family: var(--mono); font-variant-numeric: tabular-nums; white-space: nowrap; }
.up { color: var(--up); } .down { color: var(--down); }
.tag { display: inline-block; padding: 1px 8px; border-radius: 6px; font-size: 12px; font-weight: 600; white-space: nowrap; }
.tag.won { background: var(--up-bg); color: var(--up); } .tag.lost { background: var(--down-bg); color: var(--down); }
.tag.open { background: var(--wait-bg); color: var(--ink); } .tag.flat { background: var(--line); color: var(--ink); }
.side { font-weight: 600; } .side.Buy { color: var(--up); } .side.Sell { color: var(--down); } .side.Avoid { color: var(--muted); }
.reason { color: var(--muted); font-size: 13px; min-width: 220px; }
.two { display: grid; grid-template-columns: repeat(auto-fit, minmax(280px, 1fr)); gap: 18px; align-items: start; }
footer { color: var(--muted); font-size: 12px; }
a { color: var(--accent); }
"""


def status_tag(status: str) -> str:
    s = status or ""
    kind = ("won" if s.startswith("Target") else "lost" if s.startswith("Stopped") else
            "flat" if s.startswith(("Closed", "Skipped")) else "open")
    short = "Skipped (never entered)" if s.startswith("Skipped") else s
    return f'<span class="tag {kind}">{esc(short or "Waiting")}</span>'


def build(root: Path, now: datetime | None = None) -> str:
    now = now or datetime.now(ZoneInfo(UK))
    scan = scan_today(root / "ftse_scan_latest.xlsx")
    day = scan["day"] or now.strftime("%Y-%m-%d")
    picker = read_csv(root / "scan_data" / "picker" / f"{day}.csv")
    views = {r["ticker"]: r for _, r in picker.iterrows()} if not picker.empty else {}
    moves_today = read_csv(root / "scan_data" / "daily_moves" / f"{day}.csv")
    moves_day, moves_prev = latest_moves_before(root / "scan_data" / "daily_moves", day)
    picker_trades = read_csv(root / "scan_data" / "picker_trades.csv")
    total = scanner_total(root / "ftse_outcomes.xlsx")
    structure = read_csv(root / "scan_data" / "structure_live" / "trades.csv")

    trades = scan["trades"]
    nets = [t.get("After costs per £100 (£)") for t in trades if isinstance(t.get("After costs per £100 (£)"), (int, float))]
    today_net = sum(nets) if nets else None
    still_open = sum(1 for t in trades if str(t.get("Status so far", "")).startswith(("Open", "Waiting", "Price")))
    won = sum(1 for t in trades if isinstance(t.get("After costs per £100 (£)"), (int, float))
              and t["After costs per £100 (£)"] > 0 and not str(t.get("Status so far", "")).startswith("Open"))
    lost = sum(1 for t in trades if isinstance(t.get("After costs per £100 (£)"), (int, float))
               and t["After costs per £100 (£)"] < 0 and not str(t.get("Status so far", "")).startswith("Open"))
    market_open = now.strftime("%Y-%m-%d") == day and now.weekday() < 5 and "08:00" <= now.strftime("%H:%M") < "16:30"

    picker_net_total = sum(float(x) for x in picker_trades["net_gbp"]) if not picker_trades.empty else None
    picker_today = picker_trades[picker_trades["day"] == day] if not picker_trades.empty else picker_trades
    structure_closed = structure[structure.get("net_measured", pd.Series(dtype=str)).astype(str).str.len() > 0] \
        if not structure.empty and "net_measured" in structure else pd.DataFrame()
    structure_net = sum(float(x) for x in structure_closed["net_measured"]) if not structure_closed.empty else None

    out = []
    out.append(f"""<title>FTSE Trader Today</title>
<meta name="viewport" content="width=device-width, initial-scale=1, viewport-fit=cover">
<meta http-equiv="refresh" content="300">
<link rel="preconnect" href="https://fonts.googleapis.com"><link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link rel="stylesheet" href="https://fonts.googleapis.com/css2?family=Fraunces:opsz,wght@9..144,600&family=IBM+Plex+Mono:wght@400;500&family=IBM+Plex+Sans:wght@400;500;600&display=swap">
<style>{STYLE}</style>
<div class="wrap">
<header>
  <div><h1>FTSE Trader Today</h1>
  <div class="stamp">{esc(datetime.strptime(day, "%Y-%m-%d").strftime("%A %d %B %Y"))} · last scan {esc(scan["generated"][11:] or "–")} UK · page built {esc(now.strftime("%H:%M"))}</div></div>
  <span class="pill {'live' if market_open else ''}">{'Market open' if market_open else 'Market closed'}</span>
</header>
<section class="tiles">
  <div class="tile"><div class="label">Scanner today</div><div class="value {tone(today_net)}">{money(today_net)}</div>
    <div class="sub">{len(trades)} trade(s): {won} won, {lost} lost, {still_open} open</div></div>
  <div class="tile"><div class="label">Picker today</div><div class="value {tone(picker_today['net_gbp'].astype(float).sum() if not picker_today.empty else None)}">{money(picker_today['net_gbp'].astype(float).sum()) if not picker_today.empty else '–'}</div>
    <div class="sub">{'scored after 17:05' if picker_today.empty else f'{len(picker_today)} call(s), open to close'}</div></div>
  <div class="tile"><div class="label">Scanner running total</div><div class="value {tone(total['net'] if total else None)}">{money(total['net']) if total else '–'}</div>
    <div class="sub">{f"{total['completed']} completed: {total['won']} won, {total['lost']} lost" if total else 'no results yet'}</div></div>
  <div class="tile"><div class="label">Picker running total</div><div class="value {tone(picker_net_total)}">{money(picker_net_total)}</div>
    <div class="sub">{len(picker_trades)} paper trade(s) since 26 Oct</div></div>
</section>""")

    # Today's trades
    rows = []
    for t in trades:
        ticker = str(t.get("Ticker") or "")
        side = str(t.get("Direction") or "")
        view = views.get(ticker)
        label = "–"
        if view is not None:
            v = view["view"]
            label = "Avoid" if v == "Avoid" else ("Agrees" if (v == "Buy") == (side == "Buy") else "Opposes")
        net = t.get("After costs per £100 (£)")
        rows.append(f"""<tr><td class="num">{esc(t.get('Signal time'))}</td><td>{esc(t.get('Company'))}</td>
<td><span class="side {esc(side)}">{esc(side)}</span> <span class="stamp">{esc(t.get('Setup'))}</span></td>
<td>{status_tag(str(t.get('Status so far') or ''))} <span class="stamp">{esc(t.get('At'))}</span></td>
<td class="num {tone(net if isinstance(net, (int, float)) else None)}">{money(net) if isinstance(net, (int, float)) else '–'}</td>
<td>{esc(label)}</td></tr>""")
    body = "".join(rows) or '<tr><td colspan="6">No trades yet today. The scanner trades only when an idea passes every rule.</td></tr>'
    out.append(f"""<section class="panel"><h2>Today's trades</h2>
<p>£100 each, paper only. Progress from completed 15-minute candles; Yahoo prices can be up to 15 minutes behind. The final result is worked out at 17:05.</p>
<div class="scroll"><table><thead><tr><th class="num">Signal</th><th>Share</th><th>Trade</th><th>Status</th><th class="num">After costs</th><th>Picker</th></tr></thead>
<tbody>{body}</tbody></table></div></section>""")

    # Picker views and yesterday's movers
    scored = {}
    if not moves_today.empty:
        scored = {r["ticker"]: r for _, r in moves_today.iterrows()}
    prow = []
    for ticker, r in views.items():
        m = scored.get(ticker)
        result = "–"
        if m is not None and r["view"] in ("Buy", "Sell") and m.get("picker_net_gbp", "") != "":
            v = float(m["picker_net_gbp"])
            result = f'<span class="{tone(v)}">{money(v)}</span>'
        prow.append(f"""<tr><td>{esc(r.get('company'))}</td><td><span class="side {esc(r['view'])}">{esc(r['view'])}</span></td>
<td class="num">{result}</td><td class="reason">{esc(r.get('reason'))}</td></tr>""")
    pbody = "".join(prow) or '<tr><td colspan="4">No picker list for today yet (it runs at about 07:06 on weekdays).</td></tr>'

    mrows = []
    if not moves_prev.empty:
        moves_prev["chg"] = pd.to_numeric(moves_prev["change_pct"], errors="coerce")
        moves_prev = moves_prev.dropna(subset=["chg"]).sort_values("chg", ascending=False)
        picks = pd.concat([moves_prev.head(5), moves_prev.tail(5)])
        for _, m in picks.iterrows():
            mrows.append(f'<tr><td>{esc(m["company"])}</td><td class="num {tone(m["chg"])}">{m["chg"]:+.1f}%</td></tr>')
    mbody = "".join(mrows) or '<tr><td colspan="2">Not available yet.</td></tr>'
    prev_label = datetime.strptime(moves_day, "%Y-%m-%d").strftime("%A %d %B") if moves_day else "previous day"
    out.append(f"""<div class="two">
<section class="panel"><h2>Morning picker</h2><p>Labels only: the picker does not change the scanner's trades. Each Buy or Sell is also paper-traded with £100 from the open to the close (shown after 17:05).</p>
<div class="scroll"><table><thead><tr><th>Share</th><th>View</th><th class="num">Result</th><th>Reason</th></tr></thead><tbody>{pbody}</tbody></table></div></section>
<section class="panel"><h2>Biggest movers, {esc(prev_label)}</h2><p>FTSE 100, change on the previous close.</p>
<div class="scroll"><table><thead><tr><th>Share</th><th class="num">Move</th></tr></thead><tbody>{mbody}</tbody></table></div></section>
</div>""")

    out.append(f"""<section class="panel"><h2>Other models</h2><div class="scroll"><table>
<thead><tr><th>Model</th><th class="num">Trades</th><th class="num">Closed, after costs</th></tr></thead><tbody>
<tr><td>Daily structure pullback (£250 a trade)</td><td class="num">{len(structure)}</td><td class="num {tone(structure_net)}">{money(structure_net)}</td></tr>
</tbody></table></div><p>The monthly momentum model trades at the end of each month; see ftse_momentum_live.xlsx in the repo.</p></section>
<footer>Paper trading research only: no real orders and not financial advice. The page refreshes itself every 5 minutes and is rebuilt after each scan and the 17:05 results.
Files: <a href="https://github.com/philwright210458-blip/ftse-trader">github.com/philwright210458-blip/ftse-trader</a></footer>
</div>""")
    return "<!doctype html><html lang=\"en-GB\"><head><meta charset=\"utf-8\">" + out[0].split("<div class=\"wrap\">")[0] + \
        "</head><body><div class=\"wrap\">" + out[0].split("<div class=\"wrap\">", 1)[1] + "".join(out[1:]) + "</body></html>"


def main() -> None:
    parser = argparse.ArgumentParser(description="Build the FTSE Trader Today page")
    parser.add_argument("--out", default="site")
    args = parser.parse_args()
    folder = Path(args.out)
    folder.mkdir(parents=True, exist_ok=True)
    (folder / "index.html").write_text(build(ROOT), encoding="utf-8")
    print(f"Page written: {folder / 'index.html'}")


if __name__ == "__main__":
    main()
