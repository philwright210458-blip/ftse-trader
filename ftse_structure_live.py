"""Paper trader for the daily structure pullback (ftse_structure.py, idea 24), from 22 Oct 2026.

Runs every weekday evening in the FTSE Outcomes workflow (17:05), after the close:

* Signals: checks every FTSE 100 and FTSE 250 share (no investment trusts) on
  the day's daily candle, with exactly the backtested rules, using only prices
  up to that day. A signal becomes a £250 paper trade unless that share already
  has one open. Days missed (e.g. a skipped run) are caught up the next time,
  still using only the prices that existed on each day.
* Entry: the next trading day's opening price.
* Exits, tracked separately on the same entry: the measured-move target (the
  report's rule) and a 2x-the-stop target. Stop counted first if both fall in one
  day; gaps fill at the open; anything still open after 60 trading days is
  closed at that day's close.
* Costs per £250 (IG share dealing): £1.50 (25p price gap + £1.25 stamp duty).

Results: ftse_structure_live.xlsx; trades kept in scan_data/structure_live/trades.csv.
Paper trading only: no real orders are placed.
"""

from __future__ import annotations

import argparse
import json
import math
import shutil
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

import ftse_ma_daily_backtest as db
import ftse_structure as st
import ftse_trade_scanner as sc
from ftse_sessions import calendar

FIRST_SIGNAL_DAY = "2026-10-22"
STAKE, COSTS = 250.0, 1.50
DATA_DIR = Path(__file__).parent / "scan_data" / "structure_live"
WORKBOOK = Path(__file__).parent / "ftse_structure_live.xlsx"
COLUMNS = ["signal_day", "recorded_at", "ticker", "company", "index", "impulse_low", "impulse_high", "pullback_low",
           "depth", "stop", "target_measured", "fill_day", "fill", "target_2x",
           "exit_day_measured", "exit_measured", "result_measured", "net_measured",
           "exit_day_2x", "exit_2x", "result_2x", "net_2x", "recording"]


def read_trades(folder: Path) -> pd.DataFrame:
    path = folder / "trades.csv"
    if not path.exists() or path.stat().st_size == 0:
        return pd.DataFrame(columns=COLUMNS)
    return pd.read_csv(path, dtype={"signal_day": str, "fill_day": str, "exit_day_measured": str, "exit_day_2x": str})


def read_state(folder: Path) -> dict:
    path = folder / "state.json"
    return json.loads(path.read_text()) if path.exists() else {}


def is_open(row: pd.Series) -> bool:
    measured_done = isinstance(row.get("result_measured"), str) and row["result_measured"] != ""
    twox_done = isinstance(row.get("result_2x"), str) and row["result_2x"] != ""
    return not (measured_done and twox_done)


def exit_for(frame: pd.DataFrame, fill_n: int, stop: float, target: float) -> tuple | None:
    """(exit day, price, result) once known from completed daily candles; None while still open."""
    o, h, l, c = (frame[k].to_numpy() for k in ("open", "high", "low", "close"))
    last = fill_n - 1 + st.MAX_HOLD                              # 60 trading days after the signal day
    expected = frame.index[fill_n]
    for k in range(fill_n, min(last, len(frame) - 1) + 1):
        if frame.index[k] != expected:
            return None
        values = (o[k], h[k], l[k], c[k])
        if not all(math.isfinite(v) and v > 0 for v in values) or l[k] > min(o[k], c[k]) or h[k] < max(o[k], c[k]):
            return None
        expected = calendar().next_session(expected)
        if l[k] <= stop:
            return frame.index[k], min(o[k], stop), "Stopped out"
        if h[k] >= target:
            return frame.index[k], max(o[k], target), "Target hit"
    if last <= len(frame) - 1:
        return frame.index[last], c[last], "Closed after 60 days"
    return None


def resolve_trades(trades: pd.DataFrame, data: dict[str, pd.DataFrame], today: pd.Timestamp) -> None:
    """Resolve only recorded positions, preserving already final results."""
    for i, row in trades.iterrows():
        if not is_open(row):
            continue
        frame = data.get(row["ticker"])
        if frame is None or frame.empty:
            continue
        frame = frame[frame.index <= today].sort_index()
        frame = frame[~frame.index.duplicated(keep="last")]
        frame = frame[frame.index.isin(calendar().sessions)]
        signal = pd.Timestamp(row["signal_day"])
        if not isinstance(row["fill_day"], str) or not row["fill_day"]:
            later = frame[frame.index > signal]
            if later.empty:
                continue                                          # waiting for the next morning
            expected = calendar().next_session(signal)
            if later.index[0] != expected:
                continue  # the actual next session is missing; never fill at a later open
            fill = float(later["open"].iloc[0])
            if not math.isfinite(fill) or fill <= 0:
                continue
            if not float(row["stop"]) < fill < float(row["target_measured"]):
                trades.loc[i, ["fill_day", "fill", "result_measured", "result_2x"]] = [
                    str(later.index[0].date()), round(fill, 4), "Skipped: opened past stop or target",
                    "Skipped: opened past stop or target"]
                continue
            trades.loc[i, ["fill_day", "fill", "target_2x"]] = [
                str(later.index[0].date()), round(fill, 4), round(fill + 2 * (fill - float(row["stop"])), 4)]
        if pd.Timestamp(trades.at[i, "fill_day"]) not in frame.index:
            continue
        fill_n = frame.index.get_loc(pd.Timestamp(trades.at[i, "fill_day"]))
        fill = float(trades.at[i, "fill"])
        for label, target in (("measured", float(trades.at[i, "target_measured"])), ("2x", float(trades.at[i, "target_2x"]))):
            if isinstance(trades.at[i, f"result_{label}"], str) and trades.at[i, f"result_{label}"]:
                continue
            done = exit_for(frame, fill_n, float(trades.at[i, "stop"]), target)
            if done:
                day, price, result = done
                gross = (price - fill) / fill * STAKE
                trades.loc[i, [f"exit_day_{label}", f"exit_{label}", f"result_{label}", f"net_{label}"]] = [
                    str(day.date()), round(price, 4), result, round(gross - COSTS, 2)]


def open_on(row: pd.Series, day: pd.Timestamp) -> bool:
    """Whether a recorded decision occupied this share at the end of a given day."""
    if pd.Timestamp(row["signal_day"]) > day:
        return False
    for label in ("measured", "2x"):
        result = row.get(f"result_{label}")
        end = row.get("fill_day") if isinstance(result, str) and result.startswith("Skipped") else row.get(f"exit_day_{label}")
        if not isinstance(end, str) or not end or pd.Timestamp(end) > day:
            return True
    return False


def step(trades: pd.DataFrame, data: dict[str, pd.DataFrame], names: dict, indexes: dict,
         checked_through: dict[str, str] | str | None, today: pd.Timestamp,
         now_text: str) -> tuple[pd.DataFrame, dict[str, str]]:
    """Catch up each share in order. A legacy global date is not proof of coverage.

    Existing decisions take precedence over reconstructed signals. A share's
    progress stops at the first missing or invalid session and resumes on recovery.
    """
    trades = trades.copy()
    for column in COLUMNS:
        if column not in trades.columns:
            trades[column] = None
        trades[column] = trades[column].astype(object)
    trades["recording"] = trades["recording"].fillna("Recorded before reliability update")
    checked = dict(checked_through) if isinstance(checked_through, dict) else {}
    resolve_trades(trades, data, today)
    for ticker, original in data.items():
        frame = original[original.index <= today].sort_index()
        frame = frame[~frame.index.duplicated(keep="last")]
        if frame.empty:
            continue
        start = max(pd.Timestamp(FIRST_SIGNAL_DAY), pd.Timestamp(checked[ticker]) + pd.Timedelta(days=1)) \
            if ticker in checked else pd.Timestamp(FIRST_SIGNAL_DAY)
        if start > today:
            continue
        last = None
        for day in calendar().sessions_in_range(start, today):
            if day not in frame.index:
                break
            bar = frame.loc[day]
            o, h, l, c = (float(bar[k]) for k in ("open", "high", "low", "close"))
            if not all(math.isfinite(v) and v > 0 for v in (o, h, l, c)) or l > min(o, c) or h < max(o, c):
                break
            last = day
        if last is None:
            continue
        visible = frame[frame.index <= last]
        for signal in st.signals(visible):
            day = visible.index[signal["t"]]
            if day < start:
                continue
            here = trades[trades["ticker"] == ticker]
            if (here["signal_day"] == str(day.date())).any() or any(open_on(r, day) for _, r in here.iterrows()):
                continue
            row = {"signal_day": str(day.date()), "recorded_at": now_text, "ticker": ticker,
                   "company": names.get(ticker, ticker), "index": indexes.get(ticker, ""),
                   "impulse_low": round(signal["L"], 4), "impulse_high": round(signal["H"], 4),
                   "pullback_low": round(signal["P"], 4), "depth": round(signal["depth"], 3),
                   "stop": round(signal["stop"], 4), "target_measured": round(signal["target"], 4),
                   "recording": "Recovered after missing data" if day < today else "Recorded on signal day"}
            candidate = pd.DataFrame([row]).reindex(columns=COLUMNS).astype(object)
            resolve_trades(candidate, {ticker: visible}, last)
            # During migration an old recorded trade may exist after this recovered
            # signal. Never invent an overlapping position or replace that decision.
            later = here[here["signal_day"] > str(day.date())]
            if any(open_on(candidate.iloc[0], pd.Timestamp(d)) for d in later["signal_day"]):
                print(f"Recovery: {ticker} {day.date()} overlaps a recorded decision; left out.")
                continue
            trades = candidate if trades.empty else pd.concat([trades, candidate], ignore_index=True)
        checked[ticker] = str(last.date())
    return trades, checked


def build_workbook(trades: pd.DataFrame, data: dict, today: pd.Timestamp, now_text: str) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.cell(1, 1, "Structure pullback (idea 24): paper trading, £250 a trade").font = Font(name="Arial", bold=True, size=13)
    filled = trades[trades["fill"].notna() & ~trades["result_measured"].fillna("").astype(str).str.startswith("Skipped")] \
        if len(trades) else trades
    rows = [("Updated", now_text), ("Prices up to", str(today.date())), ("Started", FIRST_SIGNAL_DAY),
            ("Signals so far", len(trades)), ("Trades entered", len(filled)),
            ("Recovered signals (included in totals)", int(trades["recording"].eq("Recovered after missing data").sum())
             if "recording" in trades else 0)]
    for label, name in (("measured", "Measured-move target"), ("2x", "2x stop target")):
        closed = filled[filled[f"net_{label}"].notna()] if len(filled) else filled
        total = float(closed[f"net_{label}"].astype(float).sum()) if len(closed) else 0.0
        rows += [(f"{name}: closed trades", len(closed)),
                 (f"{name}: total after costs (£)", round(total, 2)),
                 (f"{name}: average per trade (£)", round(total / len(closed), 2) if len(closed) else "none yet")]
    for r, (label, value) in enumerate(rows, 3):
        ws.cell(r, 1, label).font = Font(name="Arial", bold=True)
        ws.cell(r, 2, value).font = Font(name="Arial")
    ws.cell(len(rows) + 4, 1, "Rules as backtested: an impulse of large up days in an uptrend, a 38.2-61.8% pullback, then a "
                              "strong up day closing above the previous 3 days' highs. Bought at the next open; stop below the "
                              "pullback. Two exits tracked on the same trade. Costs £1.50 per £250 (IG share dealing). Only "
                              "about 6-20 signals a year are expected, so judge it over many months. Paper trading only.")
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 44, 30

    log = wb.create_sheet("Trades")
    shown = trades.rename(columns={
        "signal_day": "Signal day", "recorded_at": "Recorded", "ticker": "Share", "company": "Company", "index": "Index",
        "impulse_low": "Impulse low", "impulse_high": "Impulse high", "pullback_low": "Pullback low", "depth": "Pullback depth",
        "stop": "Stop", "target_measured": "Target (measured move)", "fill_day": "Bought on", "fill": "Entry price",
        "target_2x": "Target (2x stop)", "exit_day_measured": "Exit day (measured)", "exit_measured": "Exit price (measured)",
        "result_measured": "Result (measured)", "net_measured": "After costs, measured (£)", "exit_day_2x": "Exit day (2x)",
        "exit_2x": "Exit price (2x)", "result_2x": "Result (2x)", "net_2x": "After costs, 2x (£)", "recording": "When recorded"})
    log.append(list(shown.columns))
    for _, row in shown.iterrows():
        values = [None if (isinstance(v, float) and not math.isfinite(v)) else v for v in row.tolist()]
        log.append(values)
        for col, name in enumerate(shown.columns, 1):
            if "(£)" in name and isinstance(values[col - 1], (int, float)):
                log.cell(log.max_row, col).fill = PatternFill("solid", fgColor="D9EAD3" if values[col - 1] > 0 else "F4CCCC")
    for col in range(1, len(shown.columns) + 1):
        log.column_dimensions[log.cell(1, col).column_letter].width = 15
    log.column_dimensions["D"].width = 28
    log.freeze_panes = "A2"
    return wb


def run(args, now: pd.Timestamp | None = None) -> int:
    now = now or pd.Timestamp.now(tz=sc.UK_TZ)
    universe = [c for c in sc.load_universe(Path(args.universe)) if c.kind != "trust"]
    names = {c.yahoo_ticker: c.company for c in universe}
    indexes = {c.yahoo_ticker: c.index for c in universe}
    print(f"Structure paper trader: downloading daily candles for {len(universe)} shares...")
    data = db.download([c.yahoo_ticker for c in universe], args.period)
    if not data:
        print("No data downloaded (is Yahoo reachable?). Nothing changed.")
        return 1
    today = now.tz_convert(sc.UK_TZ).tz_localize(None).normalize()
    if not calendar().is_session(today) or now.tz_convert(sc.UK_TZ).strftime("%H:%M") < "16:45":
        today = today - pd.Timedelta(days=1)                      # today's candle isn't final yet: use up to yesterday
    folder = Path(args.data_dir)
    folder.mkdir(parents=True, exist_ok=True)
    state = read_state(folder)
    if state.get("schema_version") != 2:
        for name in ("trades.csv", "state.json"):
            source = folder / name
            backup = source.with_name(source.stem + ".pre-reliability" + source.suffix)
            if source.exists() and not backup.exists():
                shutil.copy2(source, backup)
    trades, checked = step(read_trades(folder), data, names, indexes,
                           state.get("checked_by_ticker", state.get("checked_through")), today,
                           now.tz_convert(sc.UK_TZ).strftime("%Y-%m-%d %H:%M"))
    trade_tmp = folder / "trades.csv.tmp"
    trades.reindex(columns=COLUMNS).to_csv(trade_tmp, index=False)
    trade_tmp.replace(folder / "trades.csv")
    state_tmp = folder / "state.json.tmp"
    state_tmp.write_text(json.dumps({"schema_version": 2, "checked_by_ticker": checked}, indent=2))
    state_tmp.replace(folder / "state.json")
    saved = sc.save_workbook(build_workbook(trades, data, today, now.tz_convert(sc.UK_TZ).strftime("%Y-%m-%d %H:%M UK")),
                             Path(args.out))
    outstanding = [ticker for ticker in names if checked.get(ticker, "") < str(today.date())]
    if outstanding:
        print(f"Warning: {len(outstanding)} share(s) still have missing sessions; progress retained for recovery: "
              + ", ".join(outstanding[:10]))
    print(f"{len(trades)} paper trade(s) on record; progress saved separately for {len(checked)} shares. Workbook saved: {saved}")
    return 0


def main() -> None:
    parser = argparse.ArgumentParser(description="Paper-trade the daily structure pullback")
    parser.add_argument("--universe", default=str(sc.UNIVERSE_PATH))
    parser.add_argument("--data-dir", default=str(DATA_DIR))
    parser.add_argument("--out", default=str(WORKBOOK))
    parser.add_argument("--period", default="2y")
    sys.exit(run(parser.parse_args()))


if __name__ == "__main__":
    main()
